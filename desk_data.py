"""
desk_data.py — AI-Prowler Service Desk, Phase 3: export and import.

Export (owner: everything; manager: the tickets they can see; private
tickets only for the owner — all via desk_db.list_tickets/can()):
  xlsx   Tickets, Messages and Events sheets
  csv    one CSV per table, zipped
  json   full export: every ticket.json + its attachments + contacts, zipped —
         imports into another Service Desk (e.g. Solo -> Business upgrade)
  pdf    one ticket: fields, public thread, resolution (internal notes optional)
Exports are one-way views: editing an exported file never changes live data.

Import (owner only, dry run by default):
  csv / xlsx  rows -> tickets, with an optional column map for other help desks
  json        a full export (.zip) from another Service Desk
  Every real import takes a backup first, gets a batch_id, sends NO alerts and
  runs no routing, and can be undone as a whole for 7 days.
"""
from __future__ import annotations

import csv
import datetime
import io
import json
import re
import secrets
import shutil
import zipfile
from pathlib import Path

import desk_db
import desk_files
import desk_rules as rules
from db_access import utcnow_iso

UNDO_DAYS = 7
TICKET_COLS = ("ticket_id", "title", "description", "kind", "type", "priority", "state",
               "assignee", "queue_id", "contact_id", "customer_id", "job_id", "source",
               "due_resolve_at", "first_response_at", "fixed_at", "closed_at", "close_reason",
               "csat_score", "created_at", "last_edited_at")
MSG_COLS = ("ticket_id", "at", "author", "author_kind", "visibility", "channel", "body",
            "delivery_status")
EVENT_COLS = ("ticket_id", "at", "actor", "kind", "from_value", "to_value", "reason")


def _scoped(db_path, actor, edition, f: dict) -> list[dict]:
    rows = desk_db.list_tickets(db_path, actor=actor, edition=edition, state=f.get("state") or None,
                                queue_id=f.get("queue_id") or None, kind=f.get("kind") or None,
                                customer_id=f.get("customer_id") or None, limit=1_000_000)
    since, until, typ = f.get("since") or "", f.get("until") or "", f.get("type") or ""
    return [t for t in rows if (not since or t["created_at"] >= since)
            and (not until or t["created_at"] < until) and (not typ or t["type"] == typ)]


def _children(db_path, ids: list[str], table: str, cols, internal_ok: bool) -> list[dict]:
    if not ids:
        return []
    conn = desk_db.get_connection(db_path)
    try:
        out = []
        for i in range(0, len(ids), 500):
            chunk = ids[i:i + 500]
            q = (f"SELECT {', '.join(cols)}{', visibility' if table == 'ticket_messages' and 'visibility' not in cols else ''} "
                 f"FROM {table} WHERE ticket_id IN ({', '.join('?' for _ in chunk)}) ORDER BY ticket_id")
            out += [dict(r) for r in conn.execute(q, chunk)]
    finally:
        conn.close()
    if table == "ticket_messages" and not internal_ok:
        out = [m for m in out if m["visibility"] == "public"]
    return out


def export(db_path: str, fmt: str, *, actor=None, edition: str = rules.BUSINESS,
           filters: dict | None = None, ticket_id: str = "", include_internal: bool = False) -> tuple[str, bytes, str]:
    """Returns (filename, bytes, mime)."""
    d = rules.can(actor, "export", edition=edition)
    if not d:
        raise desk_db.DeskError(d.reason)
    fmt = (fmt or "").lower().lstrip(".")
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M")
    if fmt == "pdf":
        doc = desk_db.get_ticket_detail(db_path, (ticket_id or "").upper(), actor=actor, edition=edition)
        return f"{doc['ticket_id']}.pdf", ticket_pdf(doc, include_internal=include_internal), "application/pdf"
    rows = _scoped(db_path, actor, edition, filters or {})
    ids = [t["ticket_id"] for t in rows]
    if fmt == "json":
        if rules.normalize_actor(actor)["role"] != rules.OWNER:
            raise desk_db.DeskError("only the owner can make a full export")
        return f"service-desk-full-{stamp}.zip", _full_zip(db_path, ids), "application/zip"
    msgs = _children(db_path, ids, "ticket_messages", MSG_COLS, include_internal)
    events = _children(db_path, ids, "ticket_events", EVENT_COLS, True)
    tables = [("Tickets", TICKET_COLS, rows), ("Messages", MSG_COLS, msgs), ("Events", EVENT_COLS, events)]
    if fmt == "csv":
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name, cols, data in tables:
                s = io.StringIO()
                w = csv.writer(s)
                w.writerow(cols)
                for r in data:
                    w.writerow([r.get(c) if r.get(c) is not None else "" for c in cols])
                z.writestr(f"{name.lower()}.csv", s.getvalue().encode("utf-8-sig"))
        return f"service-desk-{stamp}-csv.zip", buf.getvalue(), "application/zip"
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        wb.remove(wb.active)
        for name, cols, data in tables:
            ws = wb.create_sheet(name)
            ws.append(list(cols))
            for c in ws[1]:
                c.font = Font(bold=True)
            for r in data:
                ws.append([r.get(c) for c in cols])
            ws.freeze_panes = "A2"
        buf = io.BytesIO()
        wb.save(buf)
        return f"service-desk-{stamp}.xlsx", buf.getvalue(), \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    raise desk_db.DeskError("format must be xlsx, csv, json or pdf")


def _full_zip(db_path: str, ids: list[str]) -> bytes:
    root = desk_db.root_for(db_path)
    buf = io.BytesIO()
    conn = desk_db.get_connection(db_path)
    try:
        contacts = [dict(r) for r in conn.execute("SELECT * FROM customer_contacts")]
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for tid in ids:
                row = desk_db._row(conn, tid)
                folder = desk_files.ticket_folder(root, row)
                z.writestr(f"tickets/{tid}/ticket.json",
                           json.dumps(desk_files.build_ticket_json(conn, tid), indent=2, ensure_ascii=False))
                adir = folder / desk_files.ATTACHMENTS_DIR
                if adir.exists():
                    for f in adir.iterdir():
                        if f.is_file():
                            z.write(f, f"tickets/{tid}/attachments/{f.name}")
            z.writestr("contacts.json", json.dumps(contacts, indent=2, ensure_ascii=False))
            z.writestr("manifest.json", json.dumps({"format": "ai-prowler-service-desk-full", "version": 1,
                                                    "tickets": len(ids), "exported_at": utcnow_iso()}))
    finally:
        conn.close()
    return buf.getvalue()


# ── one-ticket PDF (built-in writer: no extra dependency) ────────────────────

def _pdf_escape(s: str) -> str:
    s = s.encode("latin-1", "replace").decode("latin-1")
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _wrap(text: str, width: int) -> list[str]:
    out = []
    for para in (text or "").splitlines() or [""]:
        line = ""
        for word in para.split(" "):
            if len(line) + len(word) + 1 > width:
                out.append(line)
                line = word
            else:
                line = f"{line} {word}".strip()
        out.append(line)
    return out


def text_pdf(lines: list[tuple[str, int]]) -> bytes:
    """Tiny PDF: list of (text, font_size) lines, Helvetica, Letter pages."""
    pages, cur, y = [], [], 750
    for text, size in lines:
        step = size + 4
        if y - step < 50:
            pages.append(cur)
            cur, y = [], 750
        cur.append((text, size, y))
        y -= step
    pages.append(cur)
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None,
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"]
    kids = []
    for p in pages:
        stream = "BT\n" + "".join(f"/F1 {s} Tf 1 0 0 1 50 {yy} Tm ({_pdf_escape(t)}) Tj\n" for t, s, yy in p) + "ET"
        objs.append(f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream")
        content_id = len(objs)
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>")
        kids.append(len(objs))
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def ticket_pdf(doc: dict, *, include_internal: bool = False) -> bytes:
    L = [(f"{doc['ticket_id']} - {doc['title']}", 16), ("", 8),
         (f"State: {doc['state']}   Priority: {doc['priority']}   Kind: {doc['kind']}", 10),
         (f"Created: {doc.get('created_at', '')[:16]}   Closed: {(doc.get('closed_at') or '')[:16]}", 10)]
    for k in ("customer_id", "job_id", "contact_id"):
        if doc.get(k):
            L.append((f"{k.replace('_', ' ').title()}: {doc[k]}", 10))
    if doc.get("description"):
        L += [("", 8), ("Details", 12)] + [(l, 10) for l in _wrap(doc["description"], 95)]
    msgs = [m for m in doc.get("messages") or [] if include_internal or m["visibility"] == "public"]
    if msgs:
        L += [("", 8), ("Messages", 12)]
        for m in msgs:
            L.append((f"{m['at'][:16].replace('T', ' ')}  {m['author']} ({m['visibility']}, {m['channel']})", 9))
            L += [(f"   {l}", 10) for l in _wrap(m["body"], 90)]
    if doc.get("close_reason"):
        L += [("", 8), (f"Resolution: {doc['close_reason']}", 10)]
    atts = [a for a in doc.get("attachments") or [] if include_internal or a["visibility"] == "public"]
    if atts:
        L += [("", 8), ("Attachments: " + ", ".join(a["filename"] for a in atts), 9)]
    return text_pdf(L)


# ── import ───────────────────────────────────────────────────────────────────

FIELD_ALIASES = {
    "title": ("title", "subject", "summary"), "description": ("description", "details", "body"),
    "kind": ("kind",), "priority": ("priority",), "state": ("state", "status"),
    "type": ("type", "category"), "contact_name": ("contact_name", "customer name", "requester", "name"),
    "contact_phone": ("contact_phone", "phone", "requester phone"),
    "contact_email": ("contact_email", "email", "requester email"),
    "customer_id": ("customer_id", "customerid"), "job_id": ("job_id", "jobid"),
    "due": ("due", "due date", "due_resolve_at"), "created_at": ("created_at", "created", "created date"),
    "external_id": ("external_id", "id", "ticket id", "ticket_id"),
}
STATE_MAP = {"new": rules.OPEN, "open": rules.OPEN, "pending": rules.ON_HOLD, "on-hold": rules.ON_HOLD,
             "hold": rules.ON_HOLD, "solved": rules.FIXED, "resolved": rules.FIXED, "closed": rules.CLOSED,
             "in progress": rules.IN_PROGRESS, "assigned": rules.ASSIGNED}
PRIORITY_MAP = {"urgent": "P1", "high": "P2", "normal": "P3", "medium": "P3", "low": "P4"}


def _rows_from(data: bytes, filename: str) -> list[dict]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", "replace")
        return [dict(r) for r in csv.DictReader(io.StringIO(text))]
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(it, [])]
        return [{h: ("" if v is None else str(v)) for h, v in zip(head, r)} for r in it if any(r)]
    raise desk_db.DeskError("import a .csv, .xlsx or a full-export .zip")


def _map_row(raw: dict, column_map: dict | None) -> dict:
    low = {str(k).strip().lower(): (v or "").strip() for k, v in raw.items()}
    out = {}
    for field, aliases in FIELD_ALIASES.items():
        src = (column_map or {}).get(field)
        val = low.get(src.strip().lower(), "") if src else next((low[a] for a in aliases if low.get(a)), "")
        out[field] = val
    return out


def _normalize(r: dict) -> tuple[dict, list[str]]:
    problems = []
    if not r["title"]:
        problems.append("no title")
    st = r["state"].strip()
    state = next((s for s in rules.STATES if s.lower() == st.lower()), None) or STATE_MAP.get(st.lower(), rules.OPEN if not st else None)
    if state is None:
        problems.append(f"unknown state {st!r}")
    pr = r["priority"].strip().upper()
    priority = pr if pr in rules.PRIORITIES else PRIORITY_MAP.get(pr.lower(), "P3")
    kind = (r["kind"] or ("external" if (r["contact_phone"] or r["contact_email"] or r["customer_id"]) else "internal")).lower()
    if kind not in rules.KINDS:
        problems.append(f"unknown kind {kind!r}")
    return {**r, "state": state, "priority": priority, "kind": kind}, problems


def _existing_external_ids(db_path: str) -> set:
    conn = desk_db.get_connection(db_path)
    try:
        return {r[0] for r in conn.execute("SELECT reason FROM ticket_events WHERE kind = 'imported' "
                                           "AND reason IS NOT NULL")}
    finally:
        conn.close()


def import_tickets(db_path: str, data: bytes, filename: str, *, actor=None,
                   edition: str = rules.BUSINESS, dry_run: bool = True,
                   column_map: dict | None = None, backup_fn=None) -> dict:
    d = rules.can(actor, "import", edition=edition)
    if not d:
        raise desk_db.DeskError(d.reason)
    if filename.lower().endswith(".zip"):
        return _import_full(db_path, data, actor=actor, dry_run=dry_run, backup_fn=backup_fn,
                            filename=filename)
    rows = [_map_row(r, column_map) for r in _rows_from(data, filename)]
    seen_ext = _existing_external_ids(db_path)
    plan, dupes, bad = [], [], []
    for i, raw in enumerate(rows, 2):
        r, problems = _normalize(raw)
        if problems:
            bad.append({"row": i, "problems": problems})
        elif r["external_id"] and r["external_id"] in seen_ext:
            dupes.append({"row": i, "external_id": r["external_id"]})
        else:
            if r["external_id"]:
                seen_ext.add(r["external_id"])
            plan.append(r)
    summary = {"dry_run": dry_run, "rows": len(rows), "to_add": len(plan), "duplicates": dupes,
               "invalid": bad, "new_contacts": len({(p["contact_phone"] or p["contact_email"]) for p in plan
                                                    if p["contact_phone"] or p["contact_email"]})}
    if dry_run or not plan:
        return summary
    if backup_fn:
        backup_fn("pre-import")
    batch = "IMP-" + datetime.datetime.now().strftime("%Y%m%d%H%M%S") + "-" + secrets.token_hex(2)
    who = rules.normalize_actor(actor)["id"]
    import desk_tools
    added = []
    for r in plan:
        cid = None
        if r["kind"] != rules.INTERNAL and (r["contact_phone"] or r["contact_email"]):
            cid = desk_db.upsert_contact(db_path, name=r["contact_name"], phone=r["contact_phone"] or None,
                                         email=r["contact_email"] or None,
                                         contact_type="subcontractor" if r["kind"] == rules.SUBCONTRACTOR_KIND else "customer",
                                         actor=actor)["contact_id"]
        if r["kind"] == rules.EXTERNAL and not (cid or r["customer_id"]):
            r["kind"] = rules.INTERNAL
        created = r["created_at"] or utcnow_iso()
        try:
            created = desk_db.utcnow_iso() if not r["created_at"] else \
                datetime.datetime.fromisoformat(r["created_at"].replace("Z", "+00:00")).astimezone(datetime.timezone.utc).isoformat(timespec="seconds")
        except ValueError:
            created = utcnow_iso()
        try:
            due = desk_tools.due_iso(r["due"]) if r["due"] else None
        except desk_db.DeskError:
            due = None
        now = utcnow_iso()
        with desk_db.ticket_transaction(db_path) as (conn, touched):
            cur = conn.execute(
                "INSERT INTO tickets (title, description, kind, type, priority, state, assignee, filed_by, "
                "contact_id, customer_id, job_id, source, visibility, due_resolve_at, closed_at, created_at, "
                "version, last_edited_at, last_edited_by) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'import', "
                "'normal', ?, ?, ?, 1, ?, ?)",
                (r["title"][:200], r["description"], r["kind"], r["type"], r["priority"],
                 rules.ASSIGNED if edition == rules.SOLO and r["state"] == rules.OPEN else r["state"],
                 "owner" if edition == rules.SOLO else None, who, cid,
                 r["customer_id"] or None, r["job_id"] or None, due,
                 now if r["state"] == rules.CLOSED else None, created, now, who))
            tid = desk_db.format_ticket_id(cur.lastrowid)
            conn.execute("UPDATE tickets SET ticket_id = ?, folder = ? WHERE row_id = ?",
                         (tid, f"tickets/{created[:4]}/{tid}", cur.lastrowid))
            desk_db._event(conn, tid, who, "created", to_value=r["state"], at=created)
            desk_db._event(conn, tid, who, "imported", to_value=batch, reason=r["external_id"] or None)
            touched.add(tid)
        added.append(tid)
    conn = desk_db.get_connection(db_path)
    try:
        conn.execute("INSERT INTO import_batches (batch_id, at, actor, filename, count) VALUES (?, ?, ?, ?, ?)",
                     (batch, utcnow_iso(), who, filename, len(added)))
    finally:
        conn.close()
    return {**summary, "batch_id": batch, "added": added}


def _import_full(db_path, data, *, actor, dry_run, backup_fn, filename) -> dict:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        manifest = json.loads(z.read("manifest.json"))
    except (zipfile.BadZipFile, KeyError, ValueError):
        raise desk_db.DeskError("that .zip is not a Service Desk full export")
    if manifest.get("format") != "ai-prowler-service-desk-full":
        raise desk_db.DeskError("that .zip is not a Service Desk full export")
    docs = [json.loads(z.read(n)) for n in z.namelist() if re.fullmatch(r"tickets/T-\d+/ticket\.json", n)]
    existing = {t["ticket_id"] for t in desk_db.list_tickets(db_path, limit=1_000_000)}
    plan = [d for d in docs if d["ticket_id"] not in existing]
    summary = {"dry_run": dry_run, "rows": len(docs), "to_add": len(plan),
               "duplicates": [{"ticket_id": d["ticket_id"]} for d in docs if d["ticket_id"] in existing],
               "invalid": []}
    if dry_run or not plan:
        return summary
    if backup_fn:
        backup_fn("pre-import")
    batch = "IMP-" + datetime.datetime.now().strftime("%Y%m%d%H%M%S") + "-" + secrets.token_hex(2)
    root = desk_db.root_for(db_path)
    conn = desk_db.get_connection(db_path)
    try:
        for c in json.loads(z.read("contacts.json")) if "contacts.json" in z.namelist() else []:
            conn.execute("INSERT OR IGNORE INTO customer_contacts (contact_id, contact_type, name, phone, email, "
                         "verified_at, linked_customer_id, created_at, version, last_edited_at, last_edited_by) "
                         "VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)",
                         (c["contact_id"], c.get("contact_type", "customer"), c.get("name", ""), c.get("phone"),
                          c.get("email"), c.get("verified_at"), c.get("linked_customer_id"),
                          c.get("created_at") or utcnow_iso(), utcnow_iso(), "import"))
        for doc in plan:
            desk_files._insert_from_doc(conn, doc)
            desk_db._event(conn, doc["ticket_id"], rules.normalize_actor(actor)["id"], "imported",
                           to_value=batch)
            folder = desk_files.ticket_folder(root, doc)
            for n in z.namelist():
                pre = f"tickets/{doc['ticket_id']}/attachments/"
                if n.startswith(pre) and not n.endswith("/"):
                    dest = folder / desk_files.ATTACHMENTS_DIR / Path(n).name
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    dest.write_bytes(z.read(n))
            desk_files.write_ticket_json(root, conn, doc["ticket_id"])
        conn.execute("INSERT INTO import_batches (batch_id, at, actor, filename, count) VALUES (?, ?, ?, ?, ?)",
                     (batch, utcnow_iso(), rules.normalize_actor(actor)["id"], filename, len(plan)))
    finally:
        conn.close()
    return {**summary, "batch_id": batch, "added": [d["ticket_id"] for d in plan]}


def undo_import(db_path: str, batch_id: str, *, actor=None, edition: str = rules.BUSINESS,
                now: datetime.datetime | None = None) -> dict:
    d = rules.can(actor, "import", edition=edition)
    if not d:
        raise desk_db.DeskError(d.reason)
    now = now or datetime.datetime.now(datetime.timezone.utc)
    conn = desk_db.get_connection(db_path)
    try:
        b = conn.execute("SELECT * FROM import_batches WHERE batch_id = ?", (batch_id,)).fetchone()
        if b is None or b["undone_at"]:
            raise desk_db.DeskError(f"no import batch {batch_id} to undo")
        at = datetime.datetime.fromisoformat(b["at"].replace("Z", "+00:00"))
        if now - at > datetime.timedelta(days=UNDO_DAYS):
            raise desk_db.DeskError(f"imports can be undone for {UNDO_DAYS} days; {batch_id} is older")
        ids = [r[0] for r in conn.execute("SELECT ticket_id FROM ticket_events WHERE kind = 'imported' "
                                          "AND to_value = ?", (batch_id,))]
        root = desk_db.root_for(db_path)
        conn.execute("BEGIN IMMEDIATE")
        folders = []
        for tid in ids:
            row = desk_db._row(conn, tid)
            if row:
                folders.append(desk_files.ticket_folder(root, row))
            for table in ("ticket_attachments", "ticket_messages", "ticket_events", "ticket_links", "tickets"):
                conn.execute(f"DELETE FROM {table} WHERE ticket_id = ?", (tid,))
        conn.execute("UPDATE import_batches SET undone_at = ? WHERE batch_id = ?", (utcnow_iso(), batch_id))
        conn.execute("COMMIT")
    finally:
        conn.close()
    for f in folders:
        shutil.rmtree(f, ignore_errors=True)
    return {"ok": True, "removed": ids}

"""
db_mileage_ops.py — Mileage tracker for the Jobs app (Personal + Server mode).

Spec: "Jobs App Mileage Tracking for Tax Deductions" (2026-10-24).

Builds an IRS-ready business mileage log from data the Jobs app already
captures (clock-in/out GPS, job addresses, Route Origin Mode, home and
Start/End addresses) plus trips people add by hand, text, or voice.

IRS rules enforced here (Pub. 463 ch. 4-5; Form 4562 Part V):
  * Every trip has a date, a destination, a specific business purpose, and
    miles (amount, time, place, purpose). Vague purposes are refused.
  * created_at is stamped on every trip ("at or near the time"); trips
    added or edited more than LATE_ENTRY_DAYS after the trip are flagged.
  * Commuting is separated from business by Route Origin Mode:
      Jobs Only        -> home office rule: home legs are business.
      Company Location -> home <-> company legs are commute.
  * The standard mileage rate is applied by trip date (2026 has two rates).
  * Per-vehicle yearly totals (business, commute, other personal, total)
    come from trips + odometer readings, for Form 4562 Part V.
  * Trips are never hard-deleted (void with reason) and every change is
    written to mileage_trip_edits with who, when, and why.
  * Five or more business-owned vehicles used the same day is flagged
    (standard rate not allowed); crew-owned vehicles don't count.

Everything that reaches the network (OSRM routing, Nominatim geocoding,
home-address lookup) goes through the module-level hooks ROUTER, GEOCODER,
and HOME_RESOLVER so tests can replace them without monkeypatching
requests.
"""

from __future__ import annotations

import datetime as _dt
import hashlib
import json
import re

from db_access import get_connection, utcnow_iso

# ── Constants ──────────────────────────────────────────────────────────────

LATE_ENTRY_DAYS = 7

# IRS business standard mileage rates by effective date (cents per mile).
# 2025: Pub. 463 (2025). 2026: IR-2025-128 (Jan 1) and IRB 2026-29 (Jul 1).
DEFAULT_RATES = [
    {"from": "2025-01-01", "cents": 70.0},
    {"from": "2026-01-01", "cents": 72.5},
    {"from": "2026-07-01", "cents": 76.0},
]

CATEGORIES = ["Job", "Supplies", "Equipment", "Bank", "Estimate/Quote",
              "Client meeting", "Office errand", "Training", "Other"]

# Purposes the IRS would treat as no purpose at all.
_VAGUE_PURPOSES = {"", "business", "work", "job", "trip", "errand", "misc",
                   "miscellaneous", "other", "driving", "drive", "travel",
                   "business trip", "work trip", "n/a", "na", "none", "-"}

SETTING_ENABLED = "Mileage Tracking Enabled"
SETTING_RATES = "Mileage Rates"
SETTING_DEFAULT_VEHICLES = "Mileage Default Vehicles"   # JSON {user_key: vehicle_id}

_EDITABLE_FIELDS = {
    "trip_date", "start_time", "end_time", "vehicle_id", "from_label",
    "to_label", "miles", "classification", "category", "purpose",
    "parking_usd", "tolls_usd", "receipt_photo", "notes", "job_id",
    "customer_name",
}


# ── Network hooks (replaced in tests) ──────────────────────────────────────

_ROUTE_CACHE: dict = {}


def _osrm_miles(a, b):
    """Driving miles between two (lat, lon) points via the public OSRM
    server (same service build_daily_route uses). None on any failure."""
    if not a or not b:
        return None
    key = (round(a[0], 5), round(a[1], 5), round(b[0], 5), round(b[1], 5))
    if key in _ROUTE_CACHE:
        return _ROUTE_CACHE[key]
    try:
        import requests
        coord = f"{a[1]},{a[0]};{b[1]},{b[0]}"
        resp = requests.get(
            f"http://router.project-osrm.org/route/v1/driving/{coord}",
            params={"overview": "false", "annotations": "false"},
            timeout=20,
        ).json()
        if resp.get("code") != "Ok":
            return None
        miles = round(resp["routes"][0]["distance"] * 0.000621371, 1)
    except Exception:
        return None
    _ROUTE_CACHE[key] = miles
    return miles


def _default_geocoder(address: str):
    try:
        from db_write_ops import _geocode
        return _geocode(address) if address else None
    except Exception:
        return None


def _default_home_resolver(crew_name: str) -> str:
    """Server mode: the crew member's own Home address (Admin -> Users).
    Personal mode / not found: the owner's home address from config.json."""
    try:
        from db_write_ops import (db_read_owner_home_address,
                                  db_read_user_home_address_by_crew_name)
        if crew_name:
            addr = db_read_user_home_address_by_crew_name(crew_name)
            if addr:
                return addr
        return db_read_owner_home_address()
    except Exception:
        return ""


ROUTER = _osrm_miles
GEOCODER = _default_geocoder
HOME_RESOLVER = _default_home_resolver


# ── Small helpers ──────────────────────────────────────────────────────────

def _today() -> str:
    return _dt.date.today().isoformat()


def _parse_date(s: str) -> str:
    s = (s or "").strip()
    if not s or s.lower() == "today":
        return _today()
    if s.lower() == "yesterday":
        return (_dt.date.today() - _dt.timedelta(days=1)).isoformat()
    return _dt.datetime.strptime(s[:10], "%Y-%m-%d").date().isoformat()


def _parse_gps(raw):
    if not raw or "," not in str(raw):
        return None
    try:
        lat_s, lon_s = str(raw).split(",", 1)
        lat, lon = float(lat_s.strip()), float(lon_s.strip())
        if abs(lat) > 90 or abs(lon) > 180:
            return None
        return (lat, lon)
    except ValueError:
        return None


def _hhmm(ts) -> str:
    """'YYYY-MM-DD HH:MM:SS' or 'HH:MM[:SS]' -> 'HH:MM'."""
    if not ts:
        return ""
    s = str(ts).strip()
    if " " in s:
        s = s.split(" ", 1)[1]
    return s[:5]


def user_key(crew_user_id: str = "", crew: str = "") -> str:
    """Stable per-driver key: the server-mode user id, else the lowercased
    crew name, else 'owner' (personal mode with no crew name)."""
    uid = (crew_user_id or "").strip()
    if uid:
        return uid
    name = (crew or "").strip().lower()
    return name or "owner"


def user_id_for_name(name: str) -> str:
    """Server mode: the users.json id of the ACTIVE user with this display
    name (case-insensitive), so a trip an owner logs "for Sam" belongs to
    Sam's account and shows on Sam's phone. '' if not found / personal mode."""
    import json as _json
    want = (name or "").strip().casefold()
    if not want:
        return ""
    try:
        from db_write_ops import _state_dir_for_lookup
        path = _state_dir_for_lookup() / "users.json"
        if not path.exists():
            return ""
        data = _json.loads(path.read_text(encoding="utf-8-sig"))
        users = data.get("users") if isinstance(data, dict) else None
        if not isinstance(users, dict):
            return ""
        for uid, rec in users.items():
            if (isinstance(rec, dict) and rec.get("status", "active") == "active"
                    and str(rec.get("name") or "").strip().casefold() == want):
                return str(rec.get("id") or uid)
    except Exception:
        pass
    return ""


def entry_key(crew_user_id: str = "", crew: str = "") -> str:
    """Driver key for a time entry. Server-mode entries always carry the
    user's id. Personal-mode entries have no id and their crew text is the
    JOB's crew (e.g. "David, Jamie"), but the person clocking in is always
    the app's owner — so every personal-mode entry belongs to 'owner'."""
    uid = (crew_user_id or "").strip()
    return uid or "owner"


def validate_purpose(purpose: str) -> str:
    """'' if acceptable; otherwise the reason it's refused. IRS: the
    business purpose must be specific, not just 'business'."""
    p = (purpose or "").strip()
    if p.lower().rstrip(".") in _VAGUE_PURPOSES or len(p) < 4:
        return ("A specific business purpose is required (for example "
                "'Supplies – ladders at Home Depot'), not just 'business'.")
    return ""


def _job_address(job) -> str:
    parts = [job["street_address"] or "", job["city"] or ""]
    st_zip = f"{job['state'] or ''} {job['zip'] or ''}".strip()
    return ", ".join(p for p in (parts[0].strip(), parts[1].strip(), st_zip) if p)


def _job_point(job):
    if job is None:
        return None
    if job["latitude"] is not None and job["longitude"] is not None:
        return (float(job["latitude"]), float(job["longitude"]))
    addr = _job_address(job)
    return GEOCODER(addr) if addr else None


# ── Settings ───────────────────────────────────────────────────────────────

def _get_setting(conn, key: str):
    row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    return row["value"] if row else None


def _set_setting(conn, key: str, value: str, actor: str):
    conn.execute(
        "INSERT INTO settings (key, value, last_edited_by, last_edited_at) VALUES (?, ?, ?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value = excluded.value, "
        "last_edited_by = excluded.last_edited_by, last_edited_at = excluded.last_edited_at",
        (key, value, actor, utcnow_iso()),
    )


def tracking_enabled(db_path: str) -> bool:
    conn = get_connection(db_path)
    try:
        val = _get_setting(conn, SETTING_ENABLED)
    finally:
        conn.close()
    return _onoff(val, True)


def _onoff(val, default: bool) -> bool:
    """Read an on/off setting. The Settings tab writes Enabled/Disabled;
    older versions wrote true/false. Blank or missing = default."""
    if val is None or str(val).strip() == "":
        return default
    return str(val).strip().lower() in ("1", "true", "yes", "on", "enabled", "enable")


def set_tracking_enabled(db_path: str, enabled: bool, actor: str) -> str:
    conn = get_connection(db_path)
    try:
        _set_setting(conn, SETTING_ENABLED, "Enabled" if enabled else "Disabled", actor)
    finally:
        conn.close()
    if enabled:
        return "✅ Mileage tracking turned ON."
    return ("✅ Mileage tracking turned OFF. The Jobs app won't ask anyone for miles and adds "
            "no new trips to the mileage log; routes still show miles driven. Turn it back on "
            "in Database → Settings → Mileage Tracking Enabled.")


def get_rates(db_path: str) -> list:
    conn = get_connection(db_path)
    try:
        raw = _get_setting(conn, SETTING_RATES)
    finally:
        conn.close()
    rates = None
    if raw:
        try:
            rates = json.loads(raw)
        except Exception:
            rates = None
    if not rates:
        rates = list(DEFAULT_RATES)
    rates = [r for r in rates if r.get("from") and r.get("cents") is not None]
    return sorted(rates, key=lambda r: r["from"])


def set_rate(db_path: str, effective_from: str, cents: float, actor: str) -> str:
    eff = _parse_date(effective_from)
    cents = float(cents)
    if not 0 < cents < 500:
        return "❌ Rate must be in cents per mile, e.g. 76 for 76¢."
    rates = [r for r in get_rates(db_path) if r["from"] != eff]
    rates.append({"from": eff, "cents": cents})
    rates.sort(key=lambda r: r["from"])
    conn = get_connection(db_path)
    try:
        _set_setting(conn, SETTING_RATES, json.dumps(rates), actor)
    finally:
        conn.close()
    return f"✅ Business mileage rate {cents:g}¢/mile effective {eff}."


def rate_cents_for(rates: list, date_str: str):
    """The rate in effect on date_str, or None if before every rate."""
    cents = None
    for r in rates:
        if r["from"] <= date_str:
            cents = float(r["cents"])
        else:
            break
    return cents


# ── VIN ────────────────────────────────────────────────────────────────────

_VIN_TRANSLIT = {
    **{str(d): d for d in range(10)},
    "A": 1, "B": 2, "C": 3, "D": 4, "E": 5, "F": 6, "G": 7, "H": 8,
    "J": 1, "K": 2, "L": 3, "M": 4, "N": 5, "P": 7, "R": 9,
    "S": 2, "T": 3, "U": 4, "V": 5, "W": 6, "X": 7, "Y": 8, "Z": 9,
}
_VIN_WEIGHTS = [8, 7, 6, 5, 4, 3, 2, 10, 0, 9, 8, 7, 6, 5, 4, 3, 2]


def normalize_vin(vin: str) -> str:
    return re.sub(r"[\s-]", "", (vin or "")).upper()


def validate_vin(vin: str, check_digit: bool = True) -> str:
    """'' if valid; otherwise a plain-language reason. 17 characters, no
    I/O/Q, and (North American VINs) a correct check digit in position 9."""
    v = normalize_vin(vin)
    if len(v) != 17:
        return f"VIN must be 17 characters (got {len(v)})."
    if re.search(r"[IOQ]", v):
        return "VIN can't contain the letters I, O, or Q (they look like 1 and 0)."
    if not re.fullmatch(r"[A-HJ-NPR-Z0-9]{17}", v):
        return "VIN can only contain letters and digits."
    if check_digit:
        total = sum(_VIN_TRANSLIT[c] * w for c, w in zip(v, _VIN_WEIGHTS))
        expect = total % 11
        expect_c = "X" if expect == 10 else str(expect)
        if v[8] != expect_c:
            return ("VIN check digit doesn't match — please re-check the VIN "
                    "(position 9 should be " + expect_c + ").")
    return ""


# ── Locks ──────────────────────────────────────────────────────────────────

def _is_locked(conn, date_str: str) -> bool:
    month, year = date_str[:7], date_str[:4]
    row = conn.execute("SELECT 1 FROM mileage_locks WHERE period IN (?, ?)", (month, year)).fetchone()
    return row is not None


def lock_period(db_path: str, period: str, actor: str, unlock: bool = False) -> str:
    period = (period or "").strip()
    if not re.fullmatch(r"\d{4}(-\d{2})?", period):
        return "❌ Period must be a month (2026-10) or a year (2026)."
    conn = get_connection(db_path)
    try:
        if unlock:
            conn.execute("DELETE FROM mileage_locks WHERE period = ?", (period,))
            return f"🔓 Mileage for {period} unlocked."
        conn.execute("INSERT OR REPLACE INTO mileage_locks (period, locked_by, locked_at) "
                     "VALUES (?, ?, ?)", (period, actor, utcnow_iso()))
        return f"🔒 Mileage for {period} locked. Trips can't be edited until it's unlocked."
    finally:
        conn.close()


# ── Vehicles & odometer ────────────────────────────────────────────────────

def _vehicle_label(v) -> str:
    if v is None:
        return "(no vehicle)"
    name = v["nickname"] or " ".join(str(x) for x in (v["year"], v["make"], v["model"]) if x)
    return f"{name or v['vehicle_id']} …{(v['vin'] or '')[-6:]}"


def add_vehicle(db_path: str, vin: str, actor: str, *, year=None, make="", model="",
                color="", plate="", nickname="", ownership="business",
                owner_user_id="", owner_name="", default_user_id="",
                default_user_name="", placed_in_service="", odometer=None,
                skip_check_digit=False) -> str:
    from db_write_ops import generate_next_id
    v = normalize_vin(vin)
    err = validate_vin(v, check_digit=not skip_check_digit)
    if err:
        return f"❌ {err}"
    ownership = (ownership or "business").strip().lower()
    if ownership not in ("business", "crew"):
        return "❌ Ownership must be 'business' or 'crew'."
    if ownership == "crew" and not (owner_user_id or owner_name):
        return "❌ A crew-owned vehicle needs its owner."
    pis = _parse_date(placed_in_service) if placed_in_service else _today()
    conn = get_connection(db_path)
    try:
        if conn.execute("SELECT 1 FROM vehicles WHERE vin = ?", (v,)).fetchone():
            return f"❌ A vehicle with VIN {v} already exists."
        conn.execute("BEGIN IMMEDIATE")
        vid = generate_next_id(conn, "vehicles", "vehicle_id", "VEH", 4)
        now = utcnow_iso()
        conn.execute(
            "INSERT INTO vehicles (vehicle_id, vin, year, make, model, color, plate, nickname, "
            "ownership, owner_user_id, owner_name, default_user_id, default_user_name, "
            "placed_in_service, active, created_by, created_at, last_edited_by, last_edited_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,1,?,?,?,?)",
            (vid, v, int(year) if str(year or "").isdigit() else None, make, model, color,
             plate, nickname, ownership, owner_user_id, owner_name, default_user_id,
             default_user_name, pis, actor, now, actor, now))
        if odometer not in (None, ""):
            conn.execute("INSERT INTO odometer_readings (vehicle_id, reading_date, odometer, kind, "
                         "entered_by, entered_at) VALUES (?,?,?,?,?,?)",
                         (vid, pis, int(float(odometer)), "added", actor, now))
        conn.execute("COMMIT")
    except Exception as exc:
        try:
            conn.execute("ROLLBACK")
        except Exception:
            pass
        return f"❌ Could not add vehicle: {exc}"
    finally:
        conn.close()
    note = "" if odometer not in (None, "") else (
        "\n   ⚠️ Record the starting odometer reading — the IRS asks for total miles per vehicle.")
    return f"✅ Added {vid} ({ownership}-owned), VIN {v}.{note}"


def update_vehicle(db_path: str, vehicle_id: str, changes: dict, actor: str) -> str:
    allowed = {"year", "make", "model", "color", "plate", "nickname", "ownership",
               "owner_user_id", "owner_name", "default_user_id", "default_user_name",
               "placed_in_service", "active"}
    sets = {k: v for k, v in (changes or {}).items() if k in allowed}
    if not sets:
        return "❌ Nothing to change."
    if "ownership" in sets and str(sets["ownership"]).lower() not in ("business", "crew"):
        return "❌ Ownership must be 'business' or 'crew'."
    conn = get_connection(db_path)
    try:
        if not conn.execute("SELECT 1 FROM vehicles WHERE vehicle_id = ?", (vehicle_id,)).fetchone():
            return f"❌ No vehicle {vehicle_id}."
        cols = ", ".join(f"{k} = ?" for k in sets)
        conn.execute(f"UPDATE vehicles SET {cols}, last_edited_by = ?, last_edited_at = ?, "
                     f"version = version + 1 WHERE vehicle_id = ?",
                     (*sets.values(), actor, utcnow_iso(), vehicle_id))
    finally:
        conn.close()
    return f"✅ Updated {vehicle_id}."


def retire_vehicle(db_path: str, vehicle_id: str, actor: str, retired_on: str = "",
                   odometer=None) -> str:
    d = _parse_date(retired_on) if retired_on else _today()
    conn = get_connection(db_path)
    try:
        if not conn.execute("SELECT 1 FROM vehicles WHERE vehicle_id = ?", (vehicle_id,)).fetchone():
            return f"❌ No vehicle {vehicle_id}."
        conn.execute("UPDATE vehicles SET active = 0, retired_on = ?, last_edited_by = ?, "
                     "last_edited_at = ?, version = version + 1 WHERE vehicle_id = ?",
                     (d, actor, utcnow_iso(), vehicle_id))
        if odometer not in (None, ""):
            conn.execute("INSERT INTO odometer_readings (vehicle_id, reading_date, odometer, kind, "
                         "entered_by, entered_at) VALUES (?,?,?,?,?,?)",
                         (vehicle_id, d, int(float(odometer)), "retired", actor, utcnow_iso()))
    finally:
        conn.close()
    return f"✅ {vehicle_id} retired on {d}. Its trips and odometer history are kept."


def add_odometer(db_path: str, vehicle_id: str, odometer, actor: str, reading_date: str = "",
                 kind: str = "spot", photo: str = "") -> str:
    d = _parse_date(reading_date) if reading_date else _today()
    kind = (kind or "spot").strip().lower()
    if kind not in ("added", "year_start", "year_end", "retired", "spot"):
        return "❌ kind must be added, year_start, year_end, retired, or spot."
    try:
        val = int(float(odometer))
    except (TypeError, ValueError):
        return "❌ Odometer must be a number."
    conn = get_connection(db_path)
    try:
        if not conn.execute("SELECT 1 FROM vehicles WHERE vehicle_id = ?", (vehicle_id,)).fetchone():
            return f"❌ No vehicle {vehicle_id}."
        prev = conn.execute("SELECT odometer, reading_date FROM odometer_readings WHERE vehicle_id = ? "
                            "AND reading_date <= ? ORDER BY reading_date DESC, id DESC LIMIT 1",
                            (vehicle_id, d)).fetchone()
        conn.execute("INSERT INTO odometer_readings (vehicle_id, reading_date, odometer, kind, photo, "
                     "entered_by, entered_at) VALUES (?,?,?,?,?,?,?)",
                     (vehicle_id, d, val, kind, photo, actor, utcnow_iso()))
    finally:
        conn.close()
    warn = ""
    if prev and val < int(prev["odometer"]):
        warn = (f"\n   ⚠️ Lower than the {prev['reading_date']} reading "
                f"({int(prev['odometer']):,}). Please double-check.")
    return f"✅ Odometer {val:,} recorded for {vehicle_id} on {d} ({kind}).{warn}"


def list_vehicles(db_path: str, include_retired: bool = False, restrict_user_key: str = "") -> list:
    conn = get_connection(db_path)
    try:
        rows = conn.execute("SELECT * FROM vehicles ORDER BY active DESC, vehicle_id").fetchall()
        out = []
        for r in rows:
            if not include_retired and not r["active"]:
                continue
            if restrict_user_key and r["ownership"] == "crew" and \
                    user_key(r["owner_user_id"], r["owner_name"]) != restrict_user_key:
                continue
            last = conn.execute("SELECT odometer, reading_date FROM odometer_readings WHERE vehicle_id = ? "
                                "ORDER BY reading_date DESC, id DESC LIMIT 1", (r["vehicle_id"],)).fetchone()
            d = dict(r)
            d["label"] = _vehicle_label(r)
            d["last_odometer"] = int(last["odometer"]) if last else None
            d["last_odometer_date"] = last["reading_date"] if last else None
            out.append(d)
        return out
    finally:
        conn.close()


def set_default_vehicle(db_path: str, ukey: str, vehicle_id: str, actor: str) -> str:
    conn = get_connection(db_path)
    try:
        if vehicle_id and not conn.execute("SELECT 1 FROM vehicles WHERE vehicle_id = ? AND active = 1",
                                           (vehicle_id,)).fetchone():
            return f"❌ No active vehicle {vehicle_id}."
        raw = _get_setting(conn, SETTING_DEFAULT_VEHICLES)
        try:
            m = json.loads(raw) if raw else {}
        except Exception:
            m = {}
        if vehicle_id:
            m[ukey] = vehicle_id
        else:
            m.pop(ukey, None)
        _set_setting(conn, SETTING_DEFAULT_VEHICLES, json.dumps(m), actor)
    finally:
        conn.close()
    return f"✅ Default vehicle set to {vehicle_id}." if vehicle_id else "✅ Default vehicle cleared."


def set_day_vehicle(db_path: str, trip_date: str, ukey: str, vehicle_id: str, actor: str) -> str:
    d = _parse_date(trip_date)
    conn = get_connection(db_path)
    try:
        if not conn.execute("SELECT 1 FROM vehicles WHERE vehicle_id = ? AND active = 1",
                            (vehicle_id,)).fetchone():
            return f"❌ No active vehicle {vehicle_id}."
        conn.execute("INSERT INTO mileage_day_vehicles (trip_date, user_key, vehicle_id, set_by, set_at) "
                     "VALUES (?,?,?,?,?) ON CONFLICT(trip_date, user_key) DO UPDATE SET "
                     "vehicle_id = excluded.vehicle_id, set_by = excluded.set_by, set_at = excluded.set_at",
                     (d, ukey, vehicle_id, actor, utcnow_iso()))
        # Apply to that day's unedited auto trips right away.
        conn.execute("UPDATE mileage_trips SET vehicle_id = ?, last_edited_by = ?, last_edited_at = ? "
                     "WHERE trip_date = ? AND status = 'active' AND manually_edited = 0 AND locked = 0 "
                     "AND origin IN ('auto','backfill') AND (reimbursement_id IS NULL OR reimbursement_id NOT IN "
                     "(SELECT reimbursement_id FROM mileage_reimbursements WHERE status IN "
                     "('submitted','approved','paid'))) AND "
                     "(crew_user_id = ? OR (COALESCE(crew_user_id,'') = '' AND LOWER(COALESCE(crew,'')) = ?) "
                     " OR (? = 'owner' AND COALESCE(crew_user_id,'') = '' AND COALESCE(crew,'') = ''))",
                     (vehicle_id, actor, utcnow_iso(), d, ukey, ukey, ukey))
        _refresh_reimbursable(conn, d)
    finally:
        conn.close()
    return f"✅ Vehicle {vehicle_id} set for {d}."


def _resolve_vehicle(conn, trip_date: str, ukey: str):
    row = conn.execute("SELECT vehicle_id FROM mileage_day_vehicles WHERE trip_date = ? AND user_key = ?",
                       (trip_date, ukey)).fetchone()
    if row:
        return row["vehicle_id"]
    raw = _get_setting(conn, SETTING_DEFAULT_VEHICLES)
    try:
        m = json.loads(raw) if raw else {}
    except Exception:
        m = {}
    if m.get(ukey):
        return m[ukey]
    rows = conn.execute("SELECT vehicle_id, ownership, owner_user_id, owner_name, default_user_id, "
                        "default_user_name FROM vehicles WHERE active = 1").fetchall()
    mine = [r for r in rows if user_key(r["default_user_id"], r["default_user_name"]) == ukey
            or (r["ownership"] == "crew" and user_key(r["owner_user_id"], r["owner_name"]) == ukey)]
    if len(mine) == 1:
        return mine[0]["vehicle_id"]
    if len(rows) == 1:
        return rows[0]["vehicle_id"]
    return None


def _refresh_reimbursable(conn, trip_date: str = ""):
    sql = ("UPDATE mileage_trips SET reimbursable = CASE WHEN vehicle_id IN "
           "(SELECT vehicle_id FROM vehicles WHERE ownership = 'crew') THEN 1 ELSE 0 END")
    if trip_date:
        conn.execute(sql + " WHERE trip_date = ?", (trip_date,))
    else:
        conn.execute(sql)


# ── Trip writes (shared) ───────────────────────────────────────────────────

def _log_edit(conn, trip_id, field, old, new, actor, reason):
    conn.execute("INSERT INTO mileage_trip_edits (trip_id, field, old_value, new_value, edited_by, "
                 "edited_at, reason) VALUES (?,?,?,?,?,?,?)",
                 (trip_id, field, None if old is None else str(old),
                  None if new is None else str(new), actor, utcnow_iso(), reason))


def _insert_trip(conn, t: dict, actor: str) -> str:
    from db_write_ops import generate_next_id
    tid = generate_next_id(conn, "mileage_trips", "trip_id", "TRIP", 4)
    now = utcnow_iso()
    cols = ["trip_id", "trip_date", "start_time", "end_time", "crew_user_id", "crew", "vehicle_id",
            "from_label", "from_lat", "from_lon", "to_label", "to_lat", "to_lon", "job_id",
            "customer_name", "seq", "miles", "distance_source", "classification", "reimbursable",
            "origin_mode", "category", "purpose", "parking_usd", "tolls_usd", "receipt_photo",
            "origin", "build_key", "manually_edited", "corrected", "needs_fix", "status",
            "notes", "created_at", "created_by", "last_edited_by", "last_edited_at"]
    vals = {**{c: None for c in cols}, **t, "trip_id": tid, "created_at": now,
            "created_by": actor, "last_edited_by": actor, "last_edited_at": now}
    vals.setdefault("status", "active")
    if vals["status"] is None:
        vals["status"] = "active"
    for flag in ("reimbursable", "manually_edited", "corrected"):
        vals[flag] = int(vals[flag] or 0)
    conn.execute(f"INSERT INTO mileage_trips ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                 [vals[c] for c in cols])
    return tid


def _can_touch(trip, restrict: bool, ukey: str) -> bool:
    if not restrict:
        return True
    return user_key(trip["crew_user_id"], trip["crew"]) == ukey


# ── Daily build ────────────────────────────────────────────────────────────

def _home_and_company(db_path: str, crew_name: str):
    from db_write_ops import db_read_route_origin_mode, db_read_route_address
    mode = (db_read_route_origin_mode(db_path) or "Jobs Only").strip()
    home_addr = HOME_RESOLVER(crew_name) or ""
    home_pt = GEOCODER(home_addr) if home_addr else None
    company_addr, company_pt = "", None
    if mode.lower() == "company location":
        company_addr = db_read_route_address(db_path) or ""
        company_pt = GEOCODER(company_addr) if company_addr else None
        if not company_pt:
            mode = "Jobs Only"   # same fallback build_daily_route uses
    else:
        mode = "Jobs Only"
    return mode, (home_addr, home_pt), (company_addr, company_pt)


def _leg_miles(a_pt, b_pt, a_is_gps, b_is_gps, route_miles=None):
    m = ROUTER(a_pt, b_pt) if (a_pt and b_pt) else None
    if m is not None:
        return m, ("gps" if (a_is_gps and b_is_gps) else "address")
    if route_miles is not None:
        return round(float(route_miles), 1), "route"
    return None, None


def build_trips_for_day(db_path: str, trip_date: str, actor: str, *, crew_user_id: str = "",
                        crew: str = "", origin: str = "auto", force_rebuild_edited: bool = False) -> dict:
    """Build (or rebuild) one day's trips for every driver who clocked in
    that day — or only the given driver. Idempotent: re-running updates
    unedited auto/backfill trips, voids ones that no longer exist (e.g. a
    cancelled job), and never touches hand-edited, manual, void, or locked
    trips. Returns a summary dict."""
    d = _parse_date(trip_date)
    summary = {"date": d, "drivers": 0, "created": 0, "updated": 0, "voided": 0,
               "kept_edited": [], "gaps": [], "skipped_locked": False, "notes": []}
    conn = get_connection(db_path)
    try:
        if _is_locked(conn, d):
            summary["skipped_locked"] = True
            return summary
        entries = conn.execute(
            "SELECT t.*, j.street_address, j.city, j.state, j.zip, j.latitude, j.longitude, "
            "j.service_type, j.customer_name AS job_customer, j.job_status "
            "FROM time_entries t LEFT JOIN jobs j ON j.job_id = t.job_id "
            "WHERE t.entry_date = ? ORDER BY t.clock_in", (d,)).fetchall()
        target = entry_key(crew_user_id, crew) if crew_user_id else (
            user_key("", crew) if crew else None)
        groups: dict = {}
        for e in entries:
            k = entry_key(e["crew_user_id"], e["crew"])
            if target is not None and k != target and not (
                    k == "owner" and target == user_key("", e["crew"])):
                continue
            groups.setdefault(k, []).append(e)

        for ukey, ents in groups.items():
            summary["drivers"] += 1
            first = ents[0]
            # Personal mode (no user id): the driver is the owner; the
            # entry's crew text is the job's crew, not the driver.
            crew_name = (first["crew"] or "").strip() if first["crew_user_id"] else ""
            mode, (home_addr, home_pt), (comp_addr, comp_pt) = _home_and_company(db_path, crew_name)
            vehicle_id = _resolve_vehicle(conn, d, ukey)
            veh = conn.execute("SELECT ownership FROM vehicles WHERE vehicle_id = ?",
                               (vehicle_id,)).fetchone() if vehicle_id else None
            reimb = 1 if (veh and veh["ownership"] == "crew") else 0
            route_rows = conn.execute(
                "SELECT job_id, leg_drive_miles FROM route_stops WHERE route_date = ? AND job_id IS NOT NULL",
                (d,)).fetchall()
            route_miles = {r["job_id"]: r["leg_drive_miles"] for r in route_rows
                           if r["leg_drive_miles"] is not None}

            # Stops in clock order.
            stops = []
            for e in ents:
                job_pt = None
                if e["latitude"] is not None and e["longitude"] is not None:
                    job_pt = (float(e["latitude"]), float(e["longitude"]))
                in_gps, out_gps = _parse_gps(e["clock_in_gps"]), _parse_gps(e["clock_out_gps"])
                if not job_pt and not (in_gps or out_gps):
                    addr = ", ".join(p for p in (e["street_address"] or "", e["city"] or "",
                                                 f"{e['state'] or ''} {e['zip'] or ''}".strip()) if p)
                    job_pt = GEOCODER(addr) if addr else None
                cust = e["job_customer"] or e["customer_name"] or e["job_id"]
                purpose = f"{(e['service_type'] or 'Service').strip()} – {cust}"
                if (e["job_status"] or "").strip().lower() == "cancelled":
                    purpose = f"{cust} – cancelled on site"
                stops.append({
                    "key": e["job_id"], "label": cust, "job_id": e["job_id"], "customer": cust,
                    "in_pt": in_gps or job_pt, "in_gps": bool(in_gps),
                    "out_pt": out_gps or job_pt, "out_gps": bool(out_gps),
                    "in_time": _hhmm(e["clock_in"]), "out_time": _hhmm(e["clock_out"]),
                    "missing_out": not e["clock_out"], "entry_id": e["entry_id"], "purpose": purpose,
                })

            legs = []
            HOME = {"key": "HOME", "label": "Home", "pt": home_pt}
            COMP = {"key": "COMPANY", "label": "Company location", "pt": comp_pt}
            first_stop, last_stop = stops[0], stops[-1]
            # Day start.
            if mode == "Company Location":
                if home_pt:
                    legs.append((HOME, COMP, "commute", "Home to company location (commute)", None, None))
                legs.append((COMP, first_stop, "business", first_stop["purpose"], None, first_stop["in_time"]))
            elif home_pt:
                legs.append((HOME, first_stop, "business", first_stop["purpose"], None, first_stop["in_time"]))
            else:
                summary["notes"].append(f"{crew_name or 'Owner'}: no home address set — "
                                        "first and last legs of the day not counted.")
            # Job to job.
            for a, b in zip(stops, stops[1:]):
                legs.append((a, b, "business", b["purpose"], a["out_time"], b["in_time"]))
            # Day end.
            if mode == "Company Location":
                legs.append((last_stop, COMP, "business", "To company location", last_stop["out_time"], None))
                if home_pt:
                    legs.append((COMP, HOME, "commute", "Company location to home (commute)", None, None))
            elif home_pt:
                legs.append((last_stop, HOME, "business", "Return home", last_stop["out_time"], None))

            # Compute each leg.
            seen: dict = {}
            wanted = []
            for seq, (a, b, cls, purpose, t0, t1) in enumerate(legs, start=1):
                a_pt = a.get("out_pt", a.get("pt")) if "out_pt" in a else a.get("pt")
                b_pt = b.get("in_pt", b.get("pt")) if "in_pt" in b else b.get("pt")
                a_gps = a.get("out_gps", False)
                b_gps = b.get("in_gps", False)
                rmiles = route_miles.get(b.get("job_id")) if b.get("job_id") else None
                miles, src = _leg_miles(a_pt, b_pt, a_gps, b_gps, rmiles)
                gap = None
                if a.get("missing_out"):
                    gap = f"Missing clock-out at {a['label']}"
                elif miles is None:
                    gap = f"No location for {a['label'] if not a_pt else b['label']}"
                if gap:
                    summary["gaps"].append({"date": d, "driver": crew_name or ukey, "gap": gap,
                                            "entry_id": a.get("entry_id")})
                base = f"{d}|{ukey}|{a['key']}>{b['key']}"
                seen[base] = seen.get(base, 0) + 1
                bkey = f"{base}#{seen[base]}"
                wanted.append({
                    "trip_date": d, "start_time": t0 or None, "end_time": t1 or None,
                    "crew_user_id": first["crew_user_id"] or "", "crew": crew_name,
                    "vehicle_id": vehicle_id, "from_label": a["label"],
                    "from_lat": a_pt[0] if a_pt else None, "from_lon": a_pt[1] if a_pt else None,
                    "to_label": b["label"], "to_lat": b_pt[0] if b_pt else None,
                    "to_lon": b_pt[1] if b_pt else None, "job_id": b.get("job_id"),
                    "customer_name": b.get("customer"), "seq": seq,
                    "miles": miles if miles is not None else 0.0,
                    "distance_source": src or "none", "classification": cls,
                    "reimbursable": reimb, "origin_mode": mode,
                    "category": "Job" if (a.get("job_id") or b.get("job_id")) else "Office errand",
                    "purpose": purpose, "origin": origin, "build_key": bkey, "needs_fix": gap,
                })

            # Upsert against what exists.
            existing = conn.execute(
                "SELECT * FROM mileage_trips WHERE trip_date = ? AND origin IN ('auto','backfill') "
                "AND build_key LIKE ?", (d, f"{d}|{ukey}|%")).fetchall()
            by_key = {r["build_key"]: r for r in existing}
            conn.execute("BEGIN IMMEDIATE")
            try:
                for w in wanted:
                    old = by_key.pop(w["build_key"], None)
                    if old is None:
                        _insert_trip(conn, w, actor)
                        summary["created"] += 1
                        continue
                    if (old["manually_edited"] and not force_rebuild_edited) or old["locked"] \
                            or old["status"] == "void" or _frozen_status(conn, old):
                        if old["manually_edited"] and old["status"] == "active":
                            summary["kept_edited"].append(old["trip_id"])
                        continue
                    fields = ["start_time", "end_time", "vehicle_id", "from_label", "from_lat", "from_lon",
                              "to_label", "to_lat", "to_lon", "job_id", "customer_name", "seq", "miles",
                              "distance_source", "classification", "reimbursable", "origin_mode",
                              "purpose", "needs_fix"]
                    changed = {f: w[f] for f in fields if old[f] != w[f]}
                    if changed:
                        cols = ", ".join(f"{k} = ?" for k in changed)
                        conn.execute(f"UPDATE mileage_trips SET {cols}, last_edited_by = ?, last_edited_at = ?, "
                                     f"version = version + 1 WHERE trip_id = ?",
                                     (*changed.values(), actor, utcnow_iso(), old["trip_id"]))
                        summary["updated"] += 1
                for gone in by_key.values():
                    if gone["status"] != "active" or gone["locked"] or _frozen_status(conn, gone):
                        continue
                    if gone["manually_edited"] and not force_rebuild_edited:
                        summary["kept_edited"].append(gone["trip_id"])
                        continue
                    conn.execute("UPDATE mileage_trips SET status = 'void', void_reason = ?, last_edited_by = ?, "
                                 "last_edited_at = ?, version = version + 1 WHERE trip_id = ?",
                                 ("Rebuilt: this leg is no longer part of the day (job moved, cancelled, or reordered)",
                                  actor, utcnow_iso(), gone["trip_id"]))
                    _log_edit(conn, gone["trip_id"], "status", "active", "void", actor, "Automatic rebuild")
                    summary["voided"] += 1
                conn.execute("COMMIT")
            except Exception:
                conn.execute("ROLLBACK")
                raise

        # Drivers with built trips but no clock-ins that day any more (job
        # deleted or moved, entry removed): those drives didn't happen — void
        # them (FR-21). Hand-edited, locked, and submitted/paid trips are kept.
        orphan_keys = set()
        for o in conn.execute("SELECT * FROM mileage_trips WHERE trip_date = ? AND status = 'active' "
                              "AND origin IN ('auto','backfill')", (d,)).fetchall():
            parts = (o["build_key"] or "").split("|")
            k = parts[1] if len(parts) >= 3 else ""
            if not k or k in groups or (target is not None and k != target):
                continue
            orphan_keys.add(k)
            if o["locked"] or _frozen_status(conn, o):
                continue
            if o["manually_edited"] and not force_rebuild_edited:
                summary["kept_edited"].append(o["trip_id"])
                continue
            conn.execute("UPDATE mileage_trips SET status = 'void', void_reason = ?, last_edited_by = ?, "
                         "last_edited_at = ?, version = version + 1 WHERE trip_id = ?",
                         ("Rebuilt: no clock-ins for this driver that day any more (job moved, cancelled, "
                          "or entry removed)", actor, utcnow_iso(), o["trip_id"]))
            _log_edit(conn, o["trip_id"], "status", "active", "void", actor, "Automatic rebuild")
            summary["voided"] += 1
        _store_signatures(conn, d, set(groups) | orphan_keys | ({target} if (target and crew_user_id) else set()))
        return summary
    finally:
        conn.close()


def summarize_build(s: dict) -> str:
    if s.get("skipped_locked"):
        return f"🔒 {s['date']} is in a locked period — nothing rebuilt."
    if not s["drivers"] and not s.get("voided"):
        return f"ℹ️ No clock-ins on {s['date']} — no job trips to build."
    lines = [f"🚗 Mileage built for {s['date']}: {s['created']} new, {s['updated']} updated, "
             f"{s['voided']} removed ({s['drivers']} driver(s))."]
    if s["kept_edited"]:
        lines.append(f"   ✋ Kept hand-edited trip(s): {', '.join(sorted(set(s['kept_edited'])))}")
    for g in s["gaps"]:
        lines.append(f"   ⚠️ Needs fix — {g['driver']}: {g['gap']}")
    for n in s["notes"]:
        lines.append(f"   ℹ️ {n}")
    return "\n".join(lines)


_SWEEP_MARKER_KEY = "internal_mileage_sweep_last_run"


def maybe_sweep_recent_days(db_path: str, actor: str = "automatic", days: int = 14) -> int:
    """Safety net for the clock-out hook (phone offline, server restarted):
    at most once per calendar day, build any of the last `days` days that
    have clock-ins but no trips yet. Same once-a-day, triggered-by-a-read
    pattern as _maybe_sweep_recurring_jobs (there's no clock scheduler in
    personal mode). Returns days built. Never raises."""
    try:
        if not tracking_enabled(db_path):
            return 0
        today = _dt.date.today()
        conn = get_connection(db_path)
        try:
            if (_get_setting(conn, _SWEEP_MARKER_KEY) or "") == today.isoformat():
                return 0
            _set_setting(conn, _SWEEP_MARKER_KEY, today.isoformat(), actor)
            s = (today - _dt.timedelta(days=days)).isoformat()
            y = (today - _dt.timedelta(days=1)).isoformat()
            entry_days = [r[0] for r in conn.execute(
                "SELECT DISTINCT entry_date FROM time_entries WHERE entry_date BETWEEN ? AND ?", (s, y)).fetchall()]
            built = {r[0] for r in conn.execute(
                "SELECT DISTINCT trip_date FROM mileage_trips WHERE trip_date BETWEEN ? AND ? "
                "AND origin IN ('auto','backfill')", (s, y)).fetchall()}
            todo = [d for d in entry_days if d not in built and not _is_locked(conn, d)]
        finally:
            conn.close()
        for d in todo:
            build_trips_for_day(db_path, d, actor)
        return len(todo)
    except Exception:
        return 0


def backfill(db_path: str, start: str, end: str, actor: str, preview: bool = True) -> str:
    s, e = _parse_date(start), _parse_date(end)
    if s > e:
        s, e = e, s
    conn = get_connection(db_path)
    try:
        days = [r[0] for r in conn.execute(
            "SELECT DISTINCT entry_date FROM time_entries WHERE entry_date BETWEEN ? AND ? "
            "ORDER BY entry_date", (s, e)).fetchall()]
        already = {r[0] for r in conn.execute(
            "SELECT DISTINCT trip_date FROM mileage_trips WHERE trip_date BETWEEN ? AND ? "
            "AND origin IN ('auto','backfill')", (s, e)).fetchall()}
        no_gps = conn.execute(
            "SELECT COUNT(*) FROM time_entries WHERE entry_date BETWEEN ? AND ? AND "
            "(COALESCE(clock_in_gps,'') = '' OR COALESCE(clock_out_gps,'') = '')", (s, e)).fetchone()[0]
        locked = [d for d in days if _is_locked(conn, d)]
    finally:
        conn.close()
    todo = [d for d in days if d not in already and d not in locked]
    if preview:
        return (f"🔎 Backfill preview {s} → {e}\n"
                f"   Days with clock-ins: {len(days)}\n"
                f"   Already have trips: {len(already & set(days))} (left as they are)\n"
                f"   Locked days skipped: {len(locked)}\n"
                f"   Days to build: {len(todo)}\n"
                f"   Clock entries missing GPS (job address will be used): {no_gps}\n"
                f"   Nothing has been written. Confirm to build these days; they'll be marked as backfilled.")
    totals = {"created": 0, "gaps": 0}
    for d in todo:
        r = build_trips_for_day(db_path, d, actor, origin="backfill")
        totals["created"] += r["created"]
        totals["gaps"] += len(r["gaps"])
    return (f"✅ Backfill {s} → {e}: built {len(todo)} day(s), {totals['created']} trips "
            f"(marked as backfilled). {totals['gaps']} gap(s) need a fix — see the Mileage tab.")


# ── Manual trips, edits, void, split, gap fix ──────────────────────────────

def add_trip(db_path: str, actor: str, *, trip_date: str, to_label: str, purpose: str,
             miles=None, from_label: str = "", category: str = "Other", crew_user_id: str = "",
             crew: str = "", vehicle_id: str = "", start_time: str = "", end_time: str = "",
             round_trip: bool = False, from_gps: str = "", to_gps: str = "",
             odometer_start=None, odometer_end=None, parking_usd=None, tolls_usd=None,
             job_id: str = "", customer_name: str = "", notes: str = "", receipt_photo: str = "",
             classification: str = "business") -> str:
    """Manual / non-job trip (form, text, voice, Start/End buttons)."""
    d = _parse_date(trip_date)
    if not (to_label or "").strip():
        return "❌ Destination is required (IRS: where you drove)."
    err = validate_purpose(purpose)
    if err:
        return f"❌ {err}"
    cls = (classification or "business").strip().lower()
    if cls not in ("business", "commute", "personal"):
        return "❌ Classification must be business, commute, or personal."
    a_pt, b_pt = _parse_gps(from_gps), _parse_gps(to_gps)
    src = "manual"
    if odometer_start not in (None, "") and odometer_end not in (None, ""):
        miles = float(odometer_end) - float(odometer_start)
        src = "odometer"
        if miles <= 0:
            return "❌ Odometer end must be higher than odometer start."
    if miles in (None, ""):
        a_pt = a_pt or (GEOCODER(from_label) if from_label else None)
        b_pt = b_pt or GEOCODER(to_label)
        m = ROUTER(a_pt, b_pt) if (a_pt and b_pt) else None
        if m is None:
            return ("❌ I couldn't work out the miles from those places. "
                    "Please give the miles or the odometer start and end.")
        miles = m
        src = "gps" if (_parse_gps(from_gps) and _parse_gps(to_gps)) else "address"
    try:
        miles = round(float(miles), 1)
    except (TypeError, ValueError):
        return "❌ Miles must be a number."
    if miles <= 0 or miles > 1500:
        return "❌ Miles must be between 0 and 1,500 for a single trip."
    ukey = user_key(crew_user_id, crew)
    conn = get_connection(db_path)
    try:
        if _is_locked(conn, d):
            return f"🔒 {d} is in a locked period — unlock it first."
        vid = vehicle_id or _resolve_vehicle(conn, d, ukey)
        veh = conn.execute("SELECT ownership FROM vehicles WHERE vehicle_id = ?", (vid,)).fetchone() if vid else None
        if vid and veh is None:
            return f"❌ No vehicle {vid}."
        if veh and veh["ownership"] == "crew":
            pf = _period_frozen_for(conn, ukey, d)
            if pf:
                return (f"🔒 Your reimbursement for that pay period ({pf[0]}) is {pf[1]} — ask the owner to "
                        "return it before adding trips to it.")
        base = {"trip_date": d, "start_time": start_time or None, "end_time": end_time or None,
                "crew_user_id": crew_user_id or "", "crew": crew or "", "vehicle_id": vid,
                "from_label": from_label or "", "to_label": to_label.strip(),
                "from_lat": a_pt[0] if a_pt else None, "from_lon": a_pt[1] if a_pt else None,
                "to_lat": b_pt[0] if b_pt else None, "to_lon": b_pt[1] if b_pt else None,
                "job_id": job_id or None, "customer_name": customer_name or None,
                "distance_source": src, "classification": cls,
                "reimbursable": 1 if (veh and veh["ownership"] == "crew") else 0,
                "category": category if category in CATEGORIES else "Other",
                "purpose": purpose.strip(), "origin": "manual", "manually_edited": 1,
                "parking_usd": float(parking_usd) if parking_usd not in (None, "") else None,
                "tolls_usd": float(tolls_usd) if tolls_usd not in (None, "") else None,
                "notes": notes or None, "receipt_photo": receipt_photo or None}
        conn.execute("BEGIN IMMEDIATE")
        ids = []
        if round_trip:
            half = round(miles / 2, 1) if src == "odometer" else miles
            ids.append(_insert_trip(conn, {**base, "miles": half}, actor))
            back = {**base, "from_label": base["to_label"], "to_label": base["from_label"] or "Start",
                    "from_lat": base["to_lat"], "from_lon": base["to_lon"],
                    "to_lat": base["from_lat"], "to_lon": base["from_lon"],
                    "miles": round(miles - half, 1) if src == "odometer" else miles,
                    "parking_usd": None, "tolls_usd": None, "receipt_photo": None,
                    "start_time": None, "end_time": None}
            ids.append(_insert_trip(conn, back, actor))
        else:
            ids.append(_insert_trip(conn, {**base, "miles": miles}, actor))
        conn.execute("COMMIT")
    except Exception as exc:
        try:
            conn.execute("ROLLBACK")
        except Exception:
            pass
        return f"❌ Could not save the trip: {exc}"
    finally:
        conn.close()
    total = miles if not round_trip or src == "odometer" else miles * 2
    late = (_dt.date.today() - _dt.date.fromisoformat(d)).days > LATE_ENTRY_DAYS
    warn = "" if vid else "\n   ⚠️ No vehicle chosen — set one so the trip counts for the right vehicle."
    if late:
        warn += f"\n   ℹ️ Added more than {LATE_ENTRY_DAYS} days after the trip — flagged as a late entry."
    return f"✅ Saved {', '.join(ids)} — {d}, {total:g} mi to {to_label.strip()} ({purpose.strip()}).{warn}"


def update_trip(db_path: str, trip_id: str, changes: dict, reason: str, actor: str,
                restrict: bool = False, ukey: str = "") -> str:
    if not (reason or "").strip():
        return "❌ A short reason is required for every change (it's kept in the audit trail)."
    sets = {k: v for k, v in (changes or {}).items() if k in _EDITABLE_FIELDS}
    if not sets:
        return "❌ Nothing to change."
    if "purpose" in sets:
        err = validate_purpose(sets["purpose"])
        if err:
            return f"❌ {err}"
    if "to_label" in sets and not str(sets["to_label"]).strip():
        return "❌ Destination can't be blank."
    if "miles" in sets:
        try:
            sets["miles"] = round(float(sets["miles"]), 1)
        except (TypeError, ValueError):
            return "❌ Miles must be a number."
        if sets["miles"] < 0 or sets["miles"] > 1500:
            return "❌ Miles must be between 0 and 1,500."
    if "classification" in sets and sets["classification"] not in ("business", "commute", "personal"):
        return "❌ Classification must be business, commute, or personal."
    if "trip_date" in sets:
        sets["trip_date"] = _parse_date(sets["trip_date"])
    conn = get_connection(db_path)
    try:
        t = conn.execute("SELECT * FROM mileage_trips WHERE trip_id = ?", (trip_id,)).fetchone()
        if t is None:
            return f"❌ No trip {trip_id}."
        if not _can_touch(t, restrict, ukey):
            return "❌ You can only change your own trips."
        if t["status"] == "void":
            return f"❌ {trip_id} is void."
        if t["locked"] or _is_locked(conn, t["trip_date"]) or \
                ("trip_date" in sets and _is_locked(conn, sets["trip_date"])):
            return f"🔒 {trip_id} is in a locked or paid period and can't be changed."
        fz = _frozen_status(conn, t)
        if fz:
            return (f"🔒 {trip_id} is in a crew reimbursement that's {fz} — it can't be changed unless the "
                    "owner returns it.")
        if "vehicle_id" in sets and sets["vehicle_id"] and not conn.execute(
                "SELECT 1 FROM vehicles WHERE vehicle_id = ?", (sets["vehicle_id"],)).fetchone():
            return f"❌ No vehicle {sets['vehicle_id']}."
        if "vehicle_id" in sets or "trip_date" in sets:
            vid2 = sets.get("vehicle_id", t["vehicle_id"])
            own = conn.execute("SELECT ownership FROM vehicles WHERE vehicle_id = ?", (vid2,)).fetchone() \
                if vid2 else None
            if own and own["ownership"] == "crew":
                pf = _period_frozen_for(conn, user_key(t["crew_user_id"], t["crew"]),
                                        sets.get("trip_date", t["trip_date"]))
                if pf:
                    return (f"🔒 That pay period's reimbursement ({pf[0]}) is {pf[1]} — ask the owner to "
                            "return it first.")
        changed = {k: v for k, v in sets.items() if t[k] != v}
        if not changed:
            return f"ℹ️ {trip_id} already has those values."
        conn.execute("BEGIN IMMEDIATE")
        for k, v in changed.items():
            _log_edit(conn, trip_id, k, t[k], v, actor, reason.strip())
        extra = ", distance_source = 'manual'" if "miles" in changed else ""
        cols = ", ".join(f"{k} = ?" for k in changed)
        conn.execute(f"UPDATE mileage_trips SET {cols}{extra}, manually_edited = 1, needs_fix = NULL, "
                     f"last_edited_by = ?, last_edited_at = ?, version = version + 1 WHERE trip_id = ?",
                     (*changed.values(), actor, utcnow_iso(), trip_id))
        _refresh_reimbursable(conn, changed.get("trip_date", t["trip_date"]))
        conn.execute("COMMIT")
    finally:
        conn.close()
    return f"✅ {trip_id} updated ({', '.join(changed)}). Change recorded with your reason."


def void_trip(db_path: str, trip_id: str, reason: str, actor: str, restrict: bool = False,
              ukey: str = "") -> str:
    if not (reason or "").strip():
        return "❌ A reason is required to remove a trip (it stays in the record as void)."
    conn = get_connection(db_path)
    try:
        t = conn.execute("SELECT * FROM mileage_trips WHERE trip_id = ?", (trip_id,)).fetchone()
        if t is None:
            return f"❌ No trip {trip_id}."
        if not _can_touch(t, restrict, ukey):
            return "❌ You can only change your own trips."
        if t["status"] == "void":
            return f"ℹ️ {trip_id} is already void."
        if t["locked"] or _is_locked(conn, t["trip_date"]):
            return f"🔒 {trip_id} is in a locked or paid period."
        fz = _frozen_status(conn, t)
        if fz:
            return (f"🔒 {trip_id} is in a crew reimbursement that's {fz} — it can't be removed unless the "
                    "owner returns it.")
        conn.execute("BEGIN IMMEDIATE")
        _log_edit(conn, trip_id, "status", "active", "void", actor, reason.strip())
        conn.execute("UPDATE mileage_trips SET status = 'void', void_reason = ?, manually_edited = 1, "
                     "last_edited_by = ?, last_edited_at = ?, version = version + 1 WHERE trip_id = ?",
                     (reason.strip(), actor, utcnow_iso(), trip_id))
        conn.execute("COMMIT")
    finally:
        conn.close()
    return f"✅ {trip_id} removed from totals (kept in the record as void: {reason.strip()})."


def split_trip(db_path: str, trip_id: str, via_label: str, purpose: str, reason: str, actor: str,
               via_gps: str = "", miles_to_via=None, miles_from_via=None, category: str = "Supplies",
               restrict: bool = False, ukey: str = "") -> str:
    """Insert a stop (supply house, bank) into an existing leg: A→B becomes
    A→via→B. The original leg is voided; two manual trips replace it."""
    err = validate_purpose(purpose)
    if err:
        return f"❌ {err}"
    if not (reason or "").strip():
        return "❌ A short reason is required."
    conn = get_connection(db_path)
    try:
        t = conn.execute("SELECT * FROM mileage_trips WHERE trip_id = ?", (trip_id,)).fetchone()
    finally:
        conn.close()
    if t is None:
        return f"❌ No trip {trip_id}."
    if not _can_touch(t, restrict, ukey):
        return "❌ You can only change your own trips."
    a_pt = (t["from_lat"], t["from_lon"]) if t["from_lat"] is not None else None
    b_pt = (t["to_lat"], t["to_lon"]) if t["to_lat"] is not None else None
    v_pt = _parse_gps(via_gps) or GEOCODER(via_label)
    m1 = miles_to_via if miles_to_via not in (None, "") else (ROUTER(a_pt, v_pt) if a_pt and v_pt else None)
    m2 = miles_from_via if miles_from_via not in (None, "") else (ROUTER(v_pt, b_pt) if b_pt and v_pt else None)
    if m1 is None or m2 is None:
        return "❌ Couldn't work out the miles for the extra stop — please give both miles."
    res = void_trip(db_path, trip_id, f"Split: added stop at {via_label} ({reason.strip()})", actor, restrict, ukey)
    if not res.startswith("✅"):
        return res
    common = dict(trip_date=t["trip_date"], crew_user_id=t["crew_user_id"] or "", crew=t["crew"] or "",
                  vehicle_id=t["vehicle_id"] or "", classification=t["classification"])
    r1 = add_trip(db_path, actor, from_label=t["from_label"] or "", to_label=via_label, purpose=purpose,
                  miles=m1, category=category, start_time=t["start_time"] or "", **common)
    r2 = add_trip(db_path, actor, from_label=via_label, to_label=t["to_label"], purpose=t["purpose"],
                  miles=m2, category=t["category"] or "Job", end_time=t["end_time"] or "",
                  job_id=t["job_id"] or "", customer_name=t["customer_name"] or "", **common)
    return f"✅ Split {trip_id} at {via_label}.\n   {r1}\n   {r2}"


def fix_missing_clock_out(db_path: str, entry_id: str, clock_out_time: str, reason: str, actor: str,
                          restrict: bool = False, ukey: str = "") -> str:
    """Dead phone / forgotten clock-out: fill the clock-out (location
    defaults to the job address), record who/why on the time entry, then
    rebuild that day's trips and mark the affected trips corrected."""
    if not (reason or "").strip():
        return "❌ A reason is required (for example 'phone died')."
    conn = get_connection(db_path)
    try:
        e = conn.execute("SELECT * FROM time_entries WHERE entry_id = ?", (entry_id,)).fetchone()
        if e is None:
            return f"❌ No time entry {entry_id}."
        if restrict and entry_key(e["crew_user_id"], e["crew"]) != ukey:
            return "❌ You can only fix your own time entries."
        if e["clock_out"]:
            return f"ℹ️ {entry_id} already has a clock-out ({e['clock_out']})."
        t = (clock_out_time or "").strip()
        if not re.fullmatch(r"\d{1,2}:\d{2}", t):
            return "❌ Clock-out time must look like 15:30."
        hh, mm = t.split(":")
        out_ts = f"{e['entry_date']} {int(hh):02d}:{mm}:00"
        if e["clock_in"] and out_ts <= str(e["clock_in"]):
            return "❌ Clock-out must be after the clock-in time."
        if _is_locked(conn, e["entry_date"]):
            return f"🔒 {e['entry_date']} is locked."
        note = f"[Clock-out added later by {actor}: {reason.strip()}]"
        conn.execute("UPDATE time_entries SET clock_out = ?, notes = TRIM(COALESCE(notes,'') || ' ' || ?), "
                     "last_edited_by = ?, last_edited_at = ?, version = version + 1 WHERE entry_id = ?",
                     (out_ts, note, actor, utcnow_iso(), entry_id))
        day, cu, cr = e["entry_date"], e["crew_user_id"] or "", e["crew"] or ""
    finally:
        conn.close()
    s = build_trips_for_day(db_path, day, actor, crew_user_id=cu, crew=cr)
    conn = get_connection(db_path)
    try:
        conn.execute("UPDATE mileage_trips SET corrected = 1 WHERE trip_date = ? AND status = 'active' "
                     "AND build_key LIKE ? AND (from_label = (SELECT customer_name FROM time_entries "
                     "WHERE entry_id = ?) OR job_id = (SELECT job_id FROM time_entries WHERE entry_id = ?))",
                     (day, f"{day}|{entry_key(cu, cr)}|%", entry_id, entry_id))
    finally:
        conn.close()
    return f"✅ Clock-out {t} added to {entry_id} ({reason.strip()}).\n{summarize_build(s)}"


# ── Queries & reports ──────────────────────────────────────────────────────

def list_trips(db_path: str, start: str, end: str, restrict: bool = False, ukey: str = "",
               driver_key: str = "", include_void: bool = False) -> list:
    s, e = _parse_date(start), _parse_date(end)
    conn = get_connection(db_path)
    try:
        rows = conn.execute("SELECT t.*, v.nickname, v.year, v.make, v.model, v.vin, v.ownership "
                            "FROM mileage_trips t LEFT JOIN vehicles v ON v.vehicle_id = t.vehicle_id "
                            "WHERE t.trip_date BETWEEN ? AND ? ORDER BY t.trip_date, "
                            "COALESCE(t.start_time, t.end_time, ''), t.seq, t.trip_id", (s, e)).fetchall()
        rates = get_rates(db_path)
        out = []
        for r in rows:
            k = user_key(r["crew_user_id"], r["crew"])
            if restrict and k != ukey:
                continue
            if driver_key and k != driver_key:
                continue
            if not include_void and r["status"] == "void":
                continue
            d = dict(r)
            d["driver_key"] = k
            d["vehicle_label"] = (_vehicle_label(r) if r["vehicle_id"] and r["vin"] else "")
            d["rate_cents"] = rate_cents_for(rates, r["trip_date"])
            d["late_entry"] = _late(r)
            d["locked_period"] = _is_locked(conn, r["trip_date"])
            out.append(d)
        return out
    finally:
        conn.close()


def _late(r) -> bool:
    try:
        created = _dt.datetime.fromisoformat(str(r["created_at"]).replace("Z", "+00:00")).date()
        return (created - _dt.date.fromisoformat(r["trip_date"])).days > LATE_ENTRY_DAYS
    except Exception:
        return False


def trip_history(db_path: str, trip_id: str) -> list:
    conn = get_connection(db_path)
    try:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM mileage_trip_edits WHERE trip_id = ? ORDER BY id", (trip_id,)).fetchall()]
    finally:
        conn.close()


def day_gaps(db_path: str, start: str, end: str, restrict: bool = False, ukey: str = "") -> list:
    """Open clock-ins (no clock-out) and trips flagged needs_fix."""
    s, e = _parse_date(start), _parse_date(end)
    out = []
    conn = get_connection(db_path)
    try:
        for r in conn.execute("SELECT entry_id, entry_date, crew_user_id, crew, customer_name, job_id, clock_in "
                              "FROM time_entries WHERE entry_date BETWEEN ? AND ? AND clock_in IS NOT NULL "
                              "AND clock_out IS NULL AND entry_date < ?", (s, e, _today())).fetchall():
            k = entry_key(r["crew_user_id"], r["crew"])
            if restrict and k != ukey:
                continue
            out.append({"type": "missing_clock_out", "date": r["entry_date"], "entry_id": r["entry_id"],
                        "driver": r["crew"] or k, "where": r["customer_name"] or r["job_id"],
                        "clock_in": _hhmm(r["clock_in"])})
        for r in conn.execute("SELECT trip_id, trip_date, crew_user_id, crew, needs_fix FROM mileage_trips "
                              "WHERE trip_date BETWEEN ? AND ? AND status = 'active' AND needs_fix IS NOT NULL",
                              (s, e)).fetchall():
            k = user_key(r["crew_user_id"], r["crew"])
            if restrict and k != ukey:
                continue
            out.append({"type": "trip", "date": r["trip_date"], "trip_id": r["trip_id"],
                        "driver": r["crew"] or k, "gap": r["needs_fix"]})
    finally:
        conn.close()
    return out


def mileage_report(db_path: str, start: str, end: str, restrict: bool = False, ukey: str = "",
                   driver_key: str = "") -> dict:
    """Miles report, deduction summary (by IRS rate period), crew
    reimbursement summary, Form 4562 vehicle sheet, and compliance checks."""
    s, e = _parse_date(start), _parse_date(end)
    trips = list_trips(db_path, s, e, restrict, ukey, driver_key)
    rates = get_rates(db_path)
    rep = {"start": s, "end": e, "by_driver": {}, "by_vehicle": {}, "by_month": {},
           "deduction": {}, "parking_tolls": 0.0, "reimbursement": {}, "vehicles_4562": [],
           "compliance": [], "trip_count": len(trips)}

    def bucket(d, key):
        return d.setdefault(key, {"business": 0.0, "commute": 0.0, "personal": 0.0})

    for t in trips:
        cls = t["classification"] if t["classification"] in ("business", "commute", "personal") else "business"
        m = float(t["miles"] or 0)
        bucket(rep["by_driver"], t["crew"] or t["driver_key"])[cls] += m
        bucket(rep["by_vehicle"], t["vehicle_label"] or "(no vehicle)")[cls] += m
        bucket(rep["by_month"], t["trip_date"][:7])[cls] += m
        if cls == "business":
            cents = rate_cents_for(rates, t["trip_date"]) or 0.0
            if t["reimbursable"]:
                r = rep["reimbursement"].setdefault(t["crew"] or t["driver_key"], {})
                p = r.setdefault(f"{cents:g}¢", {"miles": 0.0, "amount": 0.0})
                p["miles"] += m
                p["amount"] += m * cents / 100
            else:
                p = rep["deduction"].setdefault(f"{cents:g}¢", {"miles": 0.0, "amount": 0.0,
                                                               "from": None, "to": None})
                p["miles"] += m
                p["amount"] += m * cents / 100
                p["from"] = min(p["from"] or t["trip_date"], t["trip_date"])
                p["to"] = max(p["to"] or t["trip_date"], t["trip_date"])
            rep["parking_tolls"] += float(t["parking_usd"] or 0) + float(t["tolls_usd"] or 0)

    # Form 4562 Part V, per business vehicle (calendar years within range).
    if not restrict:
        conn = get_connection(db_path)
        try:
            for v in conn.execute("SELECT * FROM vehicles WHERE ownership = 'business' ORDER BY vehicle_id").fetchall():
                for year in range(int(s[:4]), int(e[:4]) + 1):
                    ys, ye = f"{year}-01-01", f"{year}-12-31"
                    sums = {r["classification"]: r["m"] for r in conn.execute(
                        "SELECT classification, SUM(miles) AS m FROM mileage_trips WHERE vehicle_id = ? AND "
                        "status = 'active' AND trip_date BETWEEN ? AND ? GROUP BY classification",
                        (v["vehicle_id"], ys, ye)).fetchall()}
                    if not sums and not v["active"]:
                        continue
                    lo = conn.execute("SELECT odometer FROM odometer_readings WHERE vehicle_id = ? AND "
                                      "reading_date BETWEEN ? AND ? ORDER BY reading_date, id LIMIT 1",
                                      (v["vehicle_id"], ys, ye)).fetchone()
                    hi = conn.execute("SELECT odometer FROM odometer_readings WHERE vehicle_id = ? AND "
                                      "reading_date BETWEEN ? AND ? ORDER BY reading_date DESC, id DESC LIMIT 1",
                                      (v["vehicle_id"], ys, ye)).fetchone()
                    total = (int(hi["odometer"]) - int(lo["odometer"])) if (lo and hi) else None
                    bus, com = float(sums.get("business") or 0), float(sums.get("commute") or 0)
                    personal = None
                    if total is not None and total > 0:
                        personal = max(0.0, total - bus - com)
                    rep["vehicles_4562"].append({
                        "year": year, "vehicle_id": v["vehicle_id"], "label": _vehicle_label(v),
                        "vin": v["vin"], "placed_in_service": v["placed_in_service"],
                        "business": round(bus, 1), "commute": round(com, 1),
                        "personal": None if personal is None else round(personal, 1),
                        "total": total})
                    if total is None or (lo and hi and lo["odometer"] == hi["odometer"]):
                        rep["compliance"].append(f"{v['vehicle_id']} {year}: needs start-of-year and "
                                                 "end-of-year odometer readings for total miles.")
            # 5+ business vehicles same day.
            for r in conn.execute("SELECT t.trip_date, COUNT(DISTINCT t.vehicle_id) AS n FROM mileage_trips t "
                                  "JOIN vehicles v ON v.vehicle_id = t.vehicle_id WHERE v.ownership = 'business' "
                                  "AND t.status = 'active' AND t.trip_date BETWEEN ? AND ? GROUP BY t.trip_date "
                                  "HAVING n >= 5", (s, e)).fetchall():
                rep["compliance"].append(f"{r['trip_date']}: {r['n']} business vehicles used the same day — "
                                         "the IRS standard mileage rate isn't allowed with 5+ vehicles in use "
                                         "at once. Ask your accountant.")
        finally:
            conn.close()

    no_vehicle = [t["trip_id"] for t in trips if not t["vehicle_id"]]
    vague = [t["trip_id"] for t in trips if validate_purpose(t["purpose"])]
    late = [t["trip_id"] for t in trips if t["late_entry"]]
    gaps = [t["trip_id"] for t in trips if t["needs_fix"]]
    if no_vehicle:
        rep["compliance"].append(f"{len(no_vehicle)} trip(s) have no vehicle: {', '.join(no_vehicle[:8])}"
                                 + (" …" if len(no_vehicle) > 8 else ""))
    if vague:
        rep["compliance"].append(f"{len(vague)} trip(s) need a specific purpose: {', '.join(vague[:8])}")
    if gaps:
        rep["compliance"].append(f"{len(gaps)} trip(s) need a fix (missing clock-out or location).")
    if late:
        rep["compliance"].append(f"{len(late)} trip(s) were entered more than {LATE_ENTRY_DAYS} days after "
                                 "the trip — keep any supporting records (receipts, calendar).")
    if any(rate_cents_for(rates, t["trip_date"]) is None for t in trips):
        rep["compliance"].append("Some trips are before the first rate in Settings → Mileage rates.")
    try:   # owed / submitted / approved-unpaid / paid per crew member (FR-23)
        rs = reimbursement_status_summary(db_path, s, e, restrict, ukey)
        rep["reimbursement_status"] = {k: v for k, v in rs.items() if k in rep["reimbursement"]}
    except Exception:
        rep["reimbursement_status"] = {}
    return rep


def format_report(rep: dict) -> str:
    L = [f"🚗 Mileage report {rep['start']} → {rep['end']} ({rep['trip_count']} trips)", "─" * 52]

    def tot(d):
        return sum(v["business"] for v in d.values()), sum(v["commute"] for v in d.values())
    b, c = tot(rep["by_driver"]) if rep["by_driver"] else (0, 0)
    L.append(f"Business miles: {b:,.1f}   Commute miles (not deductible): {c:,.1f}")
    if rep["by_driver"]:
        L.append("\nBy driver:")
        for k, v in sorted(rep["by_driver"].items()):
            L.append(f"  {k}: {v['business']:,.1f} business, {v['commute']:,.1f} commute")
    if rep["by_vehicle"]:
        L.append("\nBy vehicle:")
        for k, v in sorted(rep["by_vehicle"].items()):
            L.append(f"  {k}: {v['business']:,.1f} business, {v['commute']:,.1f} commute")
    if rep["by_month"]:
        L.append("\nBy month:")
        for k, v in sorted(rep["by_month"].items()):
            L.append(f"  {k}: {v['business']:,.1f} business")
    L.append("\nDeduction summary (business-owned vehicles, standard mileage rate):")
    total = 0.0
    for k, p in sorted(rep["deduction"].items(), key=lambda kv: kv[1]["from"] or ""):
        L.append(f"  {p['from']} – {p['to']}: {p['miles']:,.1f} mi × {k} = ${p['amount']:,.2f}")
        total += p["amount"]
    L.append(f"  Parking and tolls: ${rep['parking_tolls']:,.2f}")
    L.append(f"  Total: ${total + rep['parking_tolls']:,.2f}")
    if rep["reimbursement"]:
        L.append("\nCrew reimbursement (crew-owned vehicles):")
        for who, per in sorted(rep["reimbursement"].items()):
            amt = sum(p["amount"] for p in per.values())
            mi = sum(p["miles"] for p in per.values())
            st = (rep.get("reimbursement_status") or {}).get(who)
            tail = ""
            if st:
                tail = (f" — not submitted ${st['not_submitted']:,.2f}, submitted ${st['submitted']:,.2f}, "
                        f"approved unpaid ${st['approved']:,.2f}, paid ${st['paid']:,.2f}"
                        " (incl. parking/tolls)")
            L.append(f"  {who}: {mi:,.1f} mi, ${amt:,.2f} owed{tail}")
    if rep["vehicles_4562"]:
        L.append("\nForm 4562 Part V — per business vehicle:")
        for v in rep["vehicles_4562"]:
            tot_s = f"{v['total']:,}" if v["total"] is not None else "unknown (odometer needed)"
            per_s = f"{v['personal']:,.1f}" if v["personal"] is not None else "unknown"
            L.append(f"  {v['year']} {v['label']} VIN {v['vin']} (in service {v['placed_in_service']}): "
                     f"business {v['business']:,.1f}, commute {v['commute']:,.1f}, other personal {per_s}, "
                     f"total {tot_s}")
    if rep["compliance"]:
        L.append("\n⚠️ IRS record checks:")
        for m in rep["compliance"]:
            L.append(f"  • {m}")
    else:
        L.append("\n✅ IRS record checks passed: every trip has a date, destination, purpose, miles, and vehicle.")
    L.append("\nNot tax advice — your accountant confirms how to claim these.")
    return "\n".join(L)


def export_rows(db_path: str, start: str, end: str, restrict: bool = False, ukey: str = "",
                driver_key: str = "") -> list:
    """One dict per trip, for CSV/Excel/PDF export (void trips included and marked)."""
    rows = []
    for t in list_trips(db_path, start, end, restrict, ukey, driver_key, include_void=True):
        cents = t["rate_cents"] or 0.0
        amt = (t["miles"] or 0) * cents / 100 if (t["classification"] == "business" and t["status"] == "active") else 0.0
        rows.append({
            "Trip ID": t["trip_id"], "Date": t["trip_date"], "Start": t["start_time"] or "",
            "End": t["end_time"] or "", "Driver": t["crew"] or t["driver_key"],
            "Vehicle": t["vehicle_label"], "VIN": t["vin"] or "",
            "Ownership": t["ownership"] or "", "From": t["from_label"] or "", "To": t["to_label"],
            "Miles": t["miles"], "Purpose": t["purpose"], "Category": t["category"] or "",
            "Classification": t["classification"], "Rate (¢/mi)": cents,
            "Amount ($)": round(amt, 2), "Reimbursable": "yes" if t["reimbursable"] else "",
            "Parking ($)": t["parking_usd"] or "", "Tolls ($)": t["tolls_usd"] or "",
            "Distance source": t["distance_source"], "Entry": t["origin"],
            "Edited": "yes" if t["manually_edited"] else "", "Corrected": "yes" if t["corrected"] else "",
            "Status": t["status"], "Void reason": t["void_reason"] or "",
            "Recorded at": t["created_at"], "Late entry": "yes" if t["late_entry"] else "",
            "Reimbursement": t["reimbursement_id"] or "",
        })
    return rows


# ── Free-text / voice parsing (fallback when no AI is in the loop) ─────────

_CAT_WORDS = [
    ("Supplies", ("supply", "supplies", "home depot", "lowe", "hardware", "parts", "materials")),
    ("Bank", ("bank", "deposit", "atm")),
    ("Equipment", ("repair", "mechanic", "equipment", "tire", "oil change")),
    ("Estimate/Quote", ("estimate", "quote", "bid", "walkthrough", "walk-through")),
    ("Client meeting", ("meeting", "client", "customer meeting")),
    ("Office errand", ("post office", "usps", "ups", "fedex", "office", "printing")),
    ("Training", ("training", "class", "course", "license", "licensing")),
]


def parse_trip_text(text: str) -> dict:
    """Best-effort parse of 'Home Depot for ladders, 12 miles, round trip'.
    Returns the fields it found plus a 'missing' list. Never guesses miles
    or a destination it can't see in the text."""
    t = (text or "").strip()
    low = t.lower()
    out = {"text": t, "round_trip": bool(re.search(r"round[\s-]?trip|there and back", low)),
           "miles": None, "to_label": "", "purpose": "", "category": "Other",
           "trip_date": _today(), "missing": []}
    m = re.search(r"(\d+(?:\.\d+)?)\s*(?:mi\b|miles?\b)", low)
    if m:
        out["miles"] = float(m.group(1))
        if out["round_trip"] and re.search(r"(total|both ways|round[\s-]?trip)\s*(?:of\s*)?" + re.escape(m.group(1)), low):
            out["miles"] = out["miles"] / 2
    if "yesterday" in low:
        out["trip_date"] = (_dt.date.today() - _dt.timedelta(days=1)).isoformat()
    dm = re.search(r"(20\d\d-\d\d-\d\d)", low)
    if dm:
        out["trip_date"] = dm.group(1)
    for cat, words in _CAT_WORDS:
        if any(w in low for w in words):
            out["category"] = cat
            break
    # Destination = the run of Capitalized words after "to"/"at" (stops at
    # the first lowercase word, e.g. "at Chase this afternoon" -> "Chase").
    dest = re.search(r"\b(?:to|at)\s+(?:the\s+)?([A-Z][\w'&.-]*(?:\s+[A-Z0-9][\w'&.-]*)*)", t)
    if dest:
        out["to_label"] = dest.group(1).strip()
    else:
        lead = re.match(r"\s*(?:drove\s+to\s+)?([A-Z][\w'&.-]*(?:\s+[A-Z][\w'&.-]*)*)", t)
        if lead:
            out["to_label"] = lead.group(1).strip()
    pm = re.search(r"\bfor\s+([^,.;]+)", t, re.I)
    if pm:
        out["purpose"] = f"{out['category'] if out['category'] != 'Other' else 'Business'} – {pm.group(1).strip()}"
    elif out["category"] != "Other" and out["to_label"]:
        out["purpose"] = f"{out['category']} – {out['to_label']}"
    if not out["to_label"]:
        out["missing"].append("destination")
    if out["miles"] is None:
        out["missing"].append("miles (or a from address so they can be calculated)")
    if validate_purpose(out["purpose"]):
        out["missing"].append("business purpose")
    return out


# ══════════════════════════════════════════════════════════════════════════
# Phases 4–5 (2026-10-24): miles follow schedule changes (FR-21), "Cancelled
# on arrival", crew reimbursement approval with a paid status (FR-23),
# exports to CSV / Excel / PDF (FR-17), reminders (FR-24), and the AI
# mileage review (FR-25).
# ══════════════════════════════════════════════════════════════════════════

SETTING_REIMB_PERIOD = "Mileage Reimbursement Period"   # weekly | biweekly | monthly
SETTING_PERIOD_ANCHOR = "Mileage Pay Period Start"      # a date weekly/biweekly periods count from
SETTING_REMINDERS = "Mileage Reminders"                 # JSON {"gaps": true, ...}
DEFAULT_PERIOD_KIND = "monthly"
DEFAULT_PERIOD_ANCHOR = "2026-01-05"                    # a Monday
PERIOD_KINDS = ("weekly", "biweekly", "monthly")
PAY_METHODS = ("check", "cash", "payroll", "transfer")
_FROZEN = ("submitted", "approved", "paid")
REMINDER_TYPES = {
    "gaps": "Daily, 7 PM — your mileage today has gaps",
    "weekly": "Monday — review last week's miles (owner gets a summary)",
    "pay_period": "End of pay period — submit, approve, and pay crew reimbursements",
    "odometer": "Dec 31 and new or retired vehicles — record the odometer",
}
# 2026-10-25: each reminder is its own Enabled/Disabled row on Database →
# Settings (older versions kept them all in one JSON row, SETTING_REMINDERS,
# which migrate_legacy_settings() converts).
REMINDER_SETTING_KEYS = {
    "gaps": "Mileage Reminder - Daily Gaps",
    "weekly": "Mileage Reminder - Weekly Review",
    "pay_period": "Mileage Reminder - Pay Period",
    "odometer": "Mileage Reminder - Odometer",
}


def migrate_legacy_settings(conn) -> int:
    """Convert pre-2026-10-25 mileage settings to Settings-tab rows, so seeding
    the new defaults never overrides a choice someone already made:
      * Mileage Tracking Enabled true/false -> Enabled/Disabled
      * the Mileage Reminders JSON row     -> one Enabled/Disabled row each
        (the JSON row is removed once its values are copied).
    Works on an open connection; the caller commits. Returns rows changed."""
    changed = 0
    now = utcnow_iso()
    row = conn.execute("SELECT value FROM settings WHERE key = ?", (SETTING_ENABLED,)).fetchone()
    if row is not None:
        raw = str(row[0] if row[0] is not None else "").strip()
        if raw and raw not in ("Enabled", "Disabled"):
            new = "Enabled" if _onoff(raw, True) else "Disabled"
            conn.execute("UPDATE settings SET value = ? WHERE key = ?", (new, SETTING_ENABLED))
            changed += 1
    row = conn.execute("SELECT value FROM settings WHERE key = ?", (SETTING_REMINDERS,)).fetchone()
    if row is not None:
        try:
            m = json.loads(row[0]) if row[0] else {}
        except Exception:
            m = {}
        if not isinstance(m, dict):
            m = {}
        for name, key in REMINDER_SETTING_KEYS.items():
            if name not in m:
                continue
            cur = conn.execute(
                "INSERT INTO settings (key, value, last_edited_by, last_edited_at) "
                "VALUES (?, ?, 'system', ?) ON CONFLICT(key) DO NOTHING",
                (key, "Enabled" if bool(m[name]) else "Disabled", now))
            changed += cur.rowcount if cur.rowcount and cur.rowcount > 0 else 0
        conn.execute("DELETE FROM settings WHERE key = ?", (SETTING_REMINDERS,))
        changed += 1
    return changed


# ── Miles follow schedule changes (FR-21) ─────────────────────────────────
# Each built day stores a fingerprint of everything its trips came from
# (that driver's clock entries, their jobs' address/status, and the day's
# route legs). When the fingerprint no longer matches — a job moved,
# cancelled, deleted, re-addressed, a clock time fixed, a route rebuilt —
# the day is rebuilt the next time anyone looks at it, whichever tool made
# the change. Hand-edited, locked, and paid/submitted trips stay as they are.

def _day_signatures(conn, d: str):
    rows = conn.execute(
        "SELECT t.entry_id, t.job_id, t.clock_in, t.clock_out, t.clock_in_gps, t.clock_out_gps, "
        "t.crew_user_id, t.crew, j.job_status, j.latitude, j.longitude, j.street_address, j.city, "
        "j.state, j.zip, j.customer_name, j.service_type FROM time_entries t "
        "LEFT JOIN jobs j ON j.job_id = t.job_id WHERE t.entry_date = ? ORDER BY t.entry_id", (d,)).fetchall()
    route = [list(r) for r in conn.execute(
        "SELECT job_id, leg_drive_miles FROM route_stops WHERE route_date = ? AND job_id IS NOT NULL "
        "ORDER BY job_id, id", (d,)).fetchall()]
    by: dict = {}
    for r in rows:
        by.setdefault(entry_key(r["crew_user_id"], r["crew"]), []).append(list(r))

    def sig(v):
        return hashlib.sha1(json.dumps([v, route], default=str).encode("utf-8")).hexdigest()
    return {k: sig(v) for k, v in by.items()}, sig([])


def _store_signatures(conn, d: str, keys) -> None:
    sigs, empty = _day_signatures(conn, d)
    now = utcnow_iso()
    for k in keys:
        conn.execute("INSERT INTO mileage_day_builds (trip_date, user_key, signature, built_at) VALUES (?,?,?,?) "
                     "ON CONFLICT(trip_date, user_key) DO UPDATE SET signature = excluded.signature, "
                     "built_at = excluded.built_at", (d, k, sigs.get(k, empty), now))


def refresh_stale_days(db_path: str, start: str, end: str, actor: str = "automatic", limit: int = 31) -> list:
    """Rebuild days in [start, end] whose source data changed since they
    were built (FR-21). Returns the (date, driver) pairs rebuilt. Days built
    before this existed get their current fingerprint as a starting point.
    Never raises."""
    try:
        if not tracking_enabled(db_path):
            return []
        s, e = _parse_date(start), _parse_date(end)
        e = min(e, _today())
        conn = get_connection(db_path)
        try:
            pairs = set()
            for r in conn.execute("SELECT DISTINCT trip_date, build_key FROM mileage_trips WHERE trip_date "
                                  "BETWEEN ? AND ? AND origin IN ('auto','backfill') AND status = 'active'",
                                  (s, e)).fetchall():
                parts = (r["build_key"] or "").split("|")
                if len(parts) >= 3:
                    pairs.add((r["trip_date"], parts[1]))
            stored = {(r["trip_date"], r["user_key"]): r["signature"] for r in conn.execute(
                "SELECT trip_date, user_key, signature FROM mileage_day_builds WHERE trip_date BETWEEN ? AND ?",
                (s, e)).fetchall()}
            todo, cache = [], {}
            for d, k in sorted(pairs):
                if d not in cache:
                    cache[d] = _day_signatures(conn, d)
                sigs, empty = cache[d]
                cur = sigs.get(k, empty)
                if (d, k) not in stored:
                    conn.execute("INSERT OR IGNORE INTO mileage_day_builds (trip_date, user_key, signature, "
                                 "built_at) VALUES (?,?,?,?)", (d, k, cur, utcnow_iso()))
                    continue
                if stored[(d, k)] != cur and not _is_locked(conn, d):
                    todo.append((d, k))
        finally:
            conn.close()
        done = []
        for d, k in todo[:limit]:
            if k == "owner":
                build_trips_for_day(db_path, d, actor)
            else:
                build_trips_for_day(db_path, d, actor, crew_user_id=k)
            done.append((d, k))
        return done
    except Exception:
        return []


def cancelled_on_arrival(db_path: str, job_id: str, actor: str, *, crew_user_id: str = "",
                         crew_display: str = "", gps: str = "", note: str = "") -> str:
    """Customer cancelled when the crew arrived (FR-21): record a short
    clock-in/out at the job (so the drive there is kept), mark the job
    Cancelled with a note, and rebuild today's trips — the leg there stays
    business with purpose '<customer> – cancelled on site'."""
    from db_write_ops import db_log_time_entry
    conn = get_connection(db_path)
    try:
        job = conn.execute("SELECT * FROM jobs WHERE job_id = ?", (job_id,)).fetchone()
    finally:
        conn.close()
    if job is None:
        return f"❌ No job {job_id}."
    st = (job["job_status"] or "").strip().lower()
    if st in ("complete", "completed"):
        return f"❌ {job_id} is already Complete."
    r1 = db_log_time_entry(db_path, job_id, "start", crew_user_id, crew_display, gps)
    if not (r1.startswith("⏱") or "already open" in r1):
        return f"❌ Could not record the visit: {r1}"
    r2 = db_log_time_entry(db_path, job_id, "stop", crew_user_id, crew_display, gps)
    if not r2.startswith("⏱"):
        return f"❌ Could not record the visit: {r2}"
    today = _today()
    tag = f"[Cancelled on arrival {today} ({actor})" + (f": {note.strip()}" if (note or "").strip() else "") + "]"
    conn = get_connection(db_path)
    try:
        conn.execute("UPDATE jobs SET job_status = 'Cancelled', service_details = TRIM(COALESCE(service_details,'') "
                     "|| ' ' || ?), last_edited_by = ?, last_edited_at = ?, version = version + 1 WHERE job_id = ?",
                     (tag, actor, utcnow_iso(), job_id))
    finally:
        conn.close()
    built = ""
    if tracking_enabled(db_path):
        s = build_trips_for_day(db_path, today, actor, crew_user_id=crew_user_id,
                                crew=crew_display if crew_user_id else "")
        built = "\n" + summarize_build(s)
    cust = job["customer_name"] or job_id
    return (f"✅ {job_id} ({cust}) marked Cancelled on arrival. The drive there stays a business trip "
            f"(\"{cust} – cancelled on site\").{built}")


# ── Crew reimbursement: Submitted → Approved / Returned → Paid (FR-23) ────

def get_reimb_settings(db_path: str) -> dict:
    conn = get_connection(db_path)
    try:
        kind = (_get_setting(conn, SETTING_REIMB_PERIOD) or DEFAULT_PERIOD_KIND).strip().lower()
        anchor = (_get_setting(conn, SETTING_PERIOD_ANCHOR) or DEFAULT_PERIOD_ANCHOR).strip()
    finally:
        conn.close()
    if kind not in PERIOD_KINDS:
        kind = DEFAULT_PERIOD_KIND
    try:
        anchor = _parse_date(anchor)
    except ValueError:
        anchor = DEFAULT_PERIOD_ANCHOR
    return {"kind": kind, "anchor": anchor}


def set_reimb_period(db_path: str, kind: str, actor: str, anchor: str = "") -> str:
    kind = (kind or "").strip().lower().replace("-", "").replace(" ", "")
    kind = {"biweekly": "biweekly", "everyotherweek": "biweekly", "weekly": "weekly",
            "monthly": "monthly"}.get(kind, kind)
    if kind not in PERIOD_KINDS:
        return "❌ Pay period must be weekly, biweekly, or monthly."
    conn = get_connection(db_path)
    try:
        _set_setting(conn, SETTING_REIMB_PERIOD, kind, actor)
        if anchor:
            _set_setting(conn, SETTING_PERIOD_ANCHOR, _parse_date(anchor), actor)
    finally:
        conn.close()
    return f"✅ Crew mileage reimbursement runs {kind}."


def period_bounds(db_path: str, date_str: str):
    """(start, end, kind) of the reimbursement pay period containing date_str."""
    cfg = get_reimb_settings(db_path)
    d = _dt.date.fromisoformat(_parse_date(date_str))
    if cfg["kind"] == "monthly":
        s = d.replace(day=1)
        e = (s + _dt.timedelta(days=32)).replace(day=1) - _dt.timedelta(days=1)
    else:
        n = 7 if cfg["kind"] == "weekly" else 14
        a = _dt.date.fromisoformat(cfg["anchor"])
        s = d - _dt.timedelta(days=(d - a).days % n)
        e = s + _dt.timedelta(days=n - 1)
    return s.isoformat(), e.isoformat(), cfg["kind"]


def _frozen_status(conn, trip):
    """'submitted' / 'approved' / 'paid' when the trip belongs to a crew
    reimbursement that's waiting for approval or done; else None."""
    try:
        rid = trip["reimbursement_id"]
    except (KeyError, IndexError):
        return None
    if not rid:
        return None
    r = conn.execute("SELECT status FROM mileage_reimbursements WHERE reimbursement_id = ?", (rid,)).fetchone()
    return r["status"] if (r and r["status"] in _FROZEN) else None


def _period_frozen_for(conn, ukey: str, d: str):
    r = conn.execute("SELECT reimbursement_id, status FROM mileage_reimbursements WHERE user_key = ? AND "
                     "period_start <= ? AND period_end >= ? AND status IN ('submitted','approved','paid')",
                     (ukey, d, d)).fetchone()
    return (r["reimbursement_id"], r["status"]) if r else None


def _reimb_event(conn, rid, frm, to, actor, note=""):
    conn.execute("INSERT INTO mileage_reimbursement_events (reimbursement_id, from_status, to_status, actor, at, note) "
                 "VALUES (?,?,?,?,?,?)", (rid, frm, to, actor, utcnow_iso(), note or None))


def _reimb_trip_rows(db_path: str, ukey: str, s: str, e: str) -> list:
    return [t for t in list_trips(db_path, s, e, False, "", ukey)
            if t["reimbursable"] and t["classification"] == "business" and t["status"] == "active"]


def _sum_trips(trips) -> dict:
    miles = sum(float(t["miles"] or 0) for t in trips)
    amount = sum(float(t["miles"] or 0) * float(t["rate_cents"] or 0) / 100 for t in trips)
    pt = sum(float(t["parking_usd"] or 0) + float(t["tolls_usd"] or 0) for t in trips)
    return {"trip_count": len(trips), "miles": round(miles, 1), "amount_usd": round(amount, 2),
            "parking_tolls_usd": round(pt, 2)}


def reimbursement_preview(db_path: str, ukey: str, period_date: str = "") -> dict:
    s, e, kind = period_bounds(db_path, period_date or _today())
    trips = _reimb_trip_rows(db_path, ukey, s, e)
    conn = get_connection(db_path)
    try:
        row = conn.execute("SELECT * FROM mileage_reimbursements WHERE user_key = ? AND period_start = ?",
                           (ukey, s)).fetchone()
    finally:
        conn.close()
    gaps = day_gaps(db_path, s, min(e, _today()), True, ukey)
    out = {"user_key": ukey, "period_start": s, "period_end": e, "kind": kind, "trips": trips,
           "gaps": gaps, "existing": dict(row) if row else None,
           "crew": (trips[0]["crew"] if trips else (row["crew"] if row else "")),
           "missing_rate": [t["trip_id"] for t in trips if t["rate_cents"] is None]}
    out.update(_sum_trips(trips))
    return out


def submit_reimbursement(db_path: str, actor: str, ukey: str, crew_user_id: str = "", crew: str = "",
                         period_date: str = "") -> str:
    from db_write_ops import generate_next_id
    p = reimbursement_preview(db_path, ukey, period_date)
    s, e, ex = p["period_start"], p["period_end"], p["existing"]
    if ex and ex["status"] in _FROZEN:
        return f"ℹ️ {ex['reimbursement_id']} for {s} – {e} is already {ex['status']}."
    if p["gaps"]:
        days = sorted({g["date"] for g in p["gaps"]})
        return ("❌ Can't submit yet — these days need a fix first (Mileage tab → Needs fix): "
                + ", ".join(days) + ".")
    if not p["trips"]:
        return f"ℹ️ No reimbursable trips (business trips in a crew-owned vehicle) for {s} – {e}."
    if p["missing_rate"]:
        return "❌ Some trips are before the first IRS rate in Settings → Mileage rates; add that rate first."
    name = crew or p["crew"] or ukey
    conn = get_connection(db_path)
    try:
        conn.execute("BEGIN IMMEDIATE")
        now = utcnow_iso()
        if ex:
            rid = ex["reimbursement_id"]
            conn.execute("UPDATE mileage_reimbursements SET status = 'submitted', trip_count = ?, miles = ?, "
                         "amount_usd = ?, parking_tolls_usd = ?, submitted_by = ?, submitted_at = ?, "
                         "approved_by = NULL, approved_at = NULL, last_edited_by = ?, last_edited_at = ?, "
                         "version = version + 1 WHERE reimbursement_id = ?",
                         (p["trip_count"], p["miles"], p["amount_usd"], p["parking_tolls_usd"], actor, now,
                          actor, now, rid))
            frm = ex["status"]
        else:
            rid = generate_next_id(conn, "mileage_reimbursements", "reimbursement_id", "REIM", 4)
            conn.execute("INSERT INTO mileage_reimbursements (reimbursement_id, user_key, crew_user_id, crew, "
                         "period_start, period_end, period_kind, status, trip_count, miles, amount_usd, "
                         "parking_tolls_usd, submitted_by, submitted_at, created_at, last_edited_by, last_edited_at) "
                         "VALUES (?,?,?,?,?,?,?,'submitted',?,?,?,?,?,?,?,?,?)",
                         (rid, ukey, crew_user_id or "", name, s, e, p["kind"], p["trip_count"], p["miles"],
                          p["amount_usd"], p["parking_tolls_usd"], actor, now, now, actor, now))
            frm = "draft"
        conn.execute("UPDATE mileage_trips SET reimbursement_id = NULL WHERE reimbursement_id = ?", (rid,))
        conn.executemany("UPDATE mileage_trips SET reimbursement_id = ? WHERE trip_id = ?",
                         [(rid, t["trip_id"]) for t in p["trips"]])
        _reimb_event(conn, rid, frm, "submitted", actor)
        conn.execute("COMMIT")
    except Exception as exc:
        try:
            conn.execute("ROLLBACK")
        except Exception:
            pass
        return f"❌ Could not submit: {exc}"
    finally:
        conn.close()
    return (f"✅ {rid} submitted for {name}: {s} – {e}, {p['trip_count']} trips, {p['miles']:,.1f} mi, "
            f"${p['amount_usd']:,.2f}" + (f" + ${p['parking_tolls_usd']:,.2f} parking/tolls"
                                          if p["parking_tolls_usd"] else "")
            + ". Waiting for the owner's approval; these trips can't be changed unless it's returned.")


def _get_reimb(conn, rid):
    return conn.execute("SELECT * FROM mileage_reimbursements WHERE reimbursement_id = ?", (rid,)).fetchone()


def approve_reimbursement(db_path: str, rid: str, actor: str, approver_key: str) -> str:
    rates = get_rates(db_path)
    conn = get_connection(db_path)
    try:
        r = _get_reimb(conn, rid)
        if r is None:
            return f"❌ No reimbursement {rid}."
        if approver_key and approver_key == r["user_key"]:
            return "❌ You can't approve your own reimbursement — another owner or manager has to."
        if r["status"] != "submitted":
            return f"❌ {rid} is {r['status']} — only a submitted reimbursement can be approved."
        trips = conn.execute("SELECT * FROM mileage_trips WHERE reimbursement_id = ? AND status = 'active'",
                             (rid,)).fetchall()
        miles = sum(float(t["miles"] or 0) for t in trips)
        amount = round(sum(float(t["miles"] or 0) * (rate_cents_for(rates, t["trip_date"]) or 0) / 100
                           for t in trips), 2)
        pt = round(sum(float(t["parking_usd"] or 0) + float(t["tolls_usd"] or 0) for t in trips), 2)
        conn.execute("BEGIN IMMEDIATE")
        now = utcnow_iso()
        conn.execute("UPDATE mileage_reimbursements SET status = 'approved', approved_by = ?, approved_at = ?, "
                     "trip_count = ?, miles = ?, amount_usd = ?, parking_tolls_usd = ?, last_edited_by = ?, "
                     "last_edited_at = ?, version = version + 1 WHERE reimbursement_id = ?",
                     (actor, now, len(trips), round(miles, 1), amount, pt, actor, now, rid))
        _reimb_event(conn, rid, "submitted", "approved", actor)
        conn.execute("COMMIT")
    finally:
        conn.close()
    return (f"✅ {rid} approved for {r['crew'] or r['user_key']}: ${amount + pt:,.2f} "
            f"({miles:,.1f} mi = ${amount:,.2f}" + (f", parking/tolls ${pt:,.2f}" if pt else "")
            + "). The amount is now fixed. Mark it paid when you pay it.")


def return_reimbursement(db_path: str, rid: str, note: str, actor: str, approver_key: str) -> str:
    if not (note or "").strip():
        return "❌ Say what needs fixing — the note goes back to the crew member."
    conn = get_connection(db_path)
    try:
        r = _get_reimb(conn, rid)
        if r is None:
            return f"❌ No reimbursement {rid}."
        if approver_key and approver_key == r["user_key"]:
            return "❌ You can't return your own reimbursement — another owner or manager has to."
        if r["status"] not in ("submitted", "approved"):
            return f"❌ {rid} is {r['status']} — only a submitted or approved (unpaid) one can be returned."
        conn.execute("BEGIN IMMEDIATE")
        now = utcnow_iso()
        conn.execute("UPDATE mileage_reimbursements SET status = 'returned', return_note = ?, approved_by = NULL, "
                     "approved_at = NULL, last_edited_by = ?, last_edited_at = ?, version = version + 1 "
                     "WHERE reimbursement_id = ?", (note.strip(), actor, now, rid))
        _reimb_event(conn, rid, r["status"], "returned", actor, note.strip())
        conn.execute("COMMIT")
    finally:
        conn.close()
    return (f"↩️ {rid} returned to {r['crew'] or r['user_key']}: {note.strip()}. "
            "Its trips can be edited again; resubmit when fixed.")


def mark_reimbursement_paid(db_path: str, rid: str, actor: str, approver_key: str, paid_date: str = "",
                            method: str = "", reference: str = "") -> str:
    m = (method or "").strip().lower()
    if m not in PAY_METHODS:
        return "❌ Payment method must be check, cash, payroll, or transfer."
    pd = _parse_date(paid_date) if paid_date else _today()
    conn = get_connection(db_path)
    try:
        r = _get_reimb(conn, rid)
        if r is None:
            return f"❌ No reimbursement {rid}."
        if approver_key and approver_key == r["user_key"]:
            return "❌ You can't mark your own reimbursement paid — another owner or manager has to."
        if r["status"] != "approved":
            return f"❌ {rid} is {r['status']} — approve it before marking it paid."
        conn.execute("BEGIN IMMEDIATE")
        now = utcnow_iso()
        conn.execute("UPDATE mileage_reimbursements SET status = 'paid', paid_date = ?, paid_method = ?, "
                     "paid_reference = ?, paid_by = ?, paid_at = ?, last_edited_by = ?, last_edited_at = ?, "
                     "version = version + 1 WHERE reimbursement_id = ?",
                     (pd, m, (reference or "").strip() or None, actor, now, actor, now, rid))
        conn.execute("UPDATE mileage_trips SET locked = 1 WHERE reimbursement_id = ?", (rid,))
        _reimb_event(conn, rid, "approved", "paid", actor, f"{m}{' ' + reference.strip() if reference else ''}")
        conn.execute("COMMIT")
    finally:
        conn.close()
    total = float(r["amount_usd"] or 0) + float(r["parking_tolls_usd"] or 0)
    return (f"✅ {rid} marked paid {pd} by {m}" + (f" (ref {reference.strip()})" if reference else "")
            + f": ${total:,.2f} to {r['crew'] or r['user_key']}. Its trips are now locked.")


def list_reimbursements(db_path: str, user_key_filter: str = "", statuses=None, limit: int = 100) -> list:
    conn = get_connection(db_path)
    try:
        sql, args = "SELECT * FROM mileage_reimbursements", []
        where = []
        if user_key_filter:
            where.append("user_key = ?")
            args.append(user_key_filter)
        if statuses:
            where.append("status IN (%s)" % ",".join("?" * len(statuses)))
            args.extend(statuses)
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY period_start DESC, reimbursement_id DESC LIMIT ?"
        args.append(int(limit))
        out = []
        for r in conn.execute(sql, args).fetchall():
            d = dict(r)
            d["total_usd"] = round(float(r["amount_usd"] or 0) + float(r["parking_tolls_usd"] or 0), 2)
            d["events"] = [dict(x) for x in conn.execute(
                "SELECT from_status, to_status, actor, at, note FROM mileage_reimbursement_events "
                "WHERE reimbursement_id = ? ORDER BY id", (r["reimbursement_id"],)).fetchall()]
            out.append(d)
        return out
    finally:
        conn.close()


def reimbursement_status_summary(db_path: str, start: str, end: str, restrict: bool = False,
                                 ukey: str = "") -> dict:
    """Per driver: amounts not yet submitted (incl. returned), submitted,
    approved but unpaid, and paid — for crew-owned-vehicle business trips."""
    trips = [t for t in list_trips(db_path, start, end, restrict, ukey)
             if t["reimbursable"] and t["classification"] == "business"]
    conn = get_connection(db_path)
    try:
        st = {r["reimbursement_id"]: r["status"] for r in conn.execute(
            "SELECT reimbursement_id, status FROM mileage_reimbursements").fetchall()}
    finally:
        conn.close()
    out: dict = {}
    for t in trips:
        s = st.get(t["reimbursement_id"] or "", "")
        bucket = {"submitted": "submitted", "approved": "approved", "paid": "paid"}.get(s, "not_submitted")
        who = t["crew"] or t["driver_key"]
        d = out.setdefault(who, {"not_submitted": 0.0, "submitted": 0.0, "approved": 0.0, "paid": 0.0})
        d[bucket] += float(t["miles"] or 0) * float(t["rate_cents"] or 0) / 100 \
            + float(t["parking_usd"] or 0) + float(t["tolls_usd"] or 0)
    for d in out.values():
        for k in d:
            d[k] = round(d[k], 2)
    return out


def format_reimbursements(rows: list, preview: dict = None) -> str:
    L = []
    if preview:
        p = preview
        st = p["existing"]["status"] if p["existing"] else "not submitted"
        L.append(f"🧾 {p['crew'] or p['user_key']} — pay period {p['period_start']} – {p['period_end']} "
                 f"({p['kind']}): {p['trip_count']} trips, {p['miles']:,.1f} mi, ${p['amount_usd']:,.2f}"
                 + (f" + ${p['parking_tolls_usd']:,.2f} parking/tolls" if p["parking_tolls_usd"] else "")
                 + f" — {st}.")
        if p["gaps"]:
            L.append("   ⚠️ Needs a fix before it can be submitted: "
                     + ", ".join(sorted({g['date'] for g in p['gaps']})))
    if rows:
        L.append("Reimbursements:")
        for r in rows:
            extra = ""
            if r["status"] == "paid":
                extra = f", paid {r['paid_date']} by {r['paid_method']}" + (
                    f" ref {r['paid_reference']}" if r.get("paid_reference") else "")
            elif r["status"] == "returned" and r.get("return_note"):
                extra = f" — returned: {r['return_note']}"
            L.append(f"  {r['reimbursement_id']} {r['crew'] or r['user_key']} {r['period_start']} – "
                     f"{r['period_end']}: ${r['total_usd']:,.2f} ({float(r['miles'] or 0):,.1f} mi) "
                     f"— {r['status']}{extra}")
    elif not preview:
        L.append("ℹ️ No crew reimbursements yet.")
    return "\n".join(L)


# ── Exports: CSV, Excel, PDF (FR-17) ──────────────────────────────────────

def _business_name(db_path: str) -> str:
    try:
        conn = get_connection(db_path)
        try:
            return (_get_setting(conn, "Business Name") or "").strip()
        finally:
            conn.close()
    except Exception:
        return ""


def export_mileage(db_path: str, start: str, end: str, fmt: str = "pdf", restrict: bool = False,
                   ukey: str = "", driver_key: str = "", prepared_for: str = ""):
    """(filename, bytes, mime) for the mileage log over [start, end]."""
    import io
    s, e = _parse_date(start), _parse_date(end)
    fmt = (fmt or "pdf").strip().lower().lstrip(".")
    fmt = {"excel": "xlsx", "xls": "xlsx"}.get(fmt, fmt)
    if fmt not in ("pdf", "csv", "xlsx"):
        raise ValueError("format must be pdf, csv, or xlsx")
    rep = mileage_report(db_path, s, e, restrict, ukey, driver_key)
    rows = export_rows(db_path, s, e, restrict, ukey, driver_key)
    base = f"Mileage_{s}_to_{e}"
    if fmt == "csv":
        import csv
        buf = io.StringIO()
        cols = list(rows[0].keys()) if rows else ["Trip ID", "Date", "Driver", "Vehicle", "VIN", "From", "To",
                                                  "Miles", "Purpose", "Classification"]
        w = csv.DictWriter(buf, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow(r)
        return base + ".csv", buf.getvalue().encode("utf-8-sig"), "text/csv"
    if fmt == "pdf":
        import mileage_pdf
        return (base + ".pdf", mileage_pdf.render_log_pdf(rep, rows, _business_name(db_path), prepared_for),
                "application/pdf")
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    bold = Font(bold=True)

    def sheet(ws, header, data, widths=None):
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        for row in data:
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, wdt in enumerate(widths or [], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = wdt

    ws = wb.active
    ws.title = "Trips"
    hdr = list(rows[0].keys()) if rows else ["Trip ID", "Date"]
    sheet(ws, hdr, [[r[k] for k in hdr] for r in rows], [11, 11, 7, 7, 14, 24, 19, 10, 18, 26, 8, 34, 14, 12])
    ws = wb.create_sheet("Summary")
    ws.append([f"Business mileage {s} to {e}", _business_name(db_path)])
    ws["A1"].font = bold
    ws.append([])
    ws.append(["Deduction (business-owned vehicles)", "Rate", "Business miles", "Amount ($)"])
    for c in ws[3]:
        c.font = bold
    tot = 0.0
    for k, p in sorted(rep["deduction"].items(), key=lambda kv: kv[1]["from"] or ""):
        ws.append([f"{p['from']} to {p['to']}", k, round(p["miles"], 1), round(p["amount"], 2)])
        tot += p["amount"]
    ws.append(["Parking and tolls", "", "", round(rep["parking_tolls"], 2)])
    ws.append(["Total deduction", "", round(sum(p["miles"] for p in rep["deduction"].values()), 1),
               round(tot + rep["parking_tolls"], 2)])
    ws.cell(ws.max_row, 1).font = bold
    for title, key in (("By driver", "by_driver"), ("By vehicle", "by_vehicle"), ("By month", "by_month")):
        ws.append([])
        ws.append([title, "Business mi", "Commute mi", "Personal mi"])
        for c in ws[ws.max_row]:
            c.font = bold
        for k, v in sorted(rep[key].items()):
            ws.append([k, round(v["business"], 1), round(v["commute"], 1), round(v["personal"], 1)])
    ws.append([])
    ws.append(["IRS record checks"])
    ws.cell(ws.max_row, 1).font = bold
    for m in rep["compliance"] or ["All trips have a date, destination, purpose, miles, and vehicle."]:
        ws.append([m])
    ws.column_dimensions["A"].width = 46
    for col in "BCD":
        ws.column_dimensions[col].width = 15
    ws = wb.create_sheet("Reimbursement")
    stat = rep.get("reimbursement_status") or {}
    data = []
    for who, per in sorted(rep["reimbursement"].items()):
        st = stat.get(who, {})
        data.append([who, round(sum(p["miles"] for p in per.values()), 1),
                     round(sum(p["amount"] for p in per.values()), 2), st.get("not_submitted", 0),
                     st.get("submitted", 0), st.get("approved", 0), st.get("paid", 0)])
    sheet(ws, ["Driver", "Miles", "Amount ($)", "Not submitted ($)", "Submitted ($)", "Approved, unpaid ($)",
               "Paid ($)"], data, [22, 10, 12, 16, 14, 18, 12])
    ws.append([])
    ws.append(["Statement", "Driver", "Period start", "Period end", "Status", "Miles", "Total ($)",
               "Paid date", "Method", "Reference"])
    for c in ws[ws.max_row]:
        c.font = bold
    for r in list_reimbursements(db_path, ukey if restrict else ""):
        if r["period_end"] < s or r["period_start"] > e:
            continue
        ws.append([r["reimbursement_id"], r["crew"] or r["user_key"], r["period_start"], r["period_end"],
                   r["status"], r["miles"], r["total_usd"], r["paid_date"] or "", r["paid_method"] or "",
                   r["paid_reference"] or ""])
    ws = wb.create_sheet("Vehicles")
    sheet(ws, ["Year", "Vehicle ID", "Vehicle", "VIN", "Placed in service", "Business mi", "Commute mi",
               "Other personal mi", "Total mi (odometer)"],
          [[v["year"], v["vehicle_id"], v["label"], v["vin"], v["placed_in_service"], v["business"],
            v["commute"], v["personal"] if v["personal"] is not None else "needs odometer",
            v["total"] if v["total"] is not None else "needs odometer"] for v in rep["vehicles_4562"]],
          [7, 11, 26, 20, 15, 12, 12, 16, 18])
    buf = io.BytesIO()
    wb.save(buf)
    return (base + ".xlsx", buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def statement_pdf(db_path: str, rid: str = "", ukey: str = "", period_date: str = ""):
    """(filename, bytes) — reimbursement statement for REIM-#### or, with
    no id, a preview of ukey's current (or given) pay period."""
    import mileage_pdf
    if rid:
        conn = get_connection(db_path)
        try:
            r = _get_reimb(conn, rid)
        finally:
            conn.close()
        if r is None:
            raise ValueError(f"No reimbursement {rid}.")
        reim = dict(r)
        trips = [t for t in list_trips(db_path, r["period_start"], r["period_end"], False, "", r["user_key"])
                 if t["reimbursement_id"] == rid and t["status"] == "active"]
        name = f"Mileage_reimbursement_{rid}.pdf"
    else:
        p = reimbursement_preview(db_path, ukey, period_date)
        reim = {"user_key": ukey, "crew": p["crew"], "period_start": p["period_start"],
                "period_end": p["period_end"], "status": "", "amount_usd": p["amount_usd"],
                "parking_tolls_usd": p["parking_tolls_usd"]}
        trips = p["trips"]
        name = f"Mileage_reimbursement_{p['period_start']}_draft.pdf"
    return name, mileage_pdf.render_statement_pdf(reim, trips, _business_name(db_path))


# ── Reminders (FR-24) ─────────────────────────────────────────────────────

def get_reminder_settings(db_path: str) -> dict:
    """Each reminder's own Settings row wins; the legacy JSON row (not yet
    migrated) is the fallback; missing = on."""
    conn = get_connection(db_path)
    try:
        raw = _get_setting(conn, SETTING_REMINDERS)
        rows = {k: _get_setting(conn, key) for k, key in REMINDER_SETTING_KEYS.items()}
    finally:
        conn.close()
    try:
        m = json.loads(raw) if raw else {}
    except Exception:
        m = {}
    if not isinstance(m, dict):
        m = {}
    return {k: _onoff(rows.get(k), bool(m.get(k, True))) for k in REMINDER_TYPES}


def set_reminder(db_path: str, name: str, enabled: bool, actor: str) -> str:
    name = (name or "").strip().lower().replace(" ", "_").replace("-", "_")
    name = {"pay": "pay_period", "payperiod": "pay_period", "gap": "gaps", "week": "weekly",
            "odo": "odometer"}.get(name, name)
    if name not in REMINDER_TYPES:
        return "❌ Reminder must be one of: " + ", ".join(REMINDER_TYPES) + "."
    conn = get_connection(db_path)
    try:
        _set_setting(conn, REMINDER_SETTING_KEYS[name], "Enabled" if enabled else "Disabled", actor)
    finally:
        conn.close()
    return f"✅ Mileage reminder '{REMINDER_TYPES[name]}' turned {'ON' if enabled else 'OFF'}."


def due_reminders(db_path: str, ukey: str, is_admin: bool, now=None) -> list:
    """Reminders due right now for this person (in-app notices on the Jobs
    app, the owner's email digest and morning brief). Each is a dict with
    type, text and (where it helps) the action to take."""
    try:
        if not tracking_enabled(db_path):
            return []
        cfg = get_reminder_settings(db_path)
        now = now or _dt.datetime.now()
        today = now.date()
        t_iso = today.isoformat()
        out = []
        if cfg["gaps"]:
            g = [x for x in day_gaps(db_path, (today - _dt.timedelta(days=7)).isoformat(), t_iso, True, ukey)
                 if x["date"] < t_iso or now.hour >= 19]
            if g:
                days = sorted({x["date"] for x in g})
                out.append({"type": "gaps", "text": f"Your mileage has gaps on {', '.join(days)} — tap Fix "
                            "so those trips count.", "action": "fix"})
        if cfg["weekly"] and today.weekday() == 0:
            ls = (today - _dt.timedelta(days=7)).isoformat()
            le = (today - _dt.timedelta(days=1)).isoformat()
            mine = list_trips(db_path, ls, le, True, ukey)
            mi = sum(float(t["miles"] or 0) for t in mine if t["classification"] == "business")
            out.append({"type": "weekly", "text": f"Review last week's miles ({ls} – {le}): {len(mine)} trips, "
                        f"{mi:,.1f} business mi.", "action": "review"})
            if is_admin:
                allt = list_trips(db_path, ls, le)
                amt = sum(float(t["miles"] or 0) for t in allt if t["classification"] == "business")
                drivers = len({t["driver_key"] for t in allt})
                out.append({"type": "weekly", "text": f"Team last week: {len(allt)} trips, {amt:,.1f} business mi, "
                            f"{drivers} driver(s).", "action": "report"})
        if cfg["pay_period"]:
            cs, ce, kind = period_bounds(db_path, t_iso)
            ps, pe, _k = period_bounds(db_path, (_dt.date.fromisoformat(cs) - _dt.timedelta(days=1)).isoformat())
            for s_, e_, label in ((ps, pe, "last"), (cs, ce, "this")):
                if label == "this" and (_dt.date.fromisoformat(ce) - today).days > 1:
                    continue
                p = reimbursement_preview(db_path, ukey, s_)
                if p["trips"] and not (p["existing"] and p["existing"]["status"] in _FROZEN):
                    out.append({"type": "pay_period", "text": f"Submit your mileage for {s_} – {e_}: "
                                f"{p['miles']:,.1f} mi, ${p['amount_usd'] + p['parking_tolls_usd']:,.2f}.",
                                "action": "submit", "period": s_})
            if is_admin:
                waiting = list_reimbursements(db_path, statuses=["submitted"])
                unpaid = list_reimbursements(db_path, statuses=["approved"])
                if waiting:
                    out.append({"type": "pay_period", "text": f"{len(waiting)} mileage reimbursement(s) waiting "
                                "for approval.", "action": "approve"})
                if unpaid:
                    out.append({"type": "pay_period", "text": f"{len(unpaid)} approved, not yet paid "
                                f"(${sum(r['total_usd'] for r in unpaid):,.2f}).", "action": "pay"})
        if cfg["odometer"]:
            vs = list_vehicles(db_path, restrict_user_key="" if is_admin else ukey)
            if not is_admin:
                vs = [v for v in vs if v["ownership"] == "crew"]
            conn = get_connection(db_path)
            try:
                for v in vs:
                    if v["last_odometer"] is None:
                        out.append({"type": "odometer", "text": f"Record the odometer for {v['vehicle_id']} "
                                    f"{v['label']} — the IRS asks for total yearly miles.", "action": "odometer",
                                    "vehicle_id": v["vehicle_id"]})
                        continue
                    if today.month == 12 and today.day >= 24:
                        has = conn.execute("SELECT 1 FROM odometer_readings WHERE vehicle_id = ? AND reading_date "
                                           ">= ?", (v["vehicle_id"], f"{today.year}-12-20")).fetchone()
                        if not has:
                            out.append({"type": "odometer", "text": f"Year end: record the odometer for "
                                        f"{v['vehicle_id']} {v['label']} on Dec 31.", "action": "odometer",
                                        "vehicle_id": v["vehicle_id"]})
                    elif today.month == 1 and today.day <= 15:
                        has = conn.execute("SELECT 1 FROM odometer_readings WHERE vehicle_id = ? AND reading_date "
                                           "BETWEEN ? AND ?", (v["vehicle_id"], f"{today.year - 1}-12-20",
                                                               f"{today.year}-01-15")).fetchone()
                        if not has:
                            out.append({"type": "odometer", "text": f"New year: record the starting odometer for "
                                        f"{v['vehicle_id']} {v['label']}.", "action": "odometer",
                                        "vehicle_id": v["vehicle_id"]})
            finally:
                conn.close()
        return out
    except Exception:
        return []


def brief_lines(db_path: str, now=None) -> list:
    """Owner's morning / Monday brief: reminders, IRS warnings (FR-19), and
    on Mondays a short AI review summary of last week (FR-25)."""
    try:
        if not tracking_enabled(db_path):
            return []
        now = now or _dt.datetime.now()
        today = now.date()
        lines = [r["text"] for r in due_reminders(db_path, "owner", True, now)]
        if today.weekday() == 0:
            ls = (today - _dt.timedelta(days=7)).isoformat()
            le = (today - _dt.timedelta(days=1)).isoformat()
            f = review_mileage(db_path, ls, le)
            if f:
                kinds: dict = {}
                for x in f:
                    kinds[x["title"]] = kinds.get(x["title"], 0) + 1
                lines.append("AI mileage review of last week: " + "; ".join(
                    f"{n} × {k}" for k, n in sorted(kinds.items(), key=lambda kv: -kv[1]))
                    + " — ask Claude to \"review last week's mileage and fix the obvious problems\".")
            else:
                lines.append("AI mileage review of last week: no problems found.")
            rep = mileage_report(db_path, ls, le)
            lines += [f"IRS check: {m}" for m in rep["compliance"]]
        return lines
    except Exception:
        return []


# ── AI mileage review (FR-25) ─────────────────────────────────────────────
# Rule-based checks produce findings with a stable id and (where it's safe)
# a proposed fix. Nothing changes until someone accepts a finding; accepted
# fixes go through the same audited edit path as a hand edit. Dismissed
# findings aren't raised again.

_REVIEW_TITLES = {
    "route_mismatch": "logged miles far from the approved route",
    "odometer": "trips don't match the odometer",
    "long_leg": "unusually long leg",
    "duplicate": "possible duplicate trip",
    "no_jobs_day": "job trip on a day with no clock-in",
    "two_drivers": "vehicle used by two drivers the same day",
    "vague_purpose": "purpose too vague for the IRS",
    "no_vehicle": "trip with no vehicle",
    "commute": "commute leg may be misclassified",
    "late_edit": "edited long after the trip",
    "paid_changed": "paid reimbursement changed afterwards",
}


def _haversine_mi(a, b) -> float:
    import math
    la1, lo1, la2, lo2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2
    return 3958.8 * 2 * math.asin(math.sqrt(h))


def review_mileage(db_path: str, start: str, end: str, restrict: bool = False, ukey: str = "",
                   driver_key: str = "", include_dismissed: bool = False) -> list:
    s, e = _parse_date(start), _parse_date(end)
    trips = list_trips(db_path, s, e, restrict, ukey, driver_key)
    found: list = []

    def add(check, d, target, msg, severity="medium", trip_id=None, driver="", fix=None):
        found.append({"id": f"{check}:{d}:{target}", "check": check, "title": _REVIEW_TITLES[check],
                      "severity": severity, "date": d, "trip_id": trip_id, "driver": driver,
                      "message": msg, "fix": fix})

    conn = get_connection(db_path)
    try:
        mode = (conn.execute("SELECT value FROM settings WHERE key = 'Route Origin Mode'").fetchone() or [""])[0]
        # 1. Logged miles vs the approved route plan (> 25 % apart).
        by_day: dict = {}
        for t in trips:
            if t["origin"] in ("auto", "backfill") and t["classification"] == "business":
                by_day.setdefault((t["trip_date"], t["driver_key"]), []).append(t)
        for (d, k), ts in by_day.items():
            name = ts[0]["crew"] or ""
            crew_ids = {k, name, name.lower()} - {""}
            if k == "owner":
                rs = conn.execute("SELECT job_id, leg_drive_miles FROM route_stops WHERE route_date = ? "
                                  "AND leg_drive_miles IS NOT NULL", (d,)).fetchall()
            else:
                rs = [r for r in conn.execute("SELECT crew_id, job_id, leg_drive_miles FROM route_stops WHERE "
                                              "route_date = ? AND leg_drive_miles IS NOT NULL", (d,)).fetchall()
                      if (r["crew_id"] or "") in crew_ids or (r["crew_id"] or "").lower() in crew_ids]
            plan = {r["job_id"]: float(r["leg_drive_miles"]) for r in rs if r["job_id"]}
            if not plan:
                continue
            legs = [t for t in ts if t["job_id"] in plan]
            if len(legs) < 1:
                continue
            logged = sum(float(t["miles"] or 0) for t in legs)
            planned = sum(plan[t["job_id"]] for t in legs)
            if planned > 0 and abs(logged - planned) / planned > 0.25:
                worst = max(legs, key=lambda t: abs(float(t["miles"] or 0) - plan[t["job_id"]]))
                fix = None
                if abs(float(worst["miles"] or 0) - plan[worst["job_id"]]) >= 1:
                    fix = {"action": "edit", "trip_id": worst["trip_id"],
                           "changes": {"miles": round(plan[worst["job_id"]], 1)},
                           "summary": f"set {worst['trip_id']} ({worst['to_label']}) to the planned "
                                      f"{plan[worst['job_id']]:.1f} mi (now {float(worst['miles'] or 0):.1f})"}
                add("route_mismatch", d, k, f"{name or k} logged {logged:.1f} mi to jobs on {d}; the approved "
                    f"route planned {planned:.1f} mi ({(logged - planned) / planned:+.0%}).", "medium",
                    worst["trip_id"], name or k, fix)
        # 3. Unusually long legs for the distance between the two places.
        for t in trips:
            if t["from_lat"] is None or t["to_lat"] is None:
                continue
            a, b = (t["from_lat"], t["from_lon"]), (t["to_lat"], t["to_lon"])
            straight = _haversine_mi(a, b)
            m = float(t["miles"] or 0)
            if m > straight * 2.2 + 3:
                est = ROUTER(a, b)
                fix = None
                if est is not None and abs(est - m) / max(m, 1) > 0.10:
                    fix = {"action": "edit", "trip_id": t["trip_id"], "changes": {"miles": round(est, 1)},
                           "summary": f"set {t['trip_id']} to the road distance {est:.1f} mi (now {m:.1f})"}
                add("long_leg", t["trip_date"], t["trip_id"], f"{t['trip_id']} {t['from_label'] or 'Start'} → "
                    f"{t['to_label']}: {m:.1f} mi for places {straight:.1f} mi apart — a missed stop, a detour, "
                    "or a wrong GPS point?", "medium", t["trip_id"], t["crew"] or t["driver_key"], fix)
        # 4. Duplicates / overlaps.
        seen: dict = {}
        for t in trips:
            key = (t["trip_date"], t["driver_key"], (t["to_label"] or "").strip().lower())
            for o in seen.get(key, []):
                if o["origin"] != "manual" and t["origin"] != "manual":
                    continue          # two built legs to the same place = a real repeat visit
                if abs(float(o["miles"] or 0) - float(t["miles"] or 0)) <= 0.5 and \
                        (o["from_label"] or "").strip().lower() == (t["from_label"] or "").strip().lower():
                    dup = t if (t["origin"] == "manual" or o["origin"] != "manual") else o
                    keep = o if dup is t else t
                    add("duplicate", t["trip_date"], dup["trip_id"], f"{dup['trip_id']} looks like a copy of "
                        f"{keep['trip_id']} ({keep['from_label'] or 'Start'} → {keep['to_label']}, "
                        f"{float(keep['miles'] or 0):.1f} mi, same day).", "high", dup["trip_id"],
                        t["crew"] or t["driver_key"],
                        {"action": "void", "trip_id": dup["trip_id"],
                         "summary": f"remove {dup['trip_id']} (duplicate of {keep['trip_id']})"})
            seen.setdefault(key, []).append(t)
        # 5. Job trips on days with no clock-in for that job.
        for t in trips:
            if t["job_id"] and t["origin"] == "manual":
                r = conn.execute("SELECT 1 FROM time_entries WHERE job_id = ? AND entry_date = ?",
                                 (t["job_id"], t["trip_date"])).fetchone()
                if not r:
                    add("no_jobs_day", t["trip_date"], t["trip_id"], f"{t['trip_id']} is a trip to job "
                        f"{t['job_id']} but nobody clocked in there that day — check the date.", "low",
                        t["trip_id"], t["crew"] or t["driver_key"])
        # 6. Same vehicle, two drivers, same day.
        veh_day: dict = {}
        for t in trips:
            if t["vehicle_id"]:
                veh_day.setdefault((t["trip_date"], t["vehicle_id"]), set()).add(t["crew"] or t["driver_key"])
        for (d, vid), who in veh_day.items():
            if len(who) > 1:
                add("two_drivers", d, vid, f"{vid} has trips by {', '.join(sorted(who))} on {d} — pick the right "
                    "vehicle for each driver.", "medium", None, ", ".join(sorted(who)))
        # 7–9. Purpose, vehicle, commute classification.
        for t in trips:
            who = t["crew"] or t["driver_key"]
            if validate_purpose(t["purpose"]):
                fix = None
                if t["customer_name"] or (t["category"] and t["category"] not in ("Other", "Job")):
                    p = (f"{t['category'] or 'Service'} – {t['customer_name']}" if t["customer_name"]
                         else f"{t['category']} – {t['to_label']}")
                    if not validate_purpose(p):
                        fix = {"action": "edit", "trip_id": t["trip_id"], "changes": {"purpose": p},
                               "summary": f"set the purpose to \"{p}\""}
                add("vague_purpose", t["trip_date"], t["trip_id"], f"{t['trip_id']} purpose \"{t['purpose']}\" "
                    "isn't specific enough.", "high", t["trip_id"], who, fix)
            if not t["vehicle_id"]:
                vid = _resolve_vehicle(conn, t["trip_date"], t["driver_key"])
                fix = ({"action": "edit", "trip_id": t["trip_id"], "changes": {"vehicle_id": vid},
                        "summary": f"set the vehicle to {vid}"} if vid else None)
                add("no_vehicle", t["trip_date"], t["trip_id"], f"{t['trip_id']} has no vehicle.", "medium",
                    t["trip_id"], who, fix)
            labels = {(t["from_label"] or "").lower(), (t["to_label"] or "").lower()}
            home_co = labels == {"home", "company location"}
            if t["origin_mode"] == "Company Location" and home_co and t["classification"] == "business":
                add("commute", t["trip_date"], t["trip_id"], f"{t['trip_id']} home ↔ company location is a "
                    "commute (not deductible) under Company Location mode.", "high", t["trip_id"], who,
                    {"action": "edit", "trip_id": t["trip_id"], "changes": {"classification": "commute"},
                     "summary": "mark it as commute"})
            if t["classification"] == "commute" and (t["origin_mode"] or mode or "Jobs Only") == "Jobs Only" \
                    and t["origin"] in ("auto", "backfill"):
                add("commute", t["trip_date"], t["trip_id"], f"{t['trip_id']} is marked commute, but under Jobs "
                    "Only (home office) the drive from home to a job is business.", "medium", t["trip_id"], who,
                    {"action": "edit", "trip_id": t["trip_id"], "changes": {"classification": "business"},
                     "summary": "mark it as business"})
        # 10. Edited long after the trip; paid periods changed afterwards.
        ids = [t["trip_id"] for t in trips]
        if ids:
            qs = ",".join("?" * len(ids))
            tmap = {t["trip_id"]: t for t in trips}
            late_seen = set()
            for r in conn.execute(f"SELECT trip_id, MAX(edited_at) AS last FROM mileage_trip_edits WHERE trip_id "
                                  f"IN ({qs}) AND COALESCE(reason,'') <> 'Automatic rebuild' GROUP BY trip_id",
                                  ids).fetchall():
                t = tmap[r["trip_id"]]
                try:
                    days = (_dt.datetime.fromisoformat(str(r["last"]).replace("Z", "+00:00")).date()
                            - _dt.date.fromisoformat(t["trip_date"])).days
                except Exception:
                    continue
                if days > LATE_ENTRY_DAYS * 4 and r["trip_id"] not in late_seen:
                    late_seen.add(r["trip_id"])
                    add("late_edit", t["trip_date"], t["trip_id"], f"{t['trip_id']} was changed {days} days after "
                        "the trip — keep the receipt or calendar entry that supports the change.", "low",
                        t["trip_id"], t["crew"] or t["driver_key"])
            for r in conn.execute(f"SELECT t.trip_id, r.reimbursement_id, r.paid_at FROM mileage_trips t JOIN "
                                  f"mileage_reimbursements r ON r.reimbursement_id = t.reimbursement_id WHERE "
                                  f"r.status = 'paid' AND t.trip_id IN ({qs})", ids).fetchall():
                ch = conn.execute("SELECT COUNT(*) FROM mileage_trip_edits WHERE trip_id = ? AND edited_at > ? "
                                  "AND COALESCE(reason,'') <> 'Automatic rebuild'",
                                  (r["trip_id"], r["paid_at"])).fetchone()[0]
                if ch:
                    t = tmap[r["trip_id"]]
                    add("paid_changed", t["trip_date"], t["trip_id"], f"{t['trip_id']} was changed after "
                        f"{r['reimbursement_id']} was paid.", "high", t["trip_id"], t["crew"] or t["driver_key"])
        # 2. Logged miles vs odometer readings.
        if not restrict:
            for v in conn.execute("SELECT vehicle_id FROM vehicles").fetchall():
                rd = conn.execute("SELECT reading_date, odometer FROM odometer_readings WHERE vehicle_id = ? AND "
                                  "reading_date BETWEEN ? AND ? ORDER BY reading_date, id", (v["vehicle_id"], s, e)
                                  ).fetchall()
                if len(rd) < 2:
                    continue
                span = int(rd[-1]["odometer"]) - int(rd[0]["odometer"])
                logged = conn.execute("SELECT COALESCE(SUM(miles),0) FROM mileage_trips WHERE vehicle_id = ? AND "
                                      "status = 'active' AND trip_date > ? AND trip_date <= ?",
                                      (v["vehicle_id"], rd[0]["reading_date"], rd[-1]["reading_date"])).fetchone()[0]
                if span >= 0 and logged > span * 1.05 + 5:
                    add("odometer", rd[-1]["reading_date"], v["vehicle_id"], f"{v['vehicle_id']} has {logged:,.1f} "
                        f"mi of trips between {rd[0]['reading_date']} and {rd[-1]['reading_date']}, but the "
                        f"odometer moved only {span:,} mi.", "high", None, "")
        dismissed = set()
        if not include_dismissed:
            dismissed = {r[0] for r in conn.execute("SELECT finding_id FROM mileage_review_actions").fetchall()}
    finally:
        conn.close()
    order = {"high": 0, "medium": 1, "low": 2}
    out = [f for f in found if f["id"] not in dismissed]
    out.sort(key=lambda f: (order[f["severity"]], f["date"], f["id"]))
    return out


def format_review(findings: list, start: str, end: str) -> str:
    if not findings:
        return f"✅ Mileage review {start} → {end}: no problems found."
    L = [f"🤖 Mileage review {start} → {end}: {len(findings)} finding(s). Nothing has been changed.", ""]
    for f in findings:
        mark = {"high": "🔴", "medium": "🟠", "low": "⚪"}[f["severity"]]
        L.append(f"{mark} [{f['id']}] {f['title']}: {f['message']}")
        if f["fix"]:
            L.append(f"     Proposed fix: {f['fix']['summary']}.")
    L.append("")
    L.append("To fix one: apply_mileage_review_fix(finding_id, action=\"accept\"); to skip it for good: "
             "action=\"dismiss\". Ask the user before accepting.")
    return "\n".join(L)


def apply_review_fix(db_path: str, finding_id: str, action: str, actor: str, restrict: bool = False,
                     ukey: str = "", note: str = "") -> str:
    act = (action or "accept").strip().lower()
    parts = (finding_id or "").split(":")
    if len(parts) < 3:
        return "❌ Unknown finding id."
    d = parts[1]
    try:
        _parse_date(d)
    except ValueError:
        return "❌ Unknown finding id."
    ys = d[:4] + "-01-01" if parts[0] == "odometer" else d
    f = next((x for x in review_mileage(db_path, ys, d, restrict, ukey, include_dismissed=True)
              if x["id"] == finding_id), None)
    if f is None:
        return f"ℹ️ {finding_id} no longer applies — it may already be fixed. Run the review again."
    conn = get_connection(db_path)
    try:
        done = conn.execute("SELECT action FROM mileage_review_actions WHERE finding_id = ?", (finding_id,)).fetchone()
    finally:
        conn.close()
    if act == "dismiss":
        res = f"✅ Dismissed {finding_id} — it won't be raised again."
    elif act == "accept":
        if done and done["action"] == "accepted":
            return f"ℹ️ {finding_id} was already accepted."
        fix = f["fix"]
        if not fix:
            return f"ℹ️ No automatic fix for this one: {f['message']} Fix it by hand, or dismiss it."
        reason = f"AI review: {f['title']}" + (f" — {note.strip()}" if (note or "").strip() else "")
        if fix["action"] == "void":
            res = void_trip(db_path, fix["trip_id"], reason, actor, restrict, ukey)
        else:
            res = update_trip(db_path, fix["trip_id"], fix["changes"], reason, actor, restrict, ukey)
        if not res.startswith("✅"):
            return res
    else:
        return "❌ action must be accept or dismiss."
    conn = get_connection(db_path)
    try:
        conn.execute("INSERT INTO mileage_review_actions (finding_id, action, acted_by, acted_at, note) "
                     "VALUES (?,?,?,?,?) ON CONFLICT(finding_id) DO UPDATE SET action = excluded.action, "
                     "acted_by = excluded.acted_by, acted_at = excluded.acted_at, note = excluded.note",
                     (finding_id, "accepted" if act == "accept" else "dismissed", actor, utcnow_iso(),
                      note or None))
    finally:
        conn.close()
    return res

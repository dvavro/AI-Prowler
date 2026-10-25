"""Mileage tracker Phases 4–5 (spec test plan MI-28..MI-45, plus exports,
PDFs, reminders and the new MCP tools).

Like test_mileage_ops.py: every test builds its own scratch SQLite database
with the real db_schema.apply_schema under tmp_path — never the live Jobs app
database — and the network hooks (OSRM, geocoder, home address) are fakes.

Run: py -m pytest tests\\mcp_tests\\test_mileage_phase45.py -v
"""
import base64
import datetime as dt
import io
import json
import re
import sys
import zlib
from pathlib import Path
from unittest.mock import MagicMock

import pytest

SRC_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(SRC_ROOT))

import db_mileage_ops as mo  # noqa: E402
from db_access import get_connection  # noqa: E402
from db_schema import apply_schema  # noqa: E402

PLACES = {"HOME ADDR": (0.0, 0.0), "Bank": (0.0, 0.3), "Home Depot": (0.5, 0.0)}
VIN_OK = "1HGCM82633A004352"
VIN_OK2 = "1M8GDM9AXKP042788"


def fake_router(a, b):
    if not a or not b:
        return None
    return round((abs(a[0] - b[0]) + abs(a[1] - b[1])) * 10, 1)


@pytest.fixture
def db(tmp_path, monkeypatch):
    path = str(tmp_path / "jobs.db")
    conn = get_connection(path)
    apply_schema(conn)
    for uid in ("u_owner", "u_sam", "u_pat"):
        conn.execute("INSERT OR IGNORE INTO users (id) VALUES (?)", (uid,))
    conn.close()
    monkeypatch.setattr(mo, "ROUTER", fake_router)
    monkeypatch.setattr(mo, "GEOCODER", lambda addr: PLACES.get(addr))
    monkeypatch.setattr(mo, "HOME_RESOLVER", lambda crew: "HOME ADDR")
    mo._ROUTE_CACHE.clear()
    return path


def _sql(path, q, args=()):
    conn = get_connection(path)
    try:
        return conn.execute(q, args).fetchall()
    finally:
        conn.close()


def _job(path, job_id, lat, lon, customer, status="Scheduled", crew="", crew_uid=None):
    _sql(path, "INSERT INTO jobs (job_id, customer_name, latitude, longitude, service_type, job_status, crew, "
               "crew_user_id) VALUES (?,?,?,?,?,?,?,?)", (job_id, customer, lat, lon, "Window", status, crew, crew_uid))


def _entry(path, eid, job_id, day, t_in, t_out, crew="Sam", uid=None, gps_in="", gps_out=""):
    _sql(path, "INSERT INTO time_entries (entry_id, job_id, customer_name, entry_date, clock_in, clock_out, crew, "
               "crew_user_id, clock_in_gps, clock_out_gps) VALUES (?,?,?,?,?,?,?,?,?,?)",
         (eid, job_id, job_id, day, f"{day} {t_in}:00", f"{day} {t_out}:00" if t_out else None, crew, uid,
          gps_in, gps_out))


def _three_job_day(path, day="2026-10-20"):
    _job(path, "JOB-0001", 1.0, 0.0, "Norwoods")
    _job(path, "JOB-0002", 1.0, 1.0, "Bamboo")
    _job(path, "JOB-0003", 0.0, 2.0, "Copperline")
    _entry(path, "TE-0001", "JOB-0001", day, "08:00", "09:00")
    _entry(path, "TE-0002", "JOB-0002", day, "09:30", "10:30")
    _entry(path, "TE-0003", "JOB-0003", day, "11:00", "12:00")
    return day


def _active(path, day=None):
    q = "SELECT * FROM mileage_trips WHERE status = 'active'"
    return _sql(path, q + (" AND trip_date = ?" if day else "") + " ORDER BY seq, trip_id", (day,) if day else ())


def _legs(path, day):
    return [(r["from_label"], r["to_label"]) for r in _active(path, day)]


def _pdf_text(b: bytes) -> bytes:
    assert b.startswith(b"%PDF-1.4") and b.rstrip().endswith(b"%%EOF")
    return b"\n".join(zlib.decompress(m) for m in re.findall(rb"stream\n(.*?)\nendstream", b, re.S))


def _crew_car(path, who="Sam", uid="u_sam", vin=VIN_OK):
    r = mo.add_vehicle(path, vin, "tester", ownership="crew", owner_user_id=uid, owner_name=who, odometer=1000)
    assert r.startswith("✅"), r
    return re.search(r"VEH-\d+", r).group(0)


def _crew_trip(path, vid, day="2026-10-05", miles=10, to="Bank", uid="u_sam", who="Sam", parking=None):
    r = mo.add_trip(path, "Sam", trip_date=day, to_label=to, purpose=f"Supplies – {to} run", miles=miles,
                    crew_user_id=uid, crew=who, vehicle_id=vid, parking_usd=parking)
    assert r.startswith("✅"), r
    return re.search(r"TRIP-\d+", r).group(0)


# ── FR-21: miles follow schedule changes ───────────────────────────────────

def test_mi28_job_cancelled_before_day_reconnects_legs(db):
    day = _three_job_day(db)
    mo.build_trips_for_day(db, day, "tester")
    _sql(db, "DELETE FROM time_entries WHERE entry_id = 'TE-0002'")
    _sql(db, "UPDATE jobs SET job_status = 'Cancelled' WHERE job_id = 'JOB-0002'")
    done = mo.refresh_stale_days(db, day, day)
    assert done == [(day, "owner")]
    assert _legs(db, day) == [("Home", "Norwoods"), ("Norwoods", "Copperline"), ("Copperline", "Home")]
    assert not any("Bamboo" in (f + t) for f, t in _legs(db, day))


def test_mi29_cancelled_on_arrival_keeps_leg_as_business(db):
    _job(db, "JOB-0009", 1.0, 0.0, "Smiths")
    r = mo.cancelled_on_arrival(db, "JOB-0009", "tester", gps="1.0,0.0", note="not home")
    assert r.startswith("✅"), r
    job = _sql(db, "SELECT job_status, service_details FROM jobs WHERE job_id = 'JOB-0009'")[0]
    assert job["job_status"] == "Cancelled" and "Cancelled on arrival" in job["service_details"]
    today = dt.date.today().isoformat()
    legs = _active(db, today)
    there = [t for t in legs if t["to_label"] == "Smiths"][0]
    assert there["classification"] == "business" and there["purpose"] == "Smiths – cancelled on site"


def test_mi30_job_moved_after_build_rebuilds_old_day(db):
    day = _three_job_day(db)
    mo.build_trips_for_day(db, day, "tester")
    _sql(db, "DELETE FROM time_entries WHERE entry_id = 'TE-0003'")        # moved to another day
    mo.refresh_stale_days(db, day, day)
    assert _legs(db, day) == [("Home", "Norwoods"), ("Norwoods", "Bamboo"), ("Bamboo", "Home")]
    void = _sql(db, "SELECT void_reason FROM mileage_trips WHERE status = 'void'")
    assert void and all("Rebuilt" in v["void_reason"] for v in void)


def test_driver_with_no_clock_ins_left_gets_trips_voided(db):
    _job(db, "JOB-0001", 1.0, 0.0, "Norwoods")
    _entry(db, "TE-0001", "JOB-0001", "2026-10-21", "08:00", "09:00")
    mo.build_trips_for_day(db, "2026-10-21", "tester")
    assert _active(db, "2026-10-21")
    _sql(db, "DELETE FROM time_entries")                                     # job deleted
    mo.refresh_stale_days(db, "2026-10-21", "2026-10-21")
    assert _active(db, "2026-10-21") == []


def test_mi31_route_reordered_after_driving_clock_order_wins(db):
    day = _three_job_day(db)
    mo.build_trips_for_day(db, day, "tester")
    before = _legs(db, day)
    for n, j in enumerate(("JOB-0003", "JOB-0002", "JOB-0001"), start=1):
        _sql(db, "INSERT INTO route_stops (route_date, crew_id, stop_number, job_id, leg_drive_miles) "
                 "VALUES (?,?,?,?,?)", (day, "Sam", n, j, 7.0))
    assert mo.refresh_stale_days(db, day, day) == [(day, "owner")]          # route change noticed
    assert _legs(db, day) == before


def test_mi32_rebuild_keeps_hand_edited_trip(db):
    day = _three_job_day(db)
    mo.build_trips_for_day(db, day, "tester")
    assert mo.update_trip(db, "TRIP-0002", {"miles": 99}, "detour for fuel", "Sam").startswith("✅")
    _sql(db, "UPDATE time_entries SET clock_in_gps = '0.0,2.5' WHERE entry_id = 'TE-0003'")
    mo.refresh_stale_days(db, day, day)
    rows = {r["trip_id"]: r for r in _active(db, day)}
    assert rows["TRIP-0002"]["miles"] == 99 and rows["TRIP-0002"]["manually_edited"] == 1
    assert rows["TRIP-0003"]["miles"] == 25.0                                 # Bamboo (1,1) -> GPS (0,2.5)


def test_unchanged_day_is_not_rebuilt_and_old_days_get_a_baseline(db):
    day = _three_job_day(db)
    mo.build_trips_for_day(db, day, "tester")
    assert mo.refresh_stale_days(db, day, day) == []
    _sql(db, "DELETE FROM mileage_day_builds")                               # built before Phase 5
    _sql(db, "UPDATE time_entries SET clock_out = '2026-10-20 12:30:00' WHERE entry_id = 'TE-0003'")
    assert mo.refresh_stale_days(db, day, day) == []                         # adopts today's state
    assert len(_sql(db, "SELECT * FROM mileage_day_builds")) == 1


def test_locked_day_is_not_rebuilt(db):
    day = _three_job_day(db)
    mo.build_trips_for_day(db, day, "tester")
    mo.lock_period(db, "2026-10", "tester")
    _sql(db, "DELETE FROM time_entries WHERE entry_id = 'TE-0002'")
    assert mo.refresh_stale_days(db, day, day) == []
    assert len(_active(db, day)) == 4


# ── FR-23: crew reimbursement ──────────────────────────────────────────────

def test_period_bounds_monthly_weekly_biweekly(db):
    assert mo.period_bounds(db, "2026-10-24") == ("2026-10-01", "2026-10-31", "monthly")
    assert mo.set_reimb_period(db, "weekly", "tester").startswith("✅")
    assert mo.period_bounds(db, "2026-10-24") == ("2026-10-19", "2026-10-25", "weekly")
    mo.set_reimb_period(db, "every other week", "tester")
    assert mo.period_bounds(db, "2026-10-24") == ("2026-10-12", "2026-10-25", "biweekly")
    assert mo.set_reimb_period(db, "daily", "tester").startswith("❌")


def test_mi35_submit_with_needs_fix_day_refused(db):
    vid = _crew_car(db)
    _crew_trip(db, vid)
    _job(db, "JOB-0001", 1.0, 0.0, "Norwoods")
    _entry(db, "TE-0001", "JOB-0001", "2026-10-06", "08:00", None, crew="Sam", uid="u_sam")
    r = mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-10-05")
    assert r.startswith("❌") and "2026-10-06" in r
    assert _sql(db, "SELECT * FROM mileage_reimbursements") == []


def test_mi36_submit_return_resubmit_approve_pay(db):
    vid = _crew_car(db)
    tid = _crew_trip(db, vid, parking=2.5)
    r = mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-10-05")
    assert r.startswith("✅ REIM-0001 submitted") and "$7.60" in r, r
    assert _sql(db, "SELECT reimbursement_id FROM mileage_trips WHERE trip_id = ?", (tid,))[0][0] == "REIM-0001"
    # Frozen while submitted.
    assert mo.update_trip(db, tid, {"miles": 12}, "x", "Sam").startswith("🔒")
    # Returned with a note -> editable again.
    assert mo.return_reimbursement(db, "REIM-0001", "", "Dave", "u_owner").startswith("❌")
    assert mo.return_reimbursement(db, "REIM-0001", "Bank run was 12 mi", "Dave", "u_owner").startswith("↩️")
    assert mo.update_trip(db, tid, {"miles": 12}, "owner asked", "Sam").startswith("✅")
    assert "REIM-0001 submitted" in mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-10-05")
    r = mo.approve_reimbursement(db, "REIM-0001", "Dave", "u_owner")
    assert r.startswith("✅") and "$11.62" in r, r                           # 12 × 76¢ + 2.50
    # Amount frozen at approval.
    row = _sql(db, "SELECT * FROM mileage_reimbursements")[0]
    assert row["status"] == "approved" and row["amount_usd"] == 9.12 and row["parking_tolls_usd"] == 2.5
    assert mo.mark_reimbursement_paid(db, "REIM-0001", "Dave", "u_owner", method="bitcoin").startswith("❌")
    r = mo.mark_reimbursement_paid(db, "REIM-0001", "Dave", "u_owner", "2026-11-03", "check", "1042")
    assert r.startswith("✅") and "check" in r
    ev = [(e["from_status"], e["to_status"], e["actor"]) for e in
          _sql(db, "SELECT * FROM mileage_reimbursement_events ORDER BY id")]
    assert ev == [("draft", "submitted", "Sam"), ("submitted", "returned", "Dave"),
                  ("returned", "submitted", "Sam"), ("submitted", "approved", "Dave"),
                  ("approved", "paid", "Dave")]
    assert all(e["at"] for e in _sql(db, "SELECT at FROM mileage_reimbursement_events"))


def test_mi37_paid_period_trips_cannot_change(db):
    vid = _crew_car(db)
    tid = _crew_trip(db, vid)
    mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-10-05")
    mo.approve_reimbursement(db, "REIM-0001", "Dave", "u_owner")
    mo.mark_reimbursement_paid(db, "REIM-0001", "Dave", "u_owner", method="cash")
    assert mo.update_trip(db, tid, {"miles": 50}, "x", "Sam").startswith("🔒")
    assert mo.void_trip(db, tid, "x", "Sam").startswith("🔒")
    r = mo.add_trip(db, "Sam", trip_date="2026-10-09", to_label="Bank", purpose="Bank – deposit", miles=5,
                    crew_user_id="u_sam", crew="Sam", vehicle_id=vid)
    assert r.startswith("🔒")                                                # period already paid
    # A change that slipped through (direct SQL) is flagged by the AI review.
    _sql(db, "INSERT INTO mileage_trip_edits (trip_id, field, old_value, new_value, edited_by, edited_at, reason) "
             "VALUES (?, 'miles', '10', '50', 'x', '2099-01-01T00:00:00Z', 'sneaky')", (tid,))
    f = mo.review_mileage(db, "2026-10-01", "2026-10-31")
    assert any(x["check"] == "paid_changed" for x in f)


def test_mi38_crew_cannot_approve_own(db):
    vid = _crew_car(db)
    _crew_trip(db, vid)
    mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-10-05")
    assert "own" in mo.approve_reimbursement(db, "REIM-0001", "Sam", "u_sam")
    assert "own" in mo.mark_reimbursement_paid(db, "REIM-0001", "Sam", "u_sam", method="cash")


def test_cannot_resubmit_or_pay_out_of_order(db):
    vid = _crew_car(db)
    _crew_trip(db, vid)
    mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-10-05")
    assert "already submitted" in mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-10-20")
    assert "approve it before" in mo.mark_reimbursement_paid(db, "REIM-0001", "Dave", "u_owner", method="cash")
    assert mo.submit_reimbursement(db, "Pat", "u_pat", "u_pat", "Pat", "2026-10-05").startswith("ℹ️")


def test_business_vehicle_trips_never_reimbursed(db):
    r = mo.add_vehicle(db, VIN_OK2, "tester", ownership="business", odometer=500)
    vid = re.search(r"VEH-\d+", r).group(0)
    mo.add_trip(db, "Sam", trip_date="2026-10-05", to_label="Bank", purpose="Bank – deposit", miles=5,
                crew_user_id="u_sam", crew="Sam", vehicle_id=vid)
    assert mo.reimbursement_preview(db, "u_sam", "2026-10-05")["trip_count"] == 0


def test_report_shows_owed_approved_and_paid(db):
    vid = _crew_car(db)
    _crew_trip(db, vid, day="2026-09-05")
    _crew_trip(db, vid, day="2026-10-05")
    mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-09-05")
    mo.approve_reimbursement(db, "REIM-0001", "Dave", "u_owner")
    rep = mo.mileage_report(db, "2026-09-01", "2026-10-31")
    st = rep["reimbursement_status"]["Sam"]
    assert st == {"not_submitted": 7.6, "submitted": 0.0, "approved": 7.6, "paid": 0.0}
    txt = mo.format_report(rep)
    assert "approved unpaid $7.60" in txt and "not submitted $7.60" in txt


# ── FR-17 / FR-23: exports and PDFs ────────────────────────────────────────

def _two_rate_year(db):
    r = mo.add_vehicle(db, VIN_OK2, "tester", ownership="business", nickname="White F-150", odometer=500,
                       placed_in_service="2026-01-01")
    vid = re.search(r"VEH-\d+", r).group(0)
    for d in ("2026-06-15", "2026-07-15"):
        assert mo.add_trip(db, "Dave", trip_date=d, to_label="Home Depot", purpose="Supplies – ladders",
                           miles=10, vehicle_id=vid).startswith("✅")
    return vid


def test_mi39_statement_pdf_matches_the_period(db):
    vid = _crew_car(db)
    _crew_trip(db, vid, miles=10)
    _crew_trip(db, vid, day="2026-10-07", miles=5, to="Home Depot")
    mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-10-05")
    mo.approve_reimbursement(db, "REIM-0001", "Dave", "u_owner")
    mo.mark_reimbursement_paid(db, "REIM-0001", "Dave", "u_owner", "2026-11-03", "check", "1042")
    name, pdf = mo.statement_pdf(db, "REIM-0001")
    assert name == "Mileage_reimbursement_REIM-0001.pdf"
    txt = _pdf_text(pdf)
    for want in (b"REIM-0001", b"Paid", b"check", b"1042", b"15.0", b"$11.40", b"Sam"):
        assert want in txt, want
    # Draft statement for a period not submitted yet.
    name, pdf = mo.statement_pdf(db, "", "u_sam", "2026-12-05")
    assert name.endswith("_draft.pdf") and b"not submitted yet" in _pdf_text(pdf)


def test_mi40_pdf_log_shows_both_rate_periods_and_vehicle_totals(db):
    _two_rate_year(db)
    name, data, mime = mo.export_mileage(db, "2026-01-01", "2026-12-31", "pdf")
    assert name == "Mileage_2026-01-01_to_2026-12-31.pdf" and mime == "application/pdf"
    txt = _pdf_text(data)
    assert b"72.5\xa2" in txt and b"76\xa2" in txt                          # cp1252 cent sign
    assert b"$7.25" in txt and b"$7.60" in txt and b"$14.85" in txt
    assert VIN_OK2.encode() in txt and b"Form 4562" in txt


def test_xlsx_export_has_four_sheets_matching_the_report(db):
    from openpyxl import load_workbook
    _two_rate_year(db)
    name, data, _ = mo.export_mileage(db, "2026-01-01", "2026-12-31", "excel")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Trips", "Summary", "Reimbursement", "Vehicles"]
    trips = list(wb["Trips"].values)
    assert trips[0][0] == "Trip ID" and len(trips) == 3
    summary = [r for r in wb["Summary"].values if r and r[0] == "Total deduction"][0]
    assert summary[2] == 20.0 and summary[3] == 14.85
    veh = list(wb["Vehicles"].values)
    assert veh[1][3] == VIN_OK2 and veh[1][5] == 20.0


def test_csv_export(db):
    _two_rate_year(db)
    name, data, mime = mo.export_mileage(db, "2026-01-01", "2026-12-31", "csv")
    text = data.decode("utf-8-sig").splitlines()
    assert mime == "text/csv" and text[0].startswith("Trip ID,Date") and len(text) == 3
    with pytest.raises(ValueError):
        mo.export_mileage(db, "2026-01-01", "2026-12-31", "docx")


# ── FR-24: reminders ───────────────────────────────────────────────────────

def test_mi41_reminder_turned_off_is_not_sent(db):
    _job(db, "JOB-0001", 1.0, 0.0, "Norwoods")
    _entry(db, "TE-0001", "JOB-0001", "2026-10-20", "08:00", None, uid="u_sam")
    now = dt.datetime(2026, 10, 21, 19, 30)
    assert any(r["type"] == "gaps" for r in mo.due_reminders(db, "u_sam", False, now))
    assert mo.set_reminder(db, "gaps", False, "Dave").startswith("✅")
    assert not any(r["type"] == "gaps" for r in mo.due_reminders(db, "u_sam", False, now))
    assert mo.set_reminder(db, "nonsense", False, "Dave").startswith("❌")


def test_mi42_gap_reminder_goes_to_that_driver_only(db):
    _job(db, "JOB-0001", 1.0, 0.0, "Norwoods")
    _entry(db, "TE-0001", "JOB-0001", "2026-10-20", "08:00", None, uid="u_sam")
    now = dt.datetime(2026, 10, 21, 19, 30)
    sam = [r for r in mo.due_reminders(db, "u_sam", False, now) if r["type"] == "gaps"]
    pat = [r for r in mo.due_reminders(db, "u_pat", False, now) if r["type"] == "gaps"]
    assert len(sam) == 1 and "2026-10-20" in sam[0]["text"] and pat == []


def test_pay_period_reminders_for_crew_and_owner(db):
    vid = _crew_car(db)
    _crew_trip(db, vid)
    nov2 = dt.datetime(2026, 11, 2, 9, 0)
    crew = [r for r in mo.due_reminders(db, "u_sam", False, nov2) if r["type"] == "pay_period"]
    assert crew and "2026-10-01" in crew[0]["text"]
    mo.submit_reimbursement(db, "Sam", "u_sam", "u_sam", "Sam", "2026-10-05")
    assert not [r for r in mo.due_reminders(db, "u_sam", False, nov2) if r["type"] == "pay_period"]
    owner = [r["text"] for r in mo.due_reminders(db, "u_owner", True, nov2)]
    assert any("waiting for approval" in t for t in owner)
    mo.approve_reimbursement(db, "REIM-0001", "Dave", "u_owner")
    assert any("approved, not yet paid" in r["text"] for r in mo.due_reminders(db, "u_owner", True, nov2))


def test_weekly_and_odometer_reminders(db):
    mo.add_vehicle(db, VIN_OK2, "tester", ownership="business")              # no odometer
    mon = dt.datetime(2026, 10, 26, 8, 0)
    types = [r["type"] for r in mo.due_reminders(db, "owner", True, mon)]
    assert "weekly" in types and "odometer" in types
    tue = dt.datetime(2026, 10, 27, 8, 0)
    assert "weekly" not in [r["type"] for r in mo.due_reminders(db, "owner", True, tue)]
    mo.add_odometer(db, "VEH-0001", 5000, "tester", "2026-10-01")
    dec = dt.datetime(2026, 12, 27, 8, 0)
    assert any("Year end" in r["text"] for r in mo.due_reminders(db, "owner", True, dec))
    mo.add_odometer(db, "VEH-0001", 6000, "tester", "2026-12-27", "year_end")
    assert not any("Year end" in r["text"] for r in mo.due_reminders(db, "owner", True, dec))


def test_monday_brief_has_ai_review_summary(db):
    for _ in range(2):
        mo.add_trip(db, "Dave", trip_date="2026-10-20", to_label="Bank", purpose="Bank – deposit", miles=6)
    lines = mo.brief_lines(db, dt.datetime(2026, 10, 26, 7, 0))
    assert any(l.startswith("AI mileage review of last week:") and "duplicate" in l for l in lines)
    mo.set_tracking_enabled(db, False, "Dave")
    assert mo.brief_lines(db, dt.datetime(2026, 10, 26, 7, 0)) == []


# ── FR-25: AI review ───────────────────────────────────────────────────────

def _route_day(db, day="2026-10-20"):
    _job(db, "JOB-0001", 1.0, 0.0, "Norwoods")
    _job(db, "JOB-0002", 1.0, 1.0, "Bamboo")
    _entry(db, "TE-0001", "JOB-0001", day, "08:00", "09:00")
    _entry(db, "TE-0002", "JOB-0002", day, "09:30", "10:30")
    mo.build_trips_for_day(db, day, "tester")
    for n, (j, m) in enumerate((("JOB-0001", 6.0), ("JOB-0002", 6.0)), start=1):
        _sql(db, "INSERT INTO route_stops (route_date, crew_id, stop_number, job_id, leg_drive_miles) "
                 "VALUES (?,?,?,?,?)", (day, "Sam", n, j, m))
    return day


def test_mi43_route_mismatch_found_nothing_changed(db):
    day = _route_day(db)                         # logged 10 + 10 to jobs, planned 6 + 6 (+67 %)
    before = [(r["trip_id"], r["miles"]) for r in _active(db, day)]
    f = [x for x in mo.review_mileage(db, day, day) if x["check"] == "route_mismatch"]
    assert len(f) == 1 and f[0]["fix"]["changes"] == {"miles": 6.0}
    assert "+67%" in f[0]["message"]
    assert [(r["trip_id"], r["miles"]) for r in _active(db, day)] == before
    assert "route_mismatch" in mo.format_review(f, day, day) and "Nothing has been changed" in \
        mo.format_review(f, day, day)


def test_mi44_accept_applies_fix_through_audited_edit(db):
    day = _route_day(db)
    f = [x for x in mo.review_mileage(db, day, day) if x["check"] == "route_mismatch"][0]
    r = mo.apply_review_fix(db, f["id"], "accept", "Dave")
    assert r.startswith("✅"), r
    t = _sql(db, "SELECT * FROM mileage_trips WHERE trip_id = ?", (f["fix"]["trip_id"],))[0]
    assert t["miles"] == 6.0 and t["manually_edited"] == 1
    ed = _sql(db, "SELECT reason FROM mileage_trip_edits WHERE trip_id = ?", (t["trip_id"],))
    assert ed and ed[-1]["reason"].startswith("AI review:")
    assert _sql(db, "SELECT action FROM mileage_review_actions")[0]["action"] == "accepted"


def test_mi45_dismissed_finding_not_raised_again(db):
    for _ in range(2):
        mo.add_trip(db, "Dave", trip_date="2026-10-20", to_label="Bank", purpose="Bank – deposit", miles=6)
    dup = [x for x in mo.review_mileage(db, "2026-10-20", "2026-10-20") if x["check"] == "duplicate"]
    assert len(dup) == 1 and dup[0]["fix"]["action"] == "void"
    assert mo.apply_review_fix(db, dup[0]["id"], "dismiss", "Dave").startswith("✅ Dismissed")
    assert not [x for x in mo.review_mileage(db, "2026-10-20", "2026-10-20") if x["check"] == "duplicate"]
    assert [x for x in mo.review_mileage(db, "2026-10-20", "2026-10-20", include_dismissed=True)
            if x["check"] == "duplicate"]


def test_accepting_duplicate_voids_the_copy(db):
    for _ in range(2):
        mo.add_trip(db, "Dave", trip_date="2026-10-20", to_label="Bank", purpose="Bank – deposit", miles=6)
    dup = [x for x in mo.review_mileage(db, "2026-10-20", "2026-10-20") if x["check"] == "duplicate"][0]
    assert mo.apply_review_fix(db, dup["id"], "accept", "Dave").startswith("✅")
    assert len(_active(db, "2026-10-20")) == 1
    assert "no longer applies" in mo.apply_review_fix(db, dup["id"], "accept", "Dave")


def test_review_other_checks(db):
    r = mo.add_vehicle(db, VIN_OK2, "tester", ownership="business", odometer=1000,
                       placed_in_service="2026-10-01")
    vid = re.search(r"VEH-\d+", r).group(0)
    mo.add_odometer(db, vid, 1010, "tester", "2026-10-31")
    # 40 mi to a place 3 mi away (manual, GPS given) -> long leg with a road-distance fix
    mo.add_trip(db, "Dave", trip_date="2026-10-21", to_label="Home Depot", purpose="Supplies – caulk",
                miles=40, from_gps="0.0,0.0", to_gps="0.0,0.05", vehicle_id=vid)
    # No vehicle anywhere -> no_vehicle has no automatic fix when it can't be resolved
    _sql(db, "UPDATE mileage_trips SET vehicle_id = NULL, purpose = 'work' WHERE trip_id = 'TRIP-0001'")
    f = {x["check"]: x for x in mo.review_mileage(db, "2026-10-01", "2026-10-31")}
    assert f["long_leg"]["fix"]["changes"] == {"miles": 0.5}
    assert "vague_purpose" in f and "no_vehicle" in f
    assert f["no_vehicle"]["fix"]["changes"] == {"vehicle_id": vid}       # the only vehicle
    _sql(db, "UPDATE mileage_trips SET vehicle_id = ? WHERE trip_id = 'TRIP-0001'", (vid,))
    f = {x["check"]: x for x in mo.review_mileage(db, "2026-10-01", "2026-10-31")}
    assert "odometer" in f and "moved only 10" in f["odometer"]["message"]


def test_review_flags_commute_misclassified(db):
    day = _three_job_day(db)
    mo.build_trips_for_day(db, day, "tester")
    _sql(db, "UPDATE mileage_trips SET classification = 'commute' WHERE trip_id = 'TRIP-0001'")
    f = [x for x in mo.review_mileage(db, day, day) if x["check"] == "commute"]
    assert f and f[0]["fix"]["changes"] == {"classification": "business"}


# ── MCP tools (both modes) ─────────────────────────────────────────────────

NEW_TOOLS = ["export_mileage_log", "mileage_reimbursement", "review_mileage_log",
             "apply_mileage_review_fix", "job_cancelled_on_arrival"]
OWNER = {"id": "u_owner", "name": "Dave Owner", "role": "owner", "status": "active", "scopes": []}
CREW_A = {"id": "u_sam", "name": "Sam", "role": "field_crew", "status": "active", "scopes": []}
CREW_B = {"id": "u_pat", "name": "Pat", "role": "field_crew", "status": "active", "scopes": []}


@pytest.fixture(scope="module")
def mcp_mod():
    import ai_prowler_mcp as ap
    ap._prewarm_event.set()
    return ap


@pytest.fixture
def env(db, monkeypatch, mcp_mod):
    monkeypatch.setattr(mcp_mod, "_resolve_job_db_path", lambda ctx, fp="": db)
    return db


def _as(monkeypatch, mcp_mod, user):
    monkeypatch.setattr(mcp_mod, "_current_user", lambda ctx: user)
    if user is None:
        return None
    c = MagicMock()
    c.request_context.request.state.user = user
    return c


def test_new_tools_registered_everywhere(mcp_mod):
    from mcp_tool_catalog import JOBS_APP_TOOLS, TOOL_CATALOG
    src = Path(mcp_mod.__file__).read_text(encoding="utf-8")
    srv = src[src.index("_srv_pa_allowed = {"):]
    srv = srv[:srv.index("}")]
    per = src[src.index("_allowed_tools = {"):]
    per = per[:per.index("}")]
    for t in NEW_TOOLS:
        assert callable(getattr(mcp_mod, t)), t
        assert t in JOBS_APP_TOOLS and t in TOOL_CATALOG, t
        assert f'"{t}"' in srv and f'"{t}"' in per, t


ALL_MILEAGE_TOOLS = ["get_mileage_log", "build_mileage_log", "add_mileage_trip", "update_mileage_trip",
                     "fix_mileage_gap", "manage_vehicles", "get_mileage_report", "mileage_settings"] + NEW_TOOLS


def test_how_to_use_and_user_guide_document_every_mileage_tool(monkeypatch, mcp_mod):
    monkeypatch.setattr(mcp_mod, "_current_user", lambda ctx: None)
    body = mcp_mod.how_to_use_ai_prowler(None)
    assert "Mileage tracker" in body and "mileage by voice or chat" in body
    guide = (SRC_ROOT / "COMPLETE_USER_GUIDE.md").read_text(encoding="utf-8")
    for t in ALL_MILEAGE_TOOLS:
        assert t in body, f"how_to_use_ai_prowler is missing {t}"
        assert f"| `{t}` |" in guide, f"COMPLETE_USER_GUIDE.md tool table is missing {t}"
        assert (getattr(mcp_mod, t).__doc__ or "").strip(), f"{t} has no description"


def test_tool_reimbursement_flow_and_role_gates(env, monkeypatch, mcp_mod):
    vid = _crew_car(env)
    _crew_trip(env, vid)
    cs = _as(monkeypatch, mcp_mod, CREW_A)
    st = json.loads(mcp_mod.mileage_reimbursement(action="status", period_date="2026-10-05", as_json=True, ctx=cs))
    assert st["preview"]["trip_count"] == 1 and st["queue"] == [] and st["is_admin"] is False
    assert "REIM-0001 submitted" in mcp_mod.mileage_reimbursement(action="submit", period_date="2026-10-05", ctx=cs)
    assert "Only the owner" in mcp_mod.mileage_reimbursement(action="approve", reimbursement_id="REIM-0001", ctx=cs)
    cp = _as(monkeypatch, mcp_mod, CREW_B)
    assert "own reimbursement statements" in mcp_mod.mileage_reimbursement(
        action="statement", reimbursement_id="REIM-0001", ctx=cp)
    co = _as(monkeypatch, mcp_mod, OWNER)
    q = json.loads(mcp_mod.mileage_reimbursement(action="status", as_json=True, ctx=co))["queue"]
    assert [r["reimbursement_id"] for r in q] == ["REIM-0001"]
    assert mcp_mod.mileage_reimbursement(action="approve", reimbursement_id="reim-0001", ctx=co).startswith("✅")
    assert mcp_mod.mileage_reimbursement(action="paid", reimbursement_id="REIM-0001", method="payroll",
                                         ctx=co).startswith("✅")
    dl = json.loads(mcp_mod.mileage_reimbursement(action="statement", reimbursement_id="REIM-0001",
                                                  deliver="download", ctx=cs))
    assert dl["filename"].endswith(".pdf") and base64.b64decode(dl["base64"]).startswith(b"%PDF")


def test_tool_export_download_and_file(env, monkeypatch, mcp_mod, tmp_path):
    _two_rate_year(env)
    co = _as(monkeypatch, mcp_mod, None)
    dl = json.loads(mcp_mod.export_mileage_log(start_date="2026-01-01", end_date="2026-12-31", format="xlsx",
                                               deliver="download", ctx=co))
    assert dl["filename"].endswith(".xlsx") and base64.b64decode(dl["base64"])[:2] == b"PK"
    r = mcp_mod.export_mileage_log(start_date="2026-01-01", end_date="2026-12-31", format="pdf", ctx=co)
    assert r.startswith("✅ Saved") and "Mileage_Exports" in r
    assert (Path(env).parent / "Mileage_Exports" / "Mileage_2026-01-01_to_2026-12-31.pdf").exists()
    assert "❌" in mcp_mod.export_mileage_log(format="docx", deliver="download", ctx=co)
    assert "email_to" in mcp_mod.export_mileage_log(format="csv", deliver="email", ctx=co)


def test_tool_crew_export_only_own_trips(env, monkeypatch, mcp_mod):
    vid = _crew_car(env)
    _crew_trip(env, vid)
    vid2 = _crew_car(env, "Pat", "u_pat", VIN_OK2)
    _crew_trip(env, vid2, uid="u_pat", who="Pat", to="Home Depot")
    cs = _as(monkeypatch, mcp_mod, CREW_A)
    dl = json.loads(mcp_mod.export_mileage_log(start_date="2026-10-01", end_date="2026-10-31", format="csv",
                                               crew="Pat", deliver="download", ctx=cs))
    rows = base64.b64decode(dl["base64"]).decode("utf-8-sig").splitlines()
    assert len(rows) == 2 and "Sam" in rows[1] and "Pat" not in rows[1]


def test_tool_review_and_apply(env, monkeypatch, mcp_mod):
    co = _as(monkeypatch, mcp_mod, None)
    for _ in range(2):
        mcp_mod.add_mileage_trip(to_place="Bank", purpose="Bank – deposit", miles="6", date="2026-10-20", ctx=co)
    j = json.loads(mcp_mod.review_mileage_log(start_date="2026-10-20", end_date="2026-10-20", as_json=True, ctx=co))
    dup = [f for f in j["findings"] if f["check"] == "duplicate"][0]
    assert "Proposed fix" in mcp_mod.review_mileage_log(start_date="2026-10-20", end_date="2026-10-20", ctx=co)
    assert mcp_mod.apply_mileage_review_fix(finding_id=dup["id"], ctx=co).startswith("✅")
    log = json.loads(mcp_mod.get_mileage_log(start_date="2026-10-20", end_date="2026-10-20", ctx=co))
    assert [t["status"] for t in log["trips"]].count("void") == 1
    assert "reminders" in log and "reimbursement" in log


def test_tool_cancelled_on_arrival_crew_scope(env, monkeypatch, mcp_mod):
    _job(env, "JOB-0007", 1.0, 0.0, "Smiths", crew="Pat", crew_uid="u_pat")
    cs = _as(monkeypatch, mcp_mod, CREW_A)
    assert "assigned to you" in mcp_mod.job_cancelled_on_arrival(job_id="JOB-0007", ctx=cs)
    cp = _as(monkeypatch, mcp_mod, CREW_B)
    r = mcp_mod.job_cancelled_on_arrival(job_id="job-0007", gps="1.0,0.0", ctx=cp)
    assert r.startswith("✅"), r
    assert _sql(env, "SELECT job_status FROM jobs WHERE job_id = 'JOB-0007'")[0][0] == "Cancelled"


def test_tool_settings_period_and_reminders(env, monkeypatch, mcp_mod):
    co = _as(monkeypatch, mcp_mod, None)
    assert mcp_mod.mileage_settings(action="set_period", pay_period="biweekly", ctx=co).startswith("✅")
    assert mcp_mod.mileage_settings(action="set_reminder", reminder="weekly", enabled=False, ctx=co).startswith("✅")
    g = json.loads(mcp_mod.mileage_settings(action="get", ctx=co))
    assert g["reimbursement"]["kind"] == "biweekly" and g["reminders"]["weekly"] is False
    cs = _as(monkeypatch, mcp_mod, CREW_A)
    assert "owner or a manager" in mcp_mod.mileage_settings(action="set_period", pay_period="weekly", ctx=cs)
    assert "reminders" in mcp_mod.mileage_settings(action="reminders", ctx=cs).lower()

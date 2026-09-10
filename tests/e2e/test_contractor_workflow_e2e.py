"""
tests/e2e/test_contractor_workflow_e2e.py
===========================================
End-to-end test suite for the broader contractor workflow: customer
lifecycle, route planning, quoting, invoicing (real email delivery), and
time logging — driven by natural-language prompts through the real
Anthropic API, exercising the same "Claude decides which tool to call"
path a contractor uses by voice or text.

This is a SEPARATE suite from test_job_tracker_e2e.py, which covers the
Jobs_Schedule lifecycle (create job -> schedule -> clock in/out -> invoice
-> payment). This suite covers what that one does NOT:
  1. Customer lifecycle — add active, retire (mark inactive) an old one
  2. Route planning — home-to-home multi-stop routing
  3. Quoting — create a quote, then update it (approve / reprice)
  4. Invoice delivery — actually SEND an email (to a real test inbox),
     not just create the Invoices-sheet row
  5. Time log — clock in/out end to end (also covered in the other suite,
     included here too since it's part of this specific ask)

WHY create_customer AND create_quote EXIST
--------------------------------------------
Before this suite was written, there was no way to add a new customer or
create a new quote at all — update_job_spreadsheet only edits EXISTING
rows, and create_job only appends to Jobs_Schedule. Attempting to test
"add a new customer" surfaced this gap directly: Claude would have had no
correct tool to call. create_customer() and create_quote() were added
specifically to close it, modeled on create_job's own append-with-auto-ID
pattern (see _append_sheet_row_impl in ai_prowler_mcp.py).

SAFETY MODEL
------------
- Runs against a SEPARATE test customer (Company Name =
  TEST_CUSTOMER_NAME below) and its associated test quote — never touches
  real customer/quote rows.
- Every write tool auto-backs up before its first write (backup=True
  default) — no separate backup step required.
- Invoice email is sent to a REAL inbox (TEST_EMAIL_TO below,
  david.vavro1@gmail.com per the test requirement) — this is a genuine
  send, not a dry run. Expect a real email to land in that inbox each run.
- At the end of a full pass, restores the pre-suite backup, removing all
  test data. On any failure, the spreadsheet is LEFT AS-IS for inspection.

REQUIREMENTS
------------
  pip install anthropic pytest openpyxl
  Set ANTHROPIC_API_KEY in the environment (or ~/.ai-prowler/claude_api_key.txt
  is NOT read automatically here — that file is Claude Code CLI's own OAuth-
  vs-key mechanism and is not usable by the anthropic Python SDK; export
  ANTHROPIC_API_KEY explicitly for this suite, same as test_job_tracker_e2e.py).
  Email must be configured in AI-Prowler Settings (SMTP or Outlook) for the
  invoice-send stage to succeed.

RUN
---
  pytest tests/e2e/test_contractor_workflow_e2e.py -v -s -x
"""
from __future__ import annotations

import asyncio
import os
import sys
from pathlib import Path

import pytest

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
INSTALL_DIR = Path(os.environ.get("AI_PROWLER_SRC",
                                   r"C:\Program Files\AI-Prowler"))
SPREADSHEET_PATH = Path(os.environ.get(
    "AI_PROWLER_JOB_TRACKER_PATH",
    r"C:\Users\david\Documents\AI-Prowler\AI-Prowler_Job_Tracker.xlsx",
))

TEST_CUSTOMER_NAME = "ZTEST Route Contractor QA"
TEST_RETIRE_CUSTOMER_NAME = "ZTEST Old Customer QA"  # a second, pre-seeded
                                                       # customer this suite
                                                       # marks inactive
TEST_EMAIL_TO = "david.vavro1@gmail.com"  # real inbox, per test requirement
MODEL = "claude-sonnet-4-6"

if str(INSTALL_DIR) not in sys.path:
    sys.path.insert(0, str(INSTALL_DIR))


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture(scope="session")
def anthropic_client():
    anthropic = pytest.importorskip("anthropic")
    if not os.environ.get("ANTHROPIC_API_KEY"):
        pytest.skip("ANTHROPIC_API_KEY not set — skipping live tool-use evals")
    return anthropic.Anthropic()


@pytest.fixture(scope="session")
def mcp_module():
    import ai_prowler_mcp as m
    return m


@pytest.fixture(scope="session")
def tool_schemas(mcp_module):
    """Real, live tool schemas from the actual FastMCP server — not
    hand-written duplicates, so a docstring/parameter change is
    automatically reflected here with no drift possible."""
    wanted = {
        "create_customer", "create_quote", "update_job_spreadsheet",
        "read_job_spreadsheet", "optimize_route", "build_maps_url",
        "get_home_address", "create_invoice", "email_invoice",
        "log_time_entry",
    }

    async def _list():
        return await mcp_module.mcp.list_tools()

    tools = asyncio.run(_list())
    schemas = [
        {"name": t.name, "description": t.description,
         "input_schema": t.inputSchema}
        for t in tools if t.name in wanted
    ]
    found = {s["name"] for s in schemas}
    missing = wanted - found
    assert not missing, f"Expected tools not found on server: {missing}"
    return schemas


@pytest.fixture(scope="session")
def pre_suite_backup_path(mcp_module):
    backup_msg = mcp_module._backup_spreadsheet(str(SPREADSHEET_PATH))
    assert "Backup saved" in backup_msg, f"Pre-suite backup failed: {backup_msg}"
    rel_path = backup_msg.split("Backup saved:")[1].strip()
    return SPREADSHEET_PATH.parent / rel_path


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def call_claude_tool(client, schemas, prompt: str):
    response = client.messages.create(
        model=MODEL,
        max_tokens=1024,
        tools=schemas,
        messages=[{"role": "user", "content": prompt}],
    )
    calls = [b for b in response.content if b.type == "tool_use"]
    return calls[0] if calls else None


def _load_sheet_rows(sheet_name: str) -> list[dict]:
    """Read every row of `sheet_name` as a list of {header: value} dicts,
    using the same header-detection convention every write tool uses
    (first row in rows 1-5 with >= 3 non-empty cells).

    BUG FIXED: previously scanned data rows from a HARDCODED min_row=6
    rather than the actual detected header row + 1 — found via
    test_route_scheduling_e2e.py seeding 3 same-day jobs and directly
    exposing that rows landing before row 6 (e.g. a 2nd or 3rd new row on
    a sheet with only 1-2 rows of preamble) were silently skipped.
    """
    import openpyxl
    wb = openpyxl.load_workbook(str(SPREADSHEET_PATH), data_only=True)
    ws = wb[sheet_name]

    header_row_num = None
    headers = None
    for row in ws.iter_rows(min_row=1, max_row=5):
        non_empty = [c for c in row if c.value is not None]
        if len(non_empty) >= 3:
            header_row_num = row[0].row
            headers = [str(c.value).strip().replace("\n", " ") if c.value else ""
                       for c in row]
            break
    assert headers, f"Could not detect header row in {sheet_name}"

    rows = []
    for row in ws.iter_rows(min_row=header_row_num + 1, values_only=True):
        if any(v is not None for v in row):
            rows.append(dict(zip(headers, row)))
    return rows


def find_row_by_name(sheet_name: str, name_col: str, name_val: str) -> "dict | None":
    for row in _load_sheet_rows(sheet_name):
        if row.get(name_col) == name_val:
            return row
    return None


# ---------------------------------------------------------------------------
# Test class — customer lifecycle, routing, quoting, invoicing, time log
# ---------------------------------------------------------------------------
@pytest.mark.job_sheet_e2e
class TestContractorWorkflowE2E:

    customer_id: "str | None" = None
    quote_id: "str | None" = None
    retire_customer_id: "str | None" = None

    # ── 1. Customer lifecycle: add active, retire an old one ───────────────

    def test_01_add_active_customer(self, anthropic_client, tool_schemas,
                                     mcp_module, pre_suite_backup_path):
        prompt = (
            f"Add a new customer: {TEST_CUSTOMER_NAME}, "
            "555 Route Test Blvd, New Smyrna Beach FL 32168, "
            "email ztest.route@example.com, monthly window washing, "
            "status active"
        )
        call = call_claude_tool(anthropic_client, tool_schemas, prompt)
        assert call is not None, "Claude called no tool for an add-customer prompt"
        assert call.name == "create_customer", (
            f"Expected create_customer, got {call.name} — this is exactly "
            f"the gap create_customer was added to close; if Claude picked "
            f"a different tool, check whether it's compensating for a "
            f"missing/misleading tool description rather than a real fit."
        )
        updates = call.input.get("updates", {})
        assert TEST_CUSTOMER_NAME in str(updates.get("Company Name", "")) or \
               TEST_CUSTOMER_NAME in str(updates), (
            f"Customer name missing/wrong in Claude's tool call: {updates}"
        )

        result = mcp_module.create_customer(updates=updates, backup=True)
        assert result.startswith("✅"), f"create_customer failed: {result}"
        assert "NEW_CUST_ID=" in result, f"No CustomerID in result: {result}"
        TestContractorWorkflowE2E.customer_id = (
            result.split("NEW_CUST_ID=")[1].splitlines()[0].strip())

        row = find_row_by_name("Customers", "Company Name", TEST_CUSTOMER_NAME)
        assert row is not None, "Test customer not found after create_customer"
        assert row.get("State") == "FL"
        assert str(row.get("Status Active/Inactive", "")).lower() == "active"

    def test_01b_retire_old_customer(self, anthropic_client, tool_schemas,
                                      mcp_module):
        """Retiring (marking inactive) an EXISTING customer uses
        update_job_spreadsheet, not create_customer — create_customer
        always appends a new row and must never be used for this."""
        # Seed a second customer to retire, so this test doesn't depend on
        # any real customer already existing in the sheet.
        seed_result = mcp_module.create_customer(
            updates={
                "Company Name": TEST_RETIRE_CUSTOMER_NAME,
                "State": "FL",
                "Status Active/Inactive": "Active",
            },
            backup=True,
        )
        assert seed_result.startswith("✅")
        TestContractorWorkflowE2E.retire_customer_id = (
            seed_result.split("NEW_CUST_ID=")[1].splitlines()[0].strip())

        prompt = f"Mark {TEST_RETIRE_CUSTOMER_NAME} as inactive — we don't service them anymore"
        call = call_claude_tool(anthropic_client, tool_schemas, prompt)
        assert call is not None
        assert call.name == "update_job_spreadsheet", (
            f"Expected update_job_spreadsheet to EDIT the existing customer, "
            f"got {call.name} — retiring a customer must never create a new "
            f"row via create_customer."
        )
        assert call.input.get("sheet_name") in ("Customers", None, ""), (
            f"Expected sheet_name='Customers' or omitted-with-correct-default, "
            f"got {call.input.get('sheet_name')!r}"
        )

        result = mcp_module.update_job_spreadsheet(
            job_identifier=TEST_RETIRE_CUSTOMER_NAME,
            id_column="Company Name",
            sheet_name="Customers",
            updates=call.input["updates"],
        )
        assert result.startswith("✅"), f"Retire update failed: {result}"

        row = find_row_by_name("Customers", "Company Name",
                                TEST_RETIRE_CUSTOMER_NAME)
        assert row is not None
        assert str(row.get("Status Active/Inactive", "")).lower() == "inactive", (
            f"Customer was not marked inactive: {row.get('Status Active/Inactive')!r}"
        )
        # The ACTIVE customer from test_01 must be unaffected by this update.
        active_row = find_row_by_name("Customers", "Company Name",
                                       TEST_CUSTOMER_NAME)
        assert str(active_row.get("Status Active/Inactive", "")).lower() == "active", (
            "REGRESSION: retiring one customer affected a different customer row"
        )

    # ── 2. Routing: home-to-home multi-stop ─────────────────────────────────

    def test_02_route_jobs_home_to_home(self, anthropic_client, tool_schemas,
                                         mcp_module):
        """optimize_route + build_maps_url, starting and ending at the
        configured home/business address. Uses get_home_address() rather
        than a hardcoded origin, matching how a contractor would actually
        phrase this — 'from my house' / 'starting from home' — with no
        address of their own included in the prompt."""
        home_call = call_claude_tool(
            anthropic_client, tool_schemas,
            "What's my home address on file?")
        # get_home_address takes no arguments — the only thing worth
        # asserting is that Claude reaches for the RIGHT tool when a prompt
        # references "home"/"my address", per that tool's own docstring
        # rationale (this used to be a real gap: no tool existed at all).
        assert home_call is not None
        assert home_call.name == "get_home_address"

        home_address = mcp_module.get_home_address()
        assert "❌" not in home_address and "not configured" not in home_address.lower(), (
            f"get_home_address() did not return a usable address: {home_address}\n"
            "Set one in AI-Prowler Settings before running this test."
        )

        # Real, geocodable addresses — a made-up test address like
        # "99 Test Ave" is not in OpenStreetMap's Nominatim database and
        # gets silently excluded from the route ("Could not geocode"),
        # which would make this test's own fixture data the failure rather
        # than the code under test. Nominatim (the free geocoding backend)
        # can also throw transient ConnectionResetError under load — if
        # this test flakes on a geocode failure, retry once before treating
        # it as a real regression.
        test_stops = [
            "412 Pelican Dr, Daytona Beach, FL 32118",
            "100 S Atlantic Ave, Ormond Beach, FL 32176",
        ]

        route_prompt = (
            f"Plan today's route starting and ending at home, "
            f"visiting {test_stops[0]} and {test_stops[1]}"
        )
        route_call = call_claude_tool(anthropic_client, tool_schemas, route_prompt)
        assert route_call is not None
        assert route_call.name == "optimize_route"
        assert route_call.input.get("return_to_origin", True) is True, (
            "Expected return_to_origin=True (or its default) for a "
            "home-to-home route prompt"
        )

        route_result = mcp_module.optimize_route(
            origin=home_address,
            stops=test_stops,
            return_to_origin=True,
        )
        assert "❌" not in route_result, f"optimize_route failed: {route_result}"
        # Confirm both test stops appear somewhere in the optimized output —
        # not asserting a specific order, since TSP ordering is legitimately
        # allowed to vary run to run based on live routing data.
        for stop in test_stops:
            street = stop.split(",")[0]
            assert street in route_result, (
                f"Stop {street!r} missing from optimize_route output"
            )

        maps_call = call_claude_tool(
            anthropic_client, tool_schemas,
            "Give me a tap-to-navigate link for that route")
        assert maps_call is not None
        assert maps_call.name == "build_maps_url"

        maps_result = mcp_module.build_maps_url(
            origin=home_address, stops=test_stops)
        assert "http" in maps_result.lower(), (
            f"build_maps_url did not return a usable URL: {maps_result}"
        )

    # ── 3. Quoting: create, then update ─────────────────────────────────────

    def test_03_create_quote(self, anthropic_client, tool_schemas, mcp_module):
        prompt = (
            f"Create a quote for {TEST_CUSTOMER_NAME} — window washing, "
            "$150, valid for 30 days, status open"
        )
        call = call_claude_tool(anthropic_client, tool_schemas, prompt)
        assert call is not None
        assert call.name == "create_quote", (
            f"Expected create_quote, got {call.name} — same gap "
            f"create_customer closed, now closed for Quotes."
        )
        updates = call.input.get("updates", {})

        result = mcp_module.create_quote(updates=updates, backup=True)
        assert result.startswith("✅"), f"create_quote failed: {result}"
        assert "NEW_QTE_ID=" in result, f"No QuoteID in result: {result}"
        TestContractorWorkflowE2E.quote_id = (
            result.split("NEW_QTE_ID=")[1].splitlines()[0].strip())

        row = find_row_by_name("Quotes", "Customer Name / Company",
                                TEST_CUSTOMER_NAME)
        assert row is not None, "Test quote not found after create_quote"
        assert str(row.get("Status (Open/Approved/Declined)", "")).lower() == "open"

    def test_03b_update_quote_to_approved(self, anthropic_client, tool_schemas,
                                           mcp_module):
        """Updating an EXISTING quote uses update_job_spreadsheet with
        sheet_name='Quotes' — create_quote must never be called again for
        an edit, since it always appends a new row."""
        assert self.quote_id, "test_03 must run first and set quote_id"
        prompt = f"Approve the quote for {TEST_CUSTOMER_NAME}"
        call = call_claude_tool(anthropic_client, tool_schemas, prompt)
        assert call is not None
        assert call.name == "update_job_spreadsheet", (
            f"Expected update_job_spreadsheet to EDIT the existing quote, "
            f"got {call.name}"
        )

        result = mcp_module.update_job_spreadsheet(
            job_identifier=self.quote_id,
            id_column="QuoteID (QTE-####)",
            sheet_name="Quotes",
            updates=call.input["updates"],
        )
        assert result.startswith("✅"), f"Quote approval update failed: {result}"

        row = find_row_by_name("Quotes", "Customer Name / Company",
                                TEST_CUSTOMER_NAME)
        assert str(row.get("Status (Open/Approved/Declined)", "")).lower() == "approved", (
            f"Quote was not marked approved: "
            f"{row.get('Status (Open/Approved/Declined)')!r}"
        )
        # A second create_quote call was NOT made — confirm exactly one
        # quote row still exists for this customer (no accidental duplicate).
        all_quotes = [r for r in _load_sheet_rows("Quotes")
                      if r.get("Customer Name / Company") == TEST_CUSTOMER_NAME]
        assert len(all_quotes) == 1, (
            f"Expected exactly 1 quote for {TEST_CUSTOMER_NAME}, "
            f"found {len(all_quotes)} — approving a quote must not "
            f"create a duplicate row."
        )

    # ── 4. Time log ──────────────────────────────────────────────────────────

    def test_04_create_job_and_time_log(self, anthropic_client, tool_schemas,
                                         mcp_module):
        """This suite's invoice test needs a real job row with a real
        price, so create one here from the approved quote — then run a
        full clock in/out cycle against it."""
        job_result = mcp_module.create_job(
            updates={
                "Customer Name / Company": TEST_CUSTOMER_NAME,
                "Service Type": "Window Washing",
                "Service Date": "2026-09-15",
                "Job Status": "Scheduled",
                "Quote Amount ($)": 150,
            },
            backup=True,
        )
        assert job_result.startswith("✅")
        job_id = job_result.split("NEW_JOB_ID=")[1].splitlines()[0].strip()
        TestContractorWorkflowE2E.job_id = job_id

        in_call = call_claude_tool(
            anthropic_client, tool_schemas,
            f"Clock in for the {TEST_CUSTOMER_NAME} job")
        assert in_call is not None and in_call.name == "log_time_entry"
        assert in_call.input.get("action") == "start"
        mcp_module.log_time_entry(job_identifier=job_id, action="start")

        out_call = call_claude_tool(
            anthropic_client, tool_schemas,
            "Clock out, job's done")
        assert out_call is not None and out_call.name == "log_time_entry"
        assert out_call.input.get("action") == "stop"
        result = mcp_module.log_time_entry(job_identifier=job_id, action="stop")
        assert "Clocked OUT" in result

        row = find_row_by_name("Jobs_Schedule", "Customer Name / Company",
                                TEST_CUSTOMER_NAME)
        assert row is not None
        assert row.get("Actual Duration Unit") == "min", (
            f"REGRESSION: Actual Duration Unit not populated — "
            f"got {row.get('Actual Duration Unit')!r}"
        )

        mcp_module.update_job_spreadsheet(
            job_identifier=job_id,
            id_column="JobID (JOB-####)",
            updates={"Job Status": "Completed"},
        )

    # ── 5. Invoicing: real send to a real inbox ─────────────────────────────

    def test_05_create_and_send_invoice(self, anthropic_client, tool_schemas,
                                         mcp_module):
        """Creates the Invoices-sheet row via create_invoice, then actually
        SENDS the invoice email to TEST_EMAIL_TO. This is a real send, not
        a dry run — check that inbox after running this suite.
        """
        assert self.job_id, "test_04 must run first and set job_id"

        invoice_call = call_claude_tool(
            anthropic_client, tool_schemas,
            f"Invoice the {TEST_CUSTOMER_NAME} job")
        assert invoice_call is not None
        assert invoice_call.name == "create_invoice"

        result = mcp_module.create_invoice(job_identifier=self.job_id)
        assert result.startswith("✅"), f"create_invoice failed: {result}"
        assert "NEW_INVOICE_ID=" in result
        invoice_id = result.split("NEW_INVOICE_ID=")[1].splitlines()[0].strip()
        TestContractorWorkflowE2E.invoice_id = invoice_id

        send_call = call_claude_tool(
            anthropic_client, tool_schemas,
            f"Email the {TEST_CUSTOMER_NAME} invoice to {TEST_EMAIL_TO}")
        assert send_call is not None
        assert send_call.name == "email_invoice"
        assert TEST_EMAIL_TO in str(send_call.input.get("to", "")), (
            f"Expected 'to' to be {TEST_EMAIL_TO}, got {send_call.input.get('to')!r}"
        )

        send_result = mcp_module.email_invoice(
            invoice_identifier=invoice_id, to=TEST_EMAIL_TO)
        assert send_result.startswith("✅"), (
            f"email_invoice failed to send — check that email is configured "
            f"in AI-Prowler Settings: {send_result}"
        )

    def test_05b_invoice_html_has_business_identity_and_no_flexbox(
            self, mcp_module):
        """Regression test for two invoice-template bugs found by manual
        inspection of a real sent email (neither is visible from a plain
        function-return-value check — both only show up in the actual
        rendered HTML):

        1. The invoice header showed no business identity at all — Settings
           sheet's "Business Name" (and Phone/Email/Address/Website) rows
           are explicitly documented as "Appears on invoices" but nothing
           read them. _read_business_info() was added to fix this.

        2. The "Bill To / Service Date / ... " and "Subtotal / Tax / TOTAL
           DUE" sections used CSS `display: flex`, which most email clients
           (Outlook especially) do not support — labels and values were
           rendered running together with zero space. Fixed by switching
           to <table> layout, the only markup reliably supported across
           email clients.

        This test captures the actual HTML the send would have used by
        monkeypatching _send_smtp to intercept the body_html argument,
        so it exercises the real template-building code path without
        requiring a second live email send (test_05 above already sent
        one — reusing that invoice here keeps this test additive rather
        than triggering another real send).
        """
        captured = {}

        def _fake_send_smtp(to, subject, body, body_html=None, **kwargs):
            captured["body_html"] = body_html
            return (True, "✅ intercepted, not actually sent")

        import ai_prowler_mcp as _m
        _real_send_smtp = _m._send_smtp
        _m._send_smtp = _fake_send_smtp
        try:
            mcp_module.email_invoice(
                invoice_identifier=self.invoice_id, to=TEST_EMAIL_TO)
        finally:
            _m._send_smtp = _real_send_smtp

        html = captured.get("body_html", "")
        assert html, "email_invoice did not produce an HTML body to inspect"

        # Bug 1 — business identity must appear in the header. Read the
        # expected name from Settings directly rather than hardcoding a
        # specific business name, so this test works against any
        # spreadsheet's own configured Business Name, not just this one.
        biz_info = mcp_module._read_business_info(str(SPREADSHEET_PATH))
        expected_name = biz_info.get("name") or "Your Business"
        assert expected_name in html, (
            f"REGRESSION: Business Name ({expected_name!r}) from Settings "
            f"sheet missing from invoice header"
        )

        # Bug 2 — flexbox must not be used for the label/value rows
        # (the specific classes that broke Outlook rendering)
        assert "display: flex" not in html, (
            "REGRESSION: invoice template uses display:flex again — "
            "this collapses label/value spacing in Outlook and most "
            "other email clients. Use <table> layout instead."
        )
        assert "<table" in html and "Bill To" in html, (
            "Expected the Bill To section to be table-based markup"
        )

    # ── 6. Cleanup ────────────────────────────────────────────────────────────

    def test_06_restore_backup_on_full_pass(self, mcp_module,
                                             pre_suite_backup_path):
        import shutil
        assert pre_suite_backup_path.exists(), (
            f"Pre-suite backup missing: {pre_suite_backup_path}"
        )
        shutil.copy2(str(pre_suite_backup_path), str(SPREADSHEET_PATH))

        assert find_row_by_name(
            "Customers", "Company Name", TEST_CUSTOMER_NAME) is None
        assert find_row_by_name(
            "Customers", "Company Name", TEST_RETIRE_CUSTOMER_NAME) is None
        assert find_row_by_name(
            "Quotes", "Customer Name / Company", TEST_CUSTOMER_NAME) is None
        assert find_row_by_name(
            "Jobs_Schedule", "Customer Name / Company", TEST_CUSTOMER_NAME) is None

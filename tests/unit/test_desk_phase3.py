"""Service Desk Phase 3 — SLA clocks, routing, reports, export/import,
backup/restore and the scheduler."""
import csv
import datetime
import io
import json
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import desk_attachments  # noqa: E402
import desk_backup  # noqa: E402
import desk_data  # noqa: E402
import desk_db  # noqa: E402
import desk_files  # noqa: E402
import desk_reports  # noqa: E402
import desk_routing  # noqa: E402
import desk_rules as r  # noqa: E402
import desk_scheduler  # noqa: E402
import desk_sla  # noqa: E402
import desk_tools  # noqa: E402

UTC = datetime.timezone.utc
OWNER_S = {"id": "dave", "role": "owner"}
MANAGER = {"id": "mgr", "role": "manager", "queues": ["field"]}
CREW = {"id": "cris", "role": "field_crew"}


@pytest.fixture
def db(tmp_path):
    path = str(tmp_path / "tickets_database" / desk_db.DEFAULT_DB_FILENAME)
    desk_db.init_db(path)
    desk_routing.availability_fn = None
    desk_routing.job_crew_fn = None
    desk_sla.working_days_fn = None
    yield path
    desk_routing.availability_fn = None
    desk_routing.job_crew_fn = None
    desk_db.listeners.clear()


def dt(s):
    return datetime.datetime.fromisoformat(s).replace(tzinfo=UTC)


# ── SLA math ─────────────────────────────────────────────────────────────────

def test_business_minutes_skip_nights_and_weekends():
    fri_4pm = dt("2026-10-30T16:00:00")          # Friday
    assert desk_sla.add_business_minutes(fri_4pm, 120, tz=UTC) == dt("2026-11-02T09:00:00")
    sat = dt("2026-10-31T10:00:00")
    assert desk_sla.add_business_minutes(sat, 60, tz=UTC) == dt("2026-11-02T09:00:00")
    early = dt("2026-10-27T06:00:00")
    assert desk_sla.add_business_minutes(early, 30, tz=UTC) == dt("2026-10-27T08:30:00")
    assert desk_sla.end_of_business_day(dt("2026-10-30T18:00:00"), tz=UTC) == dt("2026-11-02T17:00:00")


def test_due_times_per_priority():
    t0 = dt("2026-10-27T09:00:00")               # Tuesday 9am
    assert desk_sla.due_times("P1", t0, tz=UTC) == ("2026-10-27T10:00:00+00:00", "2026-10-27T17:00:00+00:00")
    r2, v2 = desk_sla.due_times("P2", t0, tz=UTC)
    assert r2 == "2026-10-27T13:00:00+00:00" and v2 == "2026-10-29T09:00:00+00:00"
    assert desk_sla.due_times("P4", t0, tz=UTC)[1] is None


def test_working_days_hook_accepts_names():
    desk_sla.working_days_fn = lambda: ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
    try:
        assert desk_sla.working_days() == (0, 1, 2, 3, 4, 5)
    finally:
        desk_sla.working_days_fn = None


def test_status_warn_breach_escalate():
    t = {"state": r.IN_PROGRESS, "kind": "internal", "created_at": "2026-10-27T09:00:00+00:00",
         "due_resolve_at": "2026-10-27T13:00:00+00:00", "first_response_at": None}
    assert desk_sla.status(t, dt("2026-10-27T10:00:00"))["resolve"] == "ok"
    assert desk_sla.status(t, dt("2026-10-27T12:30:00"))["resolve"] == "warn"
    s = desk_sla.status(t, dt("2026-10-27T14:00:00"))
    assert s["resolve"] == "breach" and not s["escalate"]
    assert desk_sla.status(t, dt("2026-10-27T18:00:00"))["escalate"]
    assert desk_sla.status({**t, "state": r.ON_HOLD}, dt("2026-10-28T00:00:00"))["resolve"] == "paused"


def test_business_ticket_gets_sla_and_solo_does_not(db):
    b = desk_db.create_ticket(db, title="Pump", kind="internal", actor=MANAGER, priority="P2")
    assert b["due_response_at"] and b["due_resolve_at"]
    s = desk_db.create_ticket(db, title="Tire", kind="internal", edition=r.SOLO, priority="P2")
    assert s["due_response_at"] is None and s["due_resolve_at"] is None
    own = desk_db.create_ticket(db, title="Hose", kind="internal", actor=MANAGER,
                                due_resolve_at="2027-01-01T00:00:00+00:00")
    assert own["due_resolve_at"] == "2027-01-01T00:00:00+00:00"          # an explicit due wins


def test_on_hold_pauses_the_clock(db):
    t = desk_db.create_ticket(db, title="Pump", kind="internal", actor=MANAGER, assignee="cris")
    due0 = desk_sla.parse(t["due_resolve_at"])
    desk_db.change_state(db, t["ticket_id"], r.IN_PROGRESS, actor=CREW)
    desk_db.change_state(db, t["ticket_id"], r.ON_HOLD, actor=CREW, reason="parts")
    con = desk_db.get_connection(db)
    try:
        con.execute("UPDATE tickets SET hold_started_at = ? WHERE ticket_id = ?",
                    ((datetime.datetime.now(UTC) - datetime.timedelta(hours=3)).isoformat(timespec="seconds"),
                     t["ticket_id"]))
    finally:
        con.close()
    t2 = desk_db.change_state(db, t["ticket_id"], r.IN_PROGRESS, actor=CREW)
    shifted = desk_sla.parse(t2["due_resolve_at"]) - due0
    assert datetime.timedelta(hours=2, minutes=59) < shifted < datetime.timedelta(hours=3, minutes=1)
    assert t2["hold_started_at"] is None


# ── routing ──────────────────────────────────────────────────────────────────

def test_round_robin_skips_unavailable(db):
    desk_routing.upsert_queue(db, "field", name="Field", members=["ann", "bob", "cris"], actor=OWNER_S)
    desk_routing.upsert_type(db, "redo", name="Re-do", kind="external", default_queue="field", actor=OWNER_S)
    desk_routing.availability_fn = lambda u: u != "bob"          # bob is on time off
    got = [desk_db.create_ticket(db, title=f"t{i}", type="redo", contact_id="CT-1", actor=MANAGER)["assignee"]
           for i in range(4)]
    assert got == ["ann", "cris", "ann", "cris"]
    ev = desk_db.get_ticket_detail(db, "T-0001")["events"]
    assert any(e["kind"] == "assigned" and e["actor"].startswith("type:redo") for e in ev)


def test_rules_in_order_and_fallbacks(db):
    desk_routing.upsert_queue(db, "office", name="Office", members=["mia"], actor=OWNER_S)
    desk_routing.upsert_rule(db, "billing", match={"keyword": "invoice"}, action={"user": "dave"},
                             sort_order=1, actor=OWNER_S)
    desk_routing.upsert_rule(db, "it", match={"type": "it"}, action={"queue": "office"}, sort_order=2,
                             actor=OWNER_S)
    assert desk_db.create_ticket(db, title="Invoice wrong", kind="internal", actor=CREW)["assignee"] == "dave"
    t = desk_db.create_ticket(db, title="Laptop", kind="internal", type="it", actor=CREW)
    assert t["assignee"] == "mia" and t["queue_id"] == "office"
    t = desk_db.create_ticket(db, title="Something", kind="internal", actor=CREW)
    assert t["assignee"] is None and t["state"] == r.OPEN
    desk_routing.job_crew_fn = lambda job: "cris" if job == "JOB-9" else None
    t = desk_db.create_ticket(db, title="Redo", contact_id="CT-1", job_id="JOB-9", actor=MANAGER)
    assert t["assignee"] == "cris"
    with pytest.raises(desk_db.DeskError, match="a rule matches"):
        desk_routing.upsert_rule(db, "bad", match={"color": "red"}, action={"user": "x"}, actor=OWNER_S)


def test_queues_are_business_and_owner_only(db):
    with pytest.raises(desk_db.DeskError, match="Business feature"):
        desk_routing.upsert_queue(db, "q", members=["a"], edition=r.SOLO)
    with pytest.raises(desk_db.DeskError, match="only the owner"):
        desk_routing.upsert_queue(db, "q", members=["a"], actor=MANAGER)
    assert "Business feature" in desk_tools.manage_ticket_queues(db, None, r.SOLO, action="list")


# ── reports and SLA check ────────────────────────────────────────────────────

def test_sla_check_flags_once_and_emits(db):
    seen = []
    desk_db.listeners.append(lambda kind, t, a, **i: seen.append((kind, t["ticket_id"], i.get("flag"))))
    t = desk_db.create_ticket(db, title="Pump", kind="internal", actor=MANAGER, assignee="cris")
    due = desk_sla.parse(t["due_resolve_at"])
    created = desk_sla.parse(t["created_at"])
    seen.clear()
    late = due + (due - created) * 0.5          # 1.5x the target: breached, not yet escalated
    assert ("T-0001", "resolve_breach") in desk_reports.sla_check(db, edition=r.BUSINESS, now=late)
    assert [s for s in seen if s[0] == "sla"]
    assert desk_reports.sla_check(db, edition=r.BUSINESS, now=late) == []        # once only
    esc = desk_reports.sla_check(db, edition=r.BUSINESS, now=due + (due - created) * 3)
    assert esc == [("T-0001", "escalated")]
    assert desk_reports.sla_check(db, edition=r.SOLO, now=late) == []


def test_report_and_groups(db):
    for who in ("cris", "cris", "ann"):
        desk_db.create_ticket(db, title="x", kind="internal", actor=OWNER_S, assignee=who)
    desk_db.change_state(db, "T-0001", r.CLOSED, actor=OWNER_S, reason="done")
    rep = desk_reports.report(db, actor=OWNER_S, group_by="assignee")
    assert rep["total"]["tickets"] == 3 and rep["total"]["closed"] == 1
    assert rep["groups"]["cris"]["tickets"] == 2
    assert "3 ticket(s)" in desk_reports.report_text(rep)
    with pytest.raises(desk_db.DeskError, match="only the owner or a manager"):
        desk_reports.report(db, actor=CREW)
    with pytest.raises(desk_db.DeskError, match="Business feature"):
        desk_reports.report(db, edition=r.SOLO, group_by="assignee")


def test_digest(db):
    assert desk_reports.digest_text(db, edition=r.SOLO) == ""
    desk_db.create_ticket(db, title="Renew license", kind="internal", edition=r.SOLO,
                          due_resolve_at="2026-01-01T00:00:00+00:00")
    text = desk_reports.digest_text(db, edition=r.SOLO)
    assert "Overdue (1)" in text and "Renew license" in text


# ── export / import ──────────────────────────────────────────────────────────

def _seed(db):
    c = desk_db.upsert_contact(db, name="Lopez", phone="386-555-0100")
    t = desk_db.create_ticket(db, title="Streaks", contact_id=c["contact_id"], edition=r.SOLO)
    desk_db.add_message(db, t["ticket_id"], "internal note", edition=r.SOLO)
    desk_db.add_message(db, t["ticket_id"], "On our way", edition=r.SOLO, visibility="public")
    desk_attachments.store_attachment(db, t["ticket_id"], b"%PDF-1.4 quote", "quote.pdf", edition=r.SOLO)
    return t


def test_export_xlsx_csv_pdf(db):
    _seed(db)
    name, data, mime = desk_data.export(db, "xlsx", edition=r.SOLO)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Tickets", "Messages", "Events"]
    assert wb["Messages"].max_row == 2                     # public only by default
    name, data, mime = desk_data.export(db, "csv", edition=r.SOLO, include_internal=True)
    z = zipfile.ZipFile(io.BytesIO(data))
    rows = list(csv.DictReader(io.StringIO(z.read("messages.csv").decode("utf-8-sig"))))
    assert len(rows) == 2
    name, data, mime = desk_data.export(db, "pdf", edition=r.SOLO, ticket_id="t-0001")
    assert data.startswith(b"%PDF") and b"Streaks" in data and b"internal note" not in data
    with pytest.raises(desk_db.DeskError, match="only the owner or a manager"):
        desk_data.export(db, "xlsx", actor=CREW)


def test_full_export_round_trips_into_another_install(db, tmp_path):
    _seed(db)
    name, data, mime = desk_data.export(db, "json", edition=r.SOLO)
    other = str(tmp_path / "other" / desk_db.DEFAULT_DB_FILENAME)
    desk_db.init_db(other)
    dry = desk_data.import_tickets(other, data, name, edition=r.BUSINESS, actor=OWNER_S)
    assert dry["dry_run"] and dry["to_add"] == 1
    out = desk_data.import_tickets(other, data, name, edition=r.BUSINESS, actor=OWNER_S, dry_run=False)
    assert out["added"] == ["T-0001"]
    d = desk_db.get_ticket_detail(other, "T-0001")
    assert [m["body"] for m in d["messages"]] == ["internal note", "On our way"]
    assert desk_db.get_contact(other, d["contact_id"])["name"] == "Lopez"
    assert desk_files.verify(other) == []                  # attachment copied with its hash
    again = desk_data.import_tickets(other, data, name, edition=r.BUSINESS, actor=OWNER_S, dry_run=False)
    assert again["to_add"] == 0 and len(again["duplicates"]) == 1


def test_csv_import_dry_run_quiet_import_and_undo(db):
    seen = []
    desk_db.listeners.append(lambda *a, **k: seen.append(a[0]))
    body = ("Subject,Status,Priority,Requester,Phone,ID\n"
            "Gutters overflowing,open,high,Ann,386-555-0101,Z-1\n"
            "Window cracked,solved,urgent,Bo,386-555-0102,Z-2\n"
            ",open,low,Cy,386-555-0103,Z-3\n").encode()
    dry = desk_data.import_tickets(db, body, "zendesk.csv", actor=OWNER_S)
    assert dry["to_add"] == 2 and dry["invalid"][0]["row"] == 4 and dry["new_contacts"] == 2
    assert desk_db.list_tickets(db) == []
    backups = []
    out = desk_data.import_tickets(db, body, "zendesk.csv", actor=OWNER_S, dry_run=False,
                                   backup_fn=lambda trig: backups.append(trig))
    assert backups == ["pre-import"] and len(out["added"]) == 2 and seen == []
    t2 = desk_db.get_ticket(db, out["added"][1])
    assert t2["state"] == r.FIXED and t2["priority"] == "P1" and t2["source"] == "import"
    dup = desk_data.import_tickets(db, body, "zendesk.csv", actor=OWNER_S)
    assert dup["to_add"] == 0 and len(dup["duplicates"]) == 2
    with pytest.raises(desk_db.DeskError, match="7 days"):
        desk_data.undo_import(db, out["batch_id"], actor=OWNER_S,
                              now=datetime.datetime.now(UTC) + datetime.timedelta(days=8))
    gone = desk_data.undo_import(db, out["batch_id"], actor=OWNER_S)
    assert sorted(gone["removed"]) == sorted(out["added"]) and desk_db.list_tickets(db) == []
    assert list(desk_files.iter_ticket_json_files(desk_db.root_for(db))) == []
    with pytest.raises(desk_db.DeskError, match="only the owner"):
        desk_data.import_tickets(db, body, "x.csv", actor=MANAGER)


# ── backup / restore ─────────────────────────────────────────────────────────

def test_backup_verify_list_and_restore_one_ticket(db):
    t = _seed(db)
    b = desk_backup.backup_now(db, "manual", now=datetime.datetime(2026, 10, 27, 10, 0, 0))
    assert b["verified"]
    names = zipfile.ZipFile(b["path"]).namelist()
    assert desk_db.DEFAULT_DB_FILENAME in names and any(n.endswith("0001_quote.pdf") for n in names)
    assert desk_backup.list_backups(db)[0]["name"] == b["name"]
    desk_db.update_ticket(db, t["ticket_id"], {"title": "CHANGED"}, edition=r.SOLO)
    pending = desk_backup.restore(db, b["name"], ticket_id=t["ticket_id"], edition=r.SOLO)
    assert pending["needs_confirm"] and desk_db.get_ticket(db, t["ticket_id"])["title"] == "CHANGED"
    out = desk_backup.restore(db, b["name"], ticket_id=t["ticket_id"], edition=r.SOLO, confirm=True)
    assert out["ok"] and desk_db.get_ticket(db, t["ticket_id"])["title"] == "Streaks"
    assert any(x["trigger"] == "pre-restore" for x in desk_backup.list_backups(db))
    assert desk_files.verify(db) == []
    with pytest.raises(desk_db.DeskError, match="only the owner"):
        desk_backup.restore(db, b["name"], actor=MANAGER, confirm=True)


def test_restore_everything(db):
    _seed(db)
    b = desk_backup.backup_now(db, "manual", now=datetime.datetime(2026, 10, 27, 10, 0, 0))
    desk_db.create_ticket(db, title="after backup", kind="internal", edition=r.SOLO)
    desk_backup.restore(db, b["name"], edition=r.SOLO, confirm=True)
    assert [t["title"] for t in desk_db.list_tickets(db)] == ["Streaks"]
    assert desk_files.verify(db) == []


def test_retention_and_due(db):
    base = datetime.datetime(2026, 10, 27, 2, 0, 0)
    for i in range(20):
        desk_backup.backup_now(db, "scheduled", now=base - datetime.timedelta(days=i))
    gone = desk_backup.prune(db, now=base)
    left = desk_backup.list_backups(db)
    assert len(gone) > 0 and len(left) < 20
    days = {x["made_at"][:10] for x in left}
    assert all((base - datetime.timedelta(days=i)).date().isoformat() in days for i in range(14))
    s = desk_db.get_settings(db)
    assert desk_backup.due(s, None, base)
    assert not desk_backup.due(s, "2026-10-27T02:00:05", datetime.datetime(2026, 10, 27, 23, 0))
    assert desk_backup.due(s, "2026-10-27T02:00:05", datetime.datetime(2026, 10, 28, 2, 1))


# ── scheduler ────────────────────────────────────────────────────────────────

def test_scheduler_digest_once_a_day_and_backup(db):
    desk_db.create_ticket(db, title="Late", kind="internal", edition=r.SOLO,
                          due_resolve_at="2026-01-01T00:00:00+00:00")
    sent = []
    t8 = datetime.datetime(2026, 10, 27, 8, 0).astimezone()
    did = desk_scheduler.run_once(db, edition=r.SOLO, now=t8, send_digest=sent.append)
    assert did["digest"] and "backup" in did and len(sent) == 1
    did2 = desk_scheduler.run_once(db, edition=r.SOLO, now=t8 + datetime.timedelta(minutes=10),
                                   send_digest=sent.append)
    assert "digest" not in did2 and "backup" not in did2 and len(sent) == 1


def test_scheduler_alerts_after_two_failed_backups(db, monkeypatch):
    def boom(*a, **k):
        raise OSError("disk full")
    monkeypatch.setattr(desk_backup, "backup_now", boom)
    alerts = []
    desk_scheduler.run_once(db, edition=r.SOLO, now=datetime.datetime(2026, 10, 27, 3, 0).astimezone(),
                            alert_owner=alerts.append)
    assert alerts == []
    desk_scheduler.run_once(db, edition=r.SOLO, now=datetime.datetime(2026, 10, 28, 3, 0).astimezone(),
                            alert_owner=alerts.append)
    assert len(alerts) == 1 and "2 times" in alerts[0]


# ── auto-close: Fixed + 3 quiet days -> Closed (spec) ────────────────────────

def _fixed_customer_ticket(db, edition=r.SOLO, title="Streaks"):
    c = desk_db.upsert_contact(db, name="Mrs Lopez", phone="386-555-0100")
    actor = None if edition == r.SOLO else OWNER_S
    t = desk_db.create_ticket(db, title=title, kind="external", contact_id=c["contact_id"],
                              actor=actor, edition=edition,
                              **({} if edition == r.SOLO else {"assignee": "dave"}))
    desk_db.change_state(db, t["ticket_id"], r.IN_PROGRESS, actor=actor, edition=edition)
    desk_db.change_state(db, t["ticket_id"], r.FIXED, actor=actor, edition=edition)
    return t["ticket_id"], c


def test_auto_close_after_three_quiet_days(db):
    tid, _c = _fixed_customer_ticket(db)
    own = desk_db.create_ticket(db, title="Renew license", kind="internal", edition=r.SOLO)
    desk_db.change_state(db, own["ticket_id"], r.FIXED, edition=r.SOLO)
    now = datetime.datetime.now(UTC)
    assert desk_reports.auto_close(db, edition=r.SOLO, now=now + datetime.timedelta(days=2, hours=23)) == []
    assert desk_reports.auto_close(db, edition=r.SOLO, now=now + datetime.timedelta(days=3, minutes=1)) == [tid]
    t = desk_db.get_ticket(db, tid)
    assert t["state"] == r.CLOSED and "automatically" in t["close_reason"]
    last = t["state_history"][-1] if "state_history" in t else None
    con = desk_db.get_connection(db)
    try:
        who = con.execute("SELECT actor FROM ticket_events WHERE ticket_id = ? AND kind = 'state_change' "
                          "ORDER BY event_id DESC LIMIT 1", (tid,)).fetchone()[0]
    finally:
        con.close()
    assert who == "system", last
    assert desk_db.get_ticket(db, own["ticket_id"])["state"] == r.FIXED     # the owner's own to-do stays
    assert desk_files.verify(db) == []


def test_a_customer_message_restarts_the_quiet_clock(db):
    tid, c = _fixed_customer_ticket(db)
    cust = {"id": "cust", "role": "customer", "contact_id": c["contact_id"]}
    desk_db.add_message(db, tid, "thanks, looks good so far", actor=cust, edition=r.SOLO, visibility="public")
    later = (datetime.datetime.now(UTC) + datetime.timedelta(days=2)).isoformat(timespec="seconds")
    con = desk_db.get_connection(db)
    try:
        con.execute("UPDATE ticket_messages SET at = ? WHERE ticket_id = ?", (later, tid))
    finally:
        con.close()
    now = datetime.datetime.now(UTC)
    assert desk_reports.auto_close(db, edition=r.SOLO, now=now + datetime.timedelta(days=3, minutes=1)) == []
    assert desk_reports.auto_close(db, edition=r.SOLO, now=now + datetime.timedelta(days=5, minutes=1)) == [tid]


def test_auto_close_zero_days_means_never_and_business_works_too(db):
    desk_db.set_setting(db, "auto_close_days", 0, actor=OWNER_S)
    tid, _c = _fixed_customer_ticket(db, edition=r.BUSINESS)
    far = datetime.datetime.now(UTC) + datetime.timedelta(days=40)
    assert desk_reports.auto_close(db, edition=r.BUSINESS, now=far) == []
    desk_db.set_setting(db, "auto_close_days", 1, actor=OWNER_S)
    did = desk_scheduler.run_once(db, edition=r.BUSINESS, now=far)
    assert did["auto_closed"] == [tid]
    with pytest.raises(desk_db.DeskError):
        desk_db.set_setting(db, "auto_close_days", "3", actor=OWNER_S)       # type is checked
    assert "auto_close_days = 5" in desk_tools.tickets_settings(db, OWNER_S, r.BUSINESS,
                                                                setting="auto_close_days", value="5")


# ── tools ────────────────────────────────────────────────────────────────────

def test_phase3_tools_text(db):
    assert "Saved queue field" in desk_tools.manage_ticket_queues(
        db, OWNER_S, r.BUSINESS, action="save", queue_id="field", name="Field", members="ann, bob")
    assert "ann, bob" in desk_tools.manage_ticket_queues(db, OWNER_S, r.BUSINESS)
    assert "Saved rule" in desk_tools.manage_ticket_rules(
        db, OWNER_S, r.BUSINESS, action="save_rule", rule_id="b", match='{"keyword": "invoice"}',
        send_to_user="dave")
    desk_db.create_ticket(db, title="x", kind="internal", actor=OWNER_S)
    assert "1 ticket(s)" in desk_tools.ticket_report(db, OWNER_S, r.BUSINESS, since="2020-01-01")
    out = desk_tools.export_tickets(db, OWNER_S, r.BUSINESS, format="xlsx")
    assert out.startswith("✅ Exported service-desk-") and ".xlsx" in out
    assert "imports folder" in desk_tools.import_tickets(db, OWNER_S, r.BUSINESS, filename="nope.csv")
    (desk_tools.imports_dir(db) / "a.csv").write_text("Title\nHello\n", encoding="utf-8")
    assert "DRY RUN" in desk_tools.import_tickets(db, OWNER_S, r.BUSINESS, filename="a.csv")
    assert "verified" in desk_tools.tickets_backup_now(db, OWNER_S, r.BUSINESS)
    assert ".zip" in desk_tools.list_tickets_backups(db, OWNER_S, r.BUSINESS)
    assert "backup_hour = 3" in desk_tools.tickets_backup_settings(db, OWNER_S, r.BUSINESS, backup_hour="3")
    assert desk_tools.tickets_backup_now(db, CREW, r.BUSINESS).startswith("❌")
